"""Render Korean workbooks from an immutable export snapshot.

Every sheet is written in everyday Korean and is also what the app previews.
"""
import math
from io import BytesIO

from .display import MISSING_CELLS, REVIEW_CELLS
from .overview import overview
from .plain import kst, plain_sheets
from .statutory_forms import _append_statutory_forms

FONT = '맑은 고딕'
DARK, SOFT, MISSING_FILL, REVIEW_FILL, GRAY = '1F3A5F', 'F4F6F8', 'FCE4D6', 'FFF2CC', '5A6578'
PREVIEW_PHOTO_ROWS = 14


def _units(text):
    """Rough display width of a text in Excel column units (a Korean letter takes two)."""
    return sum(2 if ord(char) > 127 else 1 for char in text)


def _lines(text, width):
    width = max(4, width - 2)
    return sum(max(1, math.ceil(_units(part) / width)) for part in str(text).split('\n'))


def _write_sheet(wb, sheet, subtitle):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet(sheet['name'])
    if sheet.get('tab_color'):
        ws.sheet_properties.tabColor = sheet['tab_color']
    ws.append([f"LOCO 작업 자세 분석 보고서 — {sheet['title']}"])
    ws['A1'].font = Font(name=FONT, size=17, bold=True, color=DARK)
    for text in (sheet.get('lead'), subtitle):
        if text:
            ws.append([text])
            ws.cell(ws.max_row, 1).font = Font(name=FONT, size=11, color=GRAY)
    ws.append([])
    widths = {}
    for table in sheet['tables']:
        for index, width in enumerate(table.get('widths') or [], 1):
            widths[index] = max(widths.get(index, 0), width)
    ncols = max([len(t['columns']) for t in sheet['tables']] + [1])
    plain_table = len(sheet['tables']) == 1 and not sheet['tables'][0].get('heading')
    header_row = last_row = None
    for table in sheet['tables']:
        if table.get('heading'):
            ws.append([table['heading']])
            ws.cell(ws.max_row, 1).font = Font(name=FONT, size=13, bold=True, color=DARK)
        ws.append(table['columns'])
        start = ws.max_row
        header_row = header_row or start
        for column in range(1, len(table['columns']) + 1):
            cell = ws.cell(start, column)
            cell.font = Font(name=FONT, size=11, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor=DARK)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        for number, row in enumerate(table['rows'], 1):
            ws.append(row)
            r = ws.max_row
            tallest = 1
            for column in range(1, len(row) + 1):
                cell = ws.cell(r, column)
                cell.font = Font(name=FONT, size=11, color='25324A')
                cell.fill = PatternFill('solid', fgColor=SOFT if number % 2 == 0 else 'FFFFFF')
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.value in MISSING_CELLS:
                    cell.fill = PatternFill('solid', fgColor=MISSING_FILL)
                elif cell.value in REVIEW_CELLS:
                    cell.fill = PatternFill('solid', fgColor=REVIEW_FILL)
                if isinstance(cell.value, str) and cell.value:
                    tallest = max(tallest, _lines(cell.value, widths.get(column, 16)))
            ws.row_dimensions[r].height = min(400, max(20, 15.5 * tallest + 4))
        last_row = ws.max_row
        ws.append([])
    for note in sheet.get('notes') or []:
        ws.append([note])
        r = ws.max_row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        ws.cell(r, 1).font = Font(name=FONT, size=10, italic=True, color=GRAY)
        ws.cell(r, 1).alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[r].height = max(18, 14 * _lines(note, sum(widths.get(i, 16) for i in range(1, ncols + 1))) + 4)
    for index in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(index)].width = widths.get(index, 16)
    if plain_table and header_row:
        ws.freeze_panes = ws.cell(header_row + 1, 1)
        if last_row and last_row > header_row:
            ws.auto_filter.ref = f'A{header_row}:{get_column_letter(ncols)}{last_row}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def _add_representatives(wb, snapshot, ov):
    """The representative photo goes under the posture tables, only when the person asked for images."""
    if '사진 자세' not in wb.sheetnames:
        return
    from openpyxl.drawing.image import Image as ExcelImage
    from openpyxl.styles import Font
    from ergo.engine.images import read_image
    ws = wb['사진 자세']
    analyses = {a['task_id']: a for a in snapshot.get('analyses', [])}
    for item in ov['tasks']:
        a = analyses.get(item['task_id'])
        representative = ((a or {}).get('reba') or {}).get('summary', {}).get('representative_slot_no')
        if a is None or representative is None:
            continue
        source = a['snapshot']
        slot = next((s for s in (source.get('capture') or {}).get('slots', []) if s.get('slot_no') == representative), None)
        if slot is None:
            continue
        row = ws.max_row + 2
        ws.cell(row, 1, f"대표 사진 — {item['task_name']} · {slot.get('label') or '사진 ' + str(representative)}").font = Font(name=FONT, size=13, bold=True, color=DARK)
        if not snapshot.get('include_representative'):
            ws.cell(row + 1, 1, '대표 사진 이미지는 포함하지 않았어요.')
            continue
        photo = next((p for p in source.get('photos', []) if p['photo_id'] == slot['photo_id']), None)
        try:
            if not photo:
                raise FileNotFoundError('대표 사진 파일 정보를 찾지 못했어요')
            with read_image(photo['file_path']) as image:
                image.thumbnail((300, 240))
                stream = BytesIO()
                image.save(stream, format='PNG')
                stream.seek(0)
            ws.add_image(ExcelImage(stream), f'A{row + 1}')
            for _ in range(PREVIEW_PHOTO_ROWS):
                ws.append([])
        except Exception as error:
            ws.cell(row + 1, 1, f'대표 이미지를 넣지 못했어요: {error}')


def render(snapshot, destination):
    from openpyxl import Workbook
    ov = overview(snapshot)
    survey = snapshot['survey']
    wb = Workbook()
    wb.remove(wb.active)
    subtitle = f"조사: {survey.get('title') or ''}  ·  조사일 {kst(survey.get('survey_date'), True) or ''}  ·  출력 {kst(snapshot.get('exported_at')) or ''}"
    for sheet in plain_sheets(snapshot, ov):
        _write_sheet(wb, sheet, None if sheet['name'] == '요약' else subtitle)
    _add_representatives(wb, snapshot, ov)
    if survey.get('survey_type') == 'OCCASIONAL' and survey.get('occasional_cause') == 'DISEASE_CASE':
        _append_statutory_forms(wb, snapshot)
    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    wb.save(destination)
    return dict(file_path=destination, sheets=list(wb.sheetnames), overview=ov)
