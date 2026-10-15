"""Currentness-aware v2 overview and four-sheet report; legacy snapshots remain readable."""
import json
from datetime import datetime, timedelta
from io import BytesIO
from .labels import input_label, reason_label, investigator_text

NOTICE='작업 분석 보조 보고서. 법정 유해요인조사표·증상조사표의 완성을 뜻하지 않습니다.'
HEADERS={
 '요약':['task_id','task_name','analysis_id','currentness','review_status','applicable_criteria','needs_input_count','needs_review_count','first_next_input'],
 '작업별 결과':['task_id','analysis_id','section','item_code','label','value','unit','source','evidence_text','observed_at'],
 '확인사항':['task_id','analysis_id','module','slot_no','criterion_no','kind','message','source','reviewer_memo'],
 '판정 근거':['task_id','analysis_id','criterion_no','computed_status','investigator_status','effective_status','exclusion','comparison_values','missing_fields','review_reasons','source','investigator_reason','investigator_evidence','reviewer','reviewed_at']}
HEADER_LABELS={
 '요약':['작업 ID','작업명','분석 ID','결과 현재성','검토 상태','해당 호','입력 필요 호 수','검토 필요 호 수','첫 다음 입력'],
 '작업별 결과':['작업 ID','분석 ID','구역','항목 코드','항목명','값','단위','출처','근거 내용','관찰 시각'],
 '확인사항':['작업 ID','분석 ID','모듈','사진 행 번호','기준 호','종류','안내','출처','조사자 메모'],
 '판정 근거':['작업 ID','분석 ID','기준 호','계산 초안','조사자 결론','유효 결과','제외 사유','비교값','누락 입력','검토 사유','규칙 출처','조사자 사유','조사자 근거','검토자','검토 시각']}

def currentness(task, analysis):
    return 'NONE' if analysis is None else 'CURRENT' if task['input_revision']==analysis['input_revision'] and task['latest_analysis_id']==analysis['analysis_id'] else 'STALE'

def overview(snapshot):
    tasks=[]; survey=snapshot['survey']; review=survey.get('review') or {}
    for task in snapshot['tasks']:
        a=next((x for x in snapshot['analyses'] if x['task_id']==task['task_id']),None)
        source=a['snapshot'] if a else task; cur=currentness(task,a)
        pose=a.get('pose') if a else None; tags=a.get('tags') if a else None; reba=a.get('reba') if a else None; judgment=a.get('judgment') if a else None
        judgments=(judgment or {}).get('judgments',[]); slots=(source.get('capture') or {}).get('slots',[]); v2=source.get('schema_version')=='2.0'
        targets=[s for s in slots if s.get('photo_id') is not None and s.get('analysis_target',not v2)]
        scope=source.get('analysis_scope') or dict(judgment=True,pose=True,tags=True,reba=True)
        requested={'pose':scope.get('pose',True),'tags':scope.get('tags',True),'reba':scope.get('reba',True),'judgment':scope.get('judgment',True)}
        values={'pose':pose,'tags':tags,'reba':reba,'judgment':judgment}; stored=(a or {}).get('module_states') or {}; states={}
        for module,value in values.items():
            if not requested[module]: state='NOT_REQUESTED'
            elif not a: state='NOT_CALCULATED'
            elif stored.get(module): state=stored[module]
            elif value is None: state='NOT_CALCULATED'
            elif module=='pose': state='COMPLETE' if value.get('status')=='DONE' else 'FAILED' if value.get('status')=='FAILED' else 'INCOMPLETE'
            elif module=='reba': state='COMPLETE' if value.get('status')=='COMPLETE' else 'INCOMPLETE'
            elif module=='tags': state='COMPLETE' if targets and all(x.get('observed_count')==len(targets) for x in value) else 'INCOMPLETE'
            else: state='INCOMPLETE' if any(j.get('computed_status') in ('NEEDS_INPUT','NEEDS_REVIEW') for j in (value or {}).get('judgments',[])) else 'COMPLETE'
            states[module]=state
        burden_countable=bool(a and requested['judgment'] and judgment is not None)
        reba_countable=bool(a and requested['reba'] and reba is not None)
        counts=dict(needs_input_count=sum(j['computed_status']=='NEEDS_INPUT' for j in judgments) if burden_countable else None,
                    needs_review_count=sum(j['computed_status']=='NEEDS_REVIEW' for j in judgments) if burden_countable else None,
                    incomplete_reba_count=max(0,len(targets)-(reba or {}).get('summary',{}).get('complete_count',0)) if reba_countable else None)
        reviewed=bool(a and review.get('reviewed') and a['analysis_id'] in (review.get('reviewed_analysis_ids') or []))
        incomplete=not a or any(requested[m] and states[m]!='COMPLETE' for m in requested)
        state=dict(reviewed=reviewed,review_current=reviewed and review.get('reviewed_metadata_revision')==survey.get('metadata_revision') and cur=='CURRENT',has_incomplete=incomplete,has_stale=cur=='STALE')
        next_inputs=[]
        def add(module,kind,message,slot=None,criterion=None,fields=None,reasons=None):
            next_inputs.append(dict(module=module,kind=kind,slot_no=slot,criterion_no=criterion,field_paths=fields or [],reason_codes=reasons or [],message=message))
        if not a: add('ANALYSIS','NOT_CALCULATED','분석을 아직 수행하지 않았습니다')
        else:
            labels={'pose':'MoveNet','tags':'자세 기술태그','reba':'REBA','judgment':'11개 기준'}
            kinds={'NOT_CALCULATED':'NOT_CALCULATED','RUNNING':'NOT_CALCULATED','FAILED':'REVIEW_REQUIRED','INCOMPLETE':'REVIEW_REQUIRED','CANCELED':'REVIEW_REQUIRED'}
            for module,label in labels.items():
                module_state=states[module]
                if requested[module] and module_state!='COMPLETE' and module_state!='NOT_REQUESTED':
                    kind=kinds.get(module_state,'NOT_CALCULATED')
                    if module_state in ('NOT_CALCULATED','RUNNING'): add(module.upper(),'NOT_CALCULATED',label+' 계산이 아직 완료되지 않았습니다')
                    elif module=='tags' and not targets: add('TAGS','MISSING_INPUT','자세 기술태그를 계산할 분석대상 사진이 없습니다')
                    elif module=='reba' and not targets: add('REBA','MISSING_INPUT','REBA를 계산할 분석대상 사진이 없습니다')
                    else: add(module.upper(),kind,label+' 결과가 미완료 상태입니다')
            if requested['pose'] and not targets: add('POSE','MISSING_INPUT','분석 대상으로 선택한 사진이 없습니다')
            for slot in targets:
                no=slot['slot_no']; label=slot.get('label') or f'사진 {no}'
                p=next((x for x in (pose or {}).get('slots',[]) if x['slot_no']==no),None); r=next((x for x in (reba or {}).get('slots',[]) if x['slot_no']==no),None)
                if requested['reba'] and r and r.get('missing_fields'): add('REBA','MISSING_INPUT',label+': '+', '.join(input_label(f) for f in r['missing_fields'])+' 입력 필요',no,fields=r['missing_fields'])
                if requested['pose'] and p and p['status']!='VALID': add('POSE','REVIEW_REQUIRED',label+': '+' / '.join(d['message'] for d in p.get('diagnostics',[])),no,fields=list(dict.fromkeys(d['field_path'] for d in p.get('diagnostics',[]) if d.get('field_path'))),reasons=list(dict.fromkeys(d['code'] for d in p.get('diagnostics',[]))))
                if requested['reba'] and r:
                    reasons=list(dict.fromkeys(r.get('left',{}).get('review_reasons',[])+r.get('right',{}).get('review_reasons',[])))
                    if reasons: add('REBA','REVIEW_REQUIRED',label+': '+', '.join(reason_label(x) for x in reasons),no,reasons=reasons)
            for j in judgments:
                if j.get('missing_fields'): add('BURDEN','MISSING_INPUT',f"{j['criterion_no']}호: "+', '.join(input_label(x) for x in j['missing_fields'])+' 입력 필요',criterion=j['criterion_no'],fields=j['missing_fields'])
                if j.get('review_reasons'): add('BURDEN','REVIEW_REQUIRED',f"{j['criterion_no']}호: "+', '.join(reason_label(x) for x in j['review_reasons']),criterion=j['criterion_no'],reasons=j['review_reasons'])
            for note in a.get('notes',[]): add('ANALYSIS','REVIEW_REQUIRED',note.get('message','분석 진단'),reasons=[note.get('code')] if note.get('code') else [])
            for ai in a.get('ai') or []:
                if ai.get('result') in ('FALLBACK','DISABLED'): add('EXPLANATION','REVIEW_REQUIRED','GPT 설명 '+('비활성' if ai.get('result')=='DISABLED' else '대체 문구')+': '+str(ai.get('fallback_reason') or '설정 또는 요청 상태 확인'))
        tasks.append(dict(task_id=task['task_id'],task_name=source['task_name'],analysis_id=a['analysis_id'] if a else None,input_revision=source['input_revision'],source_schema_version=source.get('schema_version','1.3'),module_states=states if a else None,currentness=cur,state=state,review_status='NO_ANALYSIS' if not a else 'UNREVIEWED' if not reviewed else 'REVIEWED_CURRENT' if state['review_current'] else 'REVIEWED_HISTORY',counts=counts,next_inputs=next_inputs))
    state={key:bool(tasks) and all(t['state'][key] for t in tasks) if key in ('reviewed','review_current') else any(t['state'][key] for t in tasks) for key in ('reviewed','review_current','has_incomplete','has_stale')}
    survey_followups=[]
    if survey.get('survey_type')=='OCCASIONAL' and not survey.get('occasional_cause'):
        survey_followups.append('기존 수시조사의 세부 사유 유형을 확인하세요.')
    participation=(survey.get('participation') or {})
    if participation.get('status')!='YES': survey_followups.append('근로자대표 또는 해당 작업 근로자의 참여 여부·기록을 확인하세요.')
    elif participation.get('role') in (None,'UNKNOWN') or participation.get('method') in (None,'UNKNOWN') or not participation.get('participant_code'):
        survey_followups.append('참여자 역할·방법·식별 코드/표시명을 보완하세요.')
    situation=survey.get('workplace_situation') or {}
    for key,label in [('equipment_status','작업설비'),('workload_status','작업량'),('speed_status','작업속도'),('work_pattern_status','업무형태')]:
        if situation.get(key,'UNKNOWN')=='UNKNOWN': survey_followups.append(f'{label}의 변경 여부와 주요 내용을 확인하세요.')
        elif situation.get(key)=='YES' and not situation.get(key.replace('_status','_notes')):
            survey_followups.append(f'{label}이 변경되었다고 기록했습니다. 주요 변경 내용을 입력하세요.')
    symptom=survey.get('symptom_survey') or {}
    if symptom.get('method') in (None,'UNKNOWN','NOT_DONE') or not symptom.get('responses') and (not symptom.get('respondent_count') or not symptom.get('summary')):
        survey_followups.append('작업 관련 근골격계 징후·증상 조사 방법과 결과를 기록하세요.')
    for response in symptom.get('responses') or []:
        if response.get('work_related_symptoms_last_year')=='YES' and not any(
            (response.get('regions') or {}).get(site,{}).get('side') not in (None,'UNKNOWN','NONE')
            for site in ('neck','shoulder','arm_elbow','hand_wrist','back','leg_foot')
        ):
            survey_followups.append(f"{response.get('response_code','익명 응답')}: 작업 관련 증상이 있다고 답한 신체 부위를 입력하세요.")
    state['survey_input_complete']=not survey_followups
    counts={key:sum(t['counts'][key] for t in tasks) if all(t['counts'][key] is not None for t in tasks) else None for key in ('needs_input_count','needs_review_count','incomplete_reba_count')}
    comparisons=[]
    analyses={x['analysis_id']:x for x in snapshot.get('analyses',[])+snapshot.get('comparison_analyses',[])}
    for spec in snapshot.get('comparison_specs',[]):
        before=analyses.get(spec.get('before_analysis_id')); after=analyses.get(spec.get('after_analysis_id'))
        if before is None or after is None: continue
        bset=before['snapshot']; aset=after['snapshot']; bsettings=bset.get('settings') or {}; asettings=aset.get('settings') or {}
        rule_ok=bool(bsettings.get('burden_rule') and asettings.get('burden_rule') and bsettings.get('reba_rule') and asettings.get('reba_rule') and bsettings.get('burden_rule')==asettings.get('burden_rule') and bsettings.get('reba_rule')==asettings.get('reba_rule'))
        bslots=[s for s in (bset.get('capture') or {}).get('slots',[]) if s.get('photo_id') is not None and s.get('analysis_target',bset.get('schema_version')!='2.0')]
        aslots=[s for s in (aset.get('capture') or {}).get('slots',[]) if s.get('photo_id') is not None and s.get('analysis_target',aset.get('schema_version')!='2.0')]
        br=before.get('reba') or {}; ar=after.get('reba') or {}
        sampling=(bset.get('capture',{}).get('protocol')==aset.get('capture',{}).get('protocol') and len(bslots)>0 and len(bslots)==len(aslots) and br.get('status')=='COMPLETE' and ar.get('status')=='COMPLETE' and br.get('summary',{}).get('complete_count')==len(bslots) and ar.get('summary',{}).get('complete_count')==len(aslots))
        delta=(ar.get('summary',{}).get('full_max')-br.get('summary',{}).get('full_max')) if rule_ok and sampling and br.get('summary',{}).get('full_max') is not None and ar.get('summary',{}).get('full_max') is not None else None
        bj={j['criterion_no']:j for j in (before.get('judgment') or {}).get('judgments',[])}; aj={j['criterion_no']:j for j in (after.get('judgment') or {}).get('judgments',[])}
        criterion_diff=[dict(criterion_no=n,before_computed_status=bj.get(n,{}).get('computed_status'),before_investigator_status=bj.get(n,{}).get('investigator',{}).get('status'),before_status=bj.get(n,{}).get('effective_status'),after_computed_status=aj.get(n,{}).get('computed_status'),after_investigator_status=aj.get(n,{}).get('investigator',{}).get('status'),after_status=aj.get(n,{}).get('effective_status')) for n in range(1,12)]
        before_task=before['task_id']; after_task=after['task_id']; changes=[]
        for label,obj1,obj2 in [('conditions',bset.get('conditions') or {},aset.get('conditions') or {}),('analysis_scope',bset.get('analysis_scope') or {},aset.get('analysis_scope') or {}),('legal',bset.get('legal') or {},aset.get('legal') or {})]:
            flat1=dict(flatten(obj1,label)); flat2=dict(flatten(obj2,label))
            for path in sorted(set(flat1)|set(flat2)):
                if flat1.get(path)!=flat2.get(path): changes.append(dict(field_path=path,before=flat1.get(path),after=flat2.get(path)))
        notes=[]
        if not rule_ok: notes.append('고시 또는 REBA 규칙 버전이 달라 점수 비교가 제한됩니다.')
        if not sampling: notes.append('촬영 프로토콜·대상 수·전체 평가 완성이 맞지 않아 REBA 최대값 차이를 계산하지 않았습니다.')
        if before_task!=after_task: notes.append('서로 다른 작업을 비교합니다. 작업 조건과 표본 대표성을 함께 확인하세요.')
        if changes: notes.append('작업 조건 또는 분석 범위가 달라졌습니다.')
        notes.append('전후 비교는 표본 자세 변화의 참고값이며 질환 감소나 인과 효과를 입증하지 않습니다.')
        comparisons.append(dict(spec=spec,before_task_id=before_task,after_task_id=after_task,rule_compatible=rule_ok,sampling_compatible=sampling,reba_max_delta=delta,criteria=criterion_diff,condition_changes=changes,notes=notes))
    return dict(report_version='2.0',state=state,counts=counts,tasks=tasks,comparisons=comparisons,survey_followups=survey_followups)

def flatten(value,prefix=''):
    if isinstance(value,dict):
        for key,child in value.items(): yield from flatten(child,(prefix+'.' if prefix else '')+key)
    elif isinstance(value,list):
        for i,child in enumerate(value):
            identity=child.get('slot_no',child.get('criterion_no',child.get('condition_code',child.get('row_id',i)))) if isinstance(child,dict) else i
            if isinstance(child,dict) and 'scope' in child: identity=str(identity)+'.'+child['scope']+'.'+child.get('field','')
            yield from flatten(child,prefix+'.'+str(identity))
    else: yield prefix,value

def render(snapshot,destination):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.image import Image as ExcelImage
    from ergo.engine.images import read_image
    ov=overview(snapshot); rows={key:[] for key in HEADERS}; survey=snapshot['survey']; representatives=[]
    labels={'worker_code':'작업자 코드','method_description':'작업 방법','work_posture':'작업 자세','work_height_cm':'작업 높이','daily_work_min':'하루 실제 작업시간','one_off':'일회성 여부','planned_duration_months':'예정 작업기간','annual_work_days':'연간 작업일수','condition_present':'해당 작업조건 존재','unsupported':'지지 여부','posture_restricted':'자세 변경 제한','pinch_mass_kg':'손가락 집기 중량','pinch_force_kg_equiv':'손가락 힘','single_hand_mass_kg':'한 손 중량','grip_force_kg_equiv':'쥐기 힘','load_mass_kg':'물체 중량','lift_count_day':'하루 들기 횟수','below_knee_count_day':'무릎 아래 횟수','above_shoulder_count_day':'어깨 위 횟수','extended_arm_count_day':'팔 뻗음 횟수','union_count_day':'중복 제거 횟수','lift_rate_min':'분당 들기 횟수','impact_rate_hour':'시간당 충격 횟수','minutes':'하루 노출시간','measured_load_kg':'측정 하중','analysis_target':'분석대상','camera_side':'카메라 측면','forward_direction':'작업자 전방','side_view_verified':'측면 확인','single_person':'단일 작업자','full_frame':'전신 확인','normal_work':'통상 작업','total':'합계','risk':'위험수준'}
    def unit_for(code):
        leaf=code.split('.')[-1]
        if leaf.endswith('value_deg'): return 'DEG'
        if leaf.endswith(('_px','x_px','y_px')): return 'PX'
        if leaf in ('confidence','min_confidence','sample_ratio'): return 'RATIO'
        if leaf in ('minutes','daily_work_min'): return 'MIN_DAY'
        if leaf in ('load_mass_kg','measured_load_kg','effective_load_kg','box_mass_kg','pinch_mass_kg','single_hand_mass_kg'): return 'KG'
        if leaf in ('pinch_force_kg_equiv','grip_force_kg_equiv'): return 'KG_EQUIV'
        if leaf.endswith('count_day') or leaf=='lift_count_day': return 'COUNT_DAY'
        if leaf=='lift_rate_min': return 'COUNT_MIN'
        if leaf=='impact_rate_hour': return 'COUNT_HOUR'
        if leaf=='planned_duration_months': return 'MONTH'
        if leaf=='annual_work_days': return 'COUNT_YEAR'
        return None
    def result_rows(tid,aid,revision,group,value):
        if value is None: return
        flat=list(flatten(value,group)); evidence_map={}
        for code,v in flat:
            for suffix,field in [('.evidence.source','source'),('.evidence.text','text'),('.evidence.observed_at','observed_at')]:
                if code.endswith(suffix): evidence_map.setdefault(code[:-len(suffix)],{})[field]=v
        for code,val in flat:
            if val is None: rendered='미입력 또는 미계산'
            elif isinstance(val,(dict,list)): rendered=json.dumps(val,ensure_ascii=False,separators=(',',':'))
            else: rendered=val
            candidates=[path for path in evidence_map if code.startswith(path+'.')]
            ev=evidence_map[max(candidates,key=len)] if candidates else {}
            label=labels.get(code.split('.')[-1],input_label(code))
            chunks=[str(rendered)[i:i+30000] for i in range(0,len(str(rendered)),30000)] if isinstance(rendered,str) and len(rendered)>30000 else [rendered]
            for index,chunk in enumerate(chunks):
                item_code=code+(f'#part{index+1}' if len(chunks)>1 else '')
                rows['작업별 결과'].append([tid,aid,group,item_code,label+(' · 계속' if index else ''),chunk,unit_for(code),ev.get('source'),ev.get('text'),ev.get('observed_at')])
    for item in ov['tasks']:
        a=next((x for x in snapshot['analyses'] if x['task_id']==item['task_id']),None); task=next(x for x in snapshot['tasks'] if x['task_id']==item['task_id']); source=a['snapshot'] if a else task
        pose=a.get('pose') if a else None; reba=a.get('reba') if a else None; judgment=a.get('judgment') if a else None; judgments=(judgment or {}).get('judgments',[]); aid=item['analysis_id']; tid=item['task_id']
        applicable=', '.join(str(j['criterion_no'])+'호' for j in judgments if j.get('effective_status')=='APPLICABLE') if judgment else None
        rows['요약'].append([tid,item['task_name'],aid,item['currentness'],item['review_status'],applicable,item['counts']['needs_input_count'],item['counts']['needs_review_count'],item['next_inputs'][0]['message'] if item['next_inputs'] else None])
        for group,key,value in [('작업조건','conditions',source.get('conditions')),('11개기준 입력','legal',source.get('legal')),('사진·관찰','capture',source.get('capture')),('수동 관찰','manual',source.get('manual')),('분석 설정','settings',source.get('settings',{})),('모듈 상태','module_states',(a or {}).get('module_states')),('MoveNet','pose',pose),('자세 기술태그','tags',dict(summary=(a or {}).get('tags'),slots=(a or {}).get('slot_tags')) if a else None),('REBA','reba',reba),('GPT 설명','ai',(a or {}).get('ai'))]: result_rows(tid,aid,item['input_revision'],group,value)
        capture=(source.get('capture') or {}).get('slots',[]); v2=source.get('schema_version')=='2.0'; targets=[s for s in capture if s.get('photo_id') is not None and s.get('analysis_target',not v2)]
        for slot in targets:
            no=slot['slot_no']; p=next((x for x in (pose or {}).get('slots',[]) if x['slot_no']==no),{}); r=next((x for x in (reba or {}).get('slots',[]) if x['slot_no']==no),{})
            for group,value in [('사진별 포즈',p),('사진별 REBA',r)]: result_rows(tid,aid,item['input_revision'],f'{group} · 행 {no}',value)
            if p.get('diagnostics'):
                for d in p['diagnostics']: rows['확인사항'].append([tid,aid,'POSE',no,None,'REVIEW_REQUIRED',d.get('message'),d.get('code'),None])
            if r.get('missing_fields'):
                rows['확인사항'].append([tid,aid,'REBA',no,None,'MISSING_INPUT',', '.join(input_label(x) for x in r['missing_fields']),None,None])
            if r.get('left',{}).get('review_reasons') or r.get('right',{}).get('review_reasons'):
                reasons=list(dict.fromkeys(r.get('left',{}).get('review_reasons',[])+r.get('right',{}).get('review_reasons',[])))
                rows['확인사항'].append([tid,aid,'REBA',no,None,'REVIEW_REQUIRED',', '.join(reason_label(x) for x in reasons),None,None])
        for notice in item['next_inputs']:
            rows['확인사항'].append([tid,aid,notice['module'],notice['slot_no'],notice['criterion_no'],notice['kind'],notice['message'],','.join(notice['reason_codes']) or None,None])
        source_criteria={x['criterion_no']:x for x in (source.get('legal') or {}).get('criteria',[])}
        for no in range(1,12):
            j=next((x for x in judgments if x['criterion_no']==no),{}); inv=j.get('investigator',{}); ci=source_criteria.get(no,{})
            comparison='\n'.join(' | '.join(str(b.get(k,'')) for k in ['item_code','value','unit','operator','threshold','source']) for b in j.get('basis',[])) or None
            evidence_data=inv.get('evidence') or {}
            rows['판정 근거'].append([tid,aid,no,j.get('computed_status'),inv.get('status'),j.get('effective_status'),j.get('exclusion'),comparison,'\n'.join(input_label(x) for x in j.get('missing_fields',[])), '\n'.join(reason_label(x) for x in j.get('review_reasons',[])),j.get('source'),inv.get('reason'),evidence_data.get('text') or (ci.get('evidence') or {}).get('text'),inv.get('reviewer'),inv.get('reviewed_at')])
        representatives.append((tid,aid,source,a,reba))
    for comparison in ov['comparisons']:
        spec=comparison['spec']; label=f"비교 #{spec['before_analysis_id']} → #{spec['after_analysis_id']}"
        if spec.get('note'): rows['확인사항'].append([None,None,'COMPARISON',None,None,'NOTE',spec['note'],None,None])
        for note in comparison['notes']: rows['확인사항'].append([None,None,'COMPARISON',None,None,'LIMITATION',note,None,None])
        rows['확인사항'].append([None,None,'COMPARISON',None,None,'SUMMARY',f"{label} · REBA 호환 {comparison['rule_compatible']} · 표본 호환 {comparison['sampling_compatible']} · 후−전 최대 점수 {comparison['reba_max_delta']}",None,None])
        for change in comparison['condition_changes']:
            rows['확인사항'].append([None,None,'COMPARISON',None,None,'CONDITION_CHANGE',f"{change['field_path']}: {change['before']} → {change['after']}",None,None])
        for diff in comparison['criteria']:
            rows['확인사항'].append([None,None,'COMPARISON',None,diff['criterion_no'],'CRITERION_CHANGE',f"{diff['before_status']} → {diff['after_status']} (계산 {diff['before_computed_status']} → {diff['after_computed_status']}; 조사자 {diff['before_investigator_status']} → {diff['after_investigator_status']})",None,None])
    # Keep the investigation context, participation, and symptom-survey coverage visible in the normal report.
    situation=survey.get('workplace_situation') or {}; participation=survey.get('participation') or {}; symptoms=survey.get('symptom_survey') or {}
    for key,label in [('survey_type','조사 구분'),('occasional_cause','수시 사유 유형'),('occasional_reason','수시 사유 상세'),('previous_survey_date','이전 조사일')]:
        if survey.get(key) is not None: rows['작업별 결과'].append([None,None,'조사 개요',key,label,survey.get(key),None,None,None,None])
    situation_labels={'equipment_status':'설비 변경 여부','equipment_notes':'설비 변경 내용','workload_status':'작업량 변경 여부','workload_notes':'작업량 변경 내용','speed_status':'작업속도 변경 여부','speed_notes':'작업속도 변경 내용','work_pattern_status':'업무형태 변경 여부','work_pattern_notes':'업무형태 변경 내용'}
    for key,label in situation_labels.items():
        value=situation.get(key)
        rows['작업별 결과'].append([None,None,'작업장 상황',key,label,value,None,None,None,None])
    for key,label in [('status','근로자 참여 여부'),('role','참여자 역할'),('participant_code','참여자 코드/표시명'),('method','참여 방법'),('date','참여 일자'),('notes','참여 의견')]:
        rows['작업별 결과'].append([None,None,'근로자 참여',key,label,participation.get(key),None,None,None,None])
    for key,label in [('method','증상조사 방법'),('date','증상조사 일자'),('respondent_count','응답자 수'),('summary','증상조사 요약')]:
        rows['작업별 결과'].append([None,None,'증상조사',key,label,symptoms.get(key),None,None,None,None])
    for response in symptoms.get('responses') or []:
        rows['작업별 결과'].append([None,None,'증상 응답',response.get('response_code'),'증상조사 응답',json.dumps(response,ensure_ascii=False,separators=(',',':')),None,'WORKER_QUESTIONNAIRE',None,symptoms.get('date')])

    wb=Workbook(); wb.remove(wb.active); exported=datetime.fromisoformat(snapshot['exported_at'])+timedelta(hours=9); exported=exported.replace(tzinfo=None)
    for message in ov.get('survey_followups',[]):
        rows['확인사항'].append([None,None,'SURVEY',None,None,'MISSING_INPUT',message,'SAFETY_RULES_657_658',None])
    for name,headers in HEADERS.items():
        ws=wb.create_sheet(name); ws.append(['LOCO 작업 자세 분석 보고서 — '+name]); ws['A1'].font=Font(name='맑은 고딕',size=17,bold=True,color='1F3A5F')
        ws.append(['조사 ID',survey['survey_id'],'보고서 버전','2.0','스키마 버전','2.0','출력 시각',exported]); ws['H2'].number_format='yyyy-mm-dd hh:mm'
        ws.append(['조사',survey.get('title'),'조사자',survey.get('investigator'),'부서',survey.get('department'),'공정',survey.get('process')]); ws.append(['조사일',datetime.fromisoformat(survey['survey_date']) if survey.get('survey_date') else None]); ws['B4'].number_format='yyyy-mm-dd'
        ws.append(HEADER_LABELS[name]); start=5
        for row in rows[name]:
            chunks=max([1]+[(len(str(value))+29999)//30000 for value in row if isinstance(value,str)])
            for index in range(chunks):
                part=[value[index*30000:(index+1)*30000] if isinstance(value,str) and len(value)>30000 else value for value in row]
                ws.append(part)
        end=max(start,ws.max_row); ws.auto_filter.ref=f'A{start}:{get_column_letter(len(headers))}{end}'; ws.freeze_panes='D6'; ws.print_title_rows='1:5'; ws.page_setup.orientation='landscape'
        for index,header in enumerate(headers,1): ws.column_dimensions[get_column_letter(index)].width=42 if header in ('message','comparison_values','evidence_text','investigator_evidence','reviewer_memo','first_next_input','label') else 22 if header in ('task_name','item_code') else 15
        for row in ws.iter_rows(min_row=1,max_row=end):
            for cell in row:
                if cell.row==1: continue
                cell.font=Font(name='맑은 고딕',size=11,color='FFFFFF' if cell.row==5 else '25324A',bold=cell.row==5); cell.fill=PatternFill('solid',fgColor='1F3A5F' if cell.row==5 else 'F4F6F8' if cell.row%2==0 else 'FFFFFF'); cell.alignment=Alignment(vertical='top',wrap_text=True)
                if cell.row>5:
                    if cell.value in ('NEEDS_INPUT','INCOMPLETE','MISSING_INPUT'): cell.fill=PatternFill('solid',fgColor='FCE4D6')
                    elif cell.value in ('NEEDS_REVIEW','REVIEW_REQUIRED'): cell.fill=PatternFill('solid',fgColor='FFF2CC')
                    if isinstance(cell.value,str): cell.data_type='s'
            if row[0].row>5: ws.row_dimensions[row[0].row].height=48
        ws.append([]); ws.append([NOTICE]); ws.append(['사진 표본 비율은 하루 노출시간이나 해당 확률이 아닙니다. 미완료 점수는 0점이 아니며, 조사자 결론은 법적 승인을 뜻하지 않습니다.'])
    result_ws=wb['작업별 결과']
    for tid,aid,source,a,reba in representatives:
        rep=(reba or {}).get('summary',{}).get('representative_slot_no')
        result_ws.append([]); result_ws.append([f'작업 {tid} / 분석 {aid} 대표 사진 행 {rep if rep is not None else "없음"}'])
        if rep is None: result_ws.append(['대표 사진 점수 없음']); continue
        slot=next((s for s in source['capture']['slots'] if s['slot_no']==rep),None); result=next((r for r in (reba or {}).get('slots',[]) if r['slot_no']==rep),None)
        if slot is None or result is None: result_ws.append(['대표 사진 상세를 찾을 수 없습니다']); continue
        result_ws.append(['대표 행',rep,'표시명',slot.get('label'),'계산 최대',reba['summary'].get('available_max'),'대상 전체 최대',reba['summary'].get('full_max'),'대표 점수',result.get('total')])
        if snapshot.get('include_representative'):
            photo=next((p for p in source.get('photos',[]) if p['photo_id']==slot['photo_id']),None)
            try:
                if not photo: raise FileNotFoundError('대표 사진 파일 메타데이터가 없습니다')
                with read_image(photo['file_path']) as image:
                    image.thumbnail((300,240)); stream=BytesIO(); image.save(stream,format='PNG'); stream.seek(0)
                row=result_ws.max_row+2; result_ws.add_image(ExcelImage(stream),f'K{row}'); result_ws.cell(row,1,'대표 사진 미리보기')
            except Exception as error: result_ws.append(['대표 이미지 생략',str(error)])
        else: result_ws.append(['대표 이미지 제외'])
    if survey.get('survey_type')=='OCCASIONAL' and survey.get('occasional_cause')=='DISEASE_CASE':
        _append_statutory_forms(wb, snapshot)
    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value,str): cell.data_type='s'
    wb.save(destination)
    return dict(file_path=destination,sheets=list(wb.sheetnames),overview=ov)


def _append_statutory_forms(workbook, snapshot):
    """Add form-oriented appendices only for the Article 657(2)(1) survey trigger."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    survey=snapshot['survey']; situation=survey.get('workplace_situation') or {}; participation=survey.get('participation') or {}; symptom=survey.get('symptom_survey') or {}
    form=workbook.create_sheet('별지1 유해요인조사표')
    form.append(['유해요인 조사표 · 고용노동부고시 제2020-12호 별지 제1호서식 기반'])
    form.append(['조사 구분', '수시조사 · 근골격계질환자 발생/업무상 질병 인정', '조사일', survey.get('survey_date'), '조사자', survey.get('investigator')])
    form.append(['조사부서명', survey.get('department'), '작업공정명', survey.get('process'), '작업명', survey.get('title')])
    form.append(['이전 조사일', survey.get('previous_survey_date'), '조사 사유', survey.get('occasional_reason')])
    form.append(['나. 작업장 상황 조사'])
    form.append(['항목','변경 여부','주요 변경 내용'])
    for key,label in [('equipment','작업 설비'),('workload','작업량'),('speed','작업 속도'),('work_pattern','업무 형태')]:
        form.append([label,situation.get(key+'_status','UNKNOWN'),situation.get(key+'_notes','')])
    form.append(['근로자 참여', participation.get('status'), participation.get('role'), participation.get('participant_code'), participation.get('method'), participation.get('date'), participation.get('notes')])
    form.append(['다. 작업조건 조사 · 작업별 작업방법/작업자세/작업시간 및 11개 기준 근거'])
    form.append(['작업명','작업자 코드','작업방법','작업자세','작업시간(분/일)','작업강도(A)','작업빈도(B)','A×B','11개 기준 계산 초안','관찰 유해요인·발생 원인','개선 제안','근거 메모'])
    analyses={x['task_id']:x for x in snapshot.get('analyses',[])}
    for task in snapshot.get('tasks',[]):
        analysis=analyses.get(task['task_id']); conditions=task.get('conditions') or {}
        statuses=[]
        for item in ((analysis or {}).get('judgment') or {}).get('judgments',[]):
            statuses.append(f"{item.get('criterion_no')}호 {item.get('effective_status') or item.get('computed_status')}")
        intensity=conditions.get('work_intensity'); frequency=conditions.get('work_frequency')
        score=intensity*frequency if isinstance(intensity,int) and isinstance(frequency,int) else None
        hazards=' / '.join(x for x in (conditions.get('hazard_factors'),conditions.get('hazard_causes')) if x)
        form.append([task.get('task_name'),conditions.get('worker_code'),conditions.get('method_description'),conditions.get('work_posture'),conditions.get('daily_work_min'),intensity,frequency,score,', '.join(statuses) or '분석 전/미출력',hazards,conditions.get('improvement_recommendations'),(conditions.get('evidence') or {}).get('text')])
    form.append(['라. 유해요인 조사 결과'])
    form.append(['조사자 결과 요약', symptom.get('summary')])
    form.append(['주의','이 시트는 입력한 값으로 작성한 별지 서식 기반 출력입니다. 빈칸 확인, 서명/날인, 원본 별지 양식과의 대조 및 최종 적합성 검토가 필요합니다.'])

    responses=symptom.get('responses') or []
    symptom_ws=workbook.create_sheet('별지2 증상조사표')
    symptom_ws.append(['근골격계질환 증상조사표 · 고용노동부고시 제2020-12호 별지 제2호서식 기반'])
    symptom_ws.append(['조사',survey.get('title'),'조사일',symptom.get('date'),'조사 방법',symptom.get('method'),'응답자 수',symptom.get('respondent_count')])
    sites=[('neck','목'),('shoulder','어깨'),('arm_elbow','팔/팔꿈치'),('hand_wrist','손/손목/손가락'),('back','허리'),('leg_foot','다리/발')]
    headers=['응답 코드','성명(선택)','연령','성별','현 직장 경력(개월)','작업부서','수행작업','현재 작업 내용','현재 작업 기간(개월)','1일 근무시간','휴식시간','이전 작업','이전 작업 기간(개월)','여가·취미','가사노동','의사진단 이력','사고·상해 이력','상해 부위','육체적 부담(1~5)','지난 1년 작업관련 증상']
    for _,label in sites:
        headers.extend([f'{label} 방향',f'{label} 지속기간',f'{label} 통증정도',f'{label} 지난 1주'])
    symptom_ws.append(headers)
    for response in responses:
        row=[response.get('response_code'),response.get('name'),response.get('age_years'),response.get('sex'),response.get('workplace_tenure_months'),response.get('department'),response.get('job'),response.get('current_work'),response.get('current_work_months'),response.get('daily_work_hours'),response.get('rest_breaks'),response.get('previous_work'),response.get('previous_work_months'),response.get('leisure_activities'),response.get('household_work'),response.get('medical_history'),response.get('injury_history'),', '.join(response.get('injury_sites') or []),response.get('physical_burden'),response.get('work_related_symptoms_last_year')]
        regions=response.get('regions') or {}
        for site,_ in sites:
            detail=regions.get(site) or {}
            row.extend([detail.get('side'),detail.get('duration'),detail.get('severity'),detail.get('last_week')])
        symptom_ws.append(row)
    symptom_ws.append(['기재 요령','응답 코드 사용을 권장합니다. 개인 식별정보·건강정보가 포함된 파일은 접근을 제한하세요. 진단을 대신하거나 질환을 판정하는 설문이 아닙니다.'])

    for ws in (form,symptom_ws):
        ws.freeze_panes='A6' if ws is form else 'B4'
        ws.sheet_view.showGridLines=False
        ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        for row in ws.iter_rows():
            for cell in row:
                cell.font=Font(name='맑은 고딕',size=10,bold=cell.row in (1,5,6) or cell.column==1 and ws is form,color='FFFFFF' if cell.row in (1,5,6) else '25324A')
                cell.fill=PatternFill('solid',fgColor='1F3A5F' if cell.row in (1,5,6) else 'F4F6F8' if cell.row%2==0 else 'FFFFFF')
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        for column in range(1,ws.max_column+1):
            ws.column_dimensions[get_column_letter(column)].width=24 if ws is form else 18
        for row in range(1,ws.max_row+1): ws.row_dimensions[row].height=32 if row in (1,5,6) else 24
