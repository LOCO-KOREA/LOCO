"""Conditional form-based sheets for disease-triggered surveys."""
def _append_statutory_forms(workbook, snapshot):
    """Add form-oriented appendices only for the Article 657(2)(1) survey trigger."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    survey = snapshot['survey']
    situation = survey.get('workplace_situation') or {}
    participation = survey.get('participation') or {}
    symptom = survey.get('symptom_survey') or {}
    form = workbook.create_sheet('별지1 유해요인조사표')
    form.append(['유해요인 조사표 · 고용노동부고시 제2020-12호 별지 제1호서식 기반'])
    form.append(['조사 구분', '수시조사 · 근골격계질환자 발생/업무상 질병 인정', '조사일', survey.get('survey_date'), '조사자', survey.get('investigator')])
    form.append(['조사부서명', survey.get('department'), '작업공정명', survey.get('process'), '작업명', survey.get('title')])
    form.append(['이전 조사일', survey.get('previous_survey_date'), '조사 사유', survey.get('occasional_reason')])
    form.append(['나. 작업장 상황 조사'])
    form.append(['항목', '변경 여부', '주요 변경 내용'])
    for key, label in [('equipment', '작업 설비'), ('workload', '작업량'), ('speed', '작업 속도'), ('work_pattern', '업무 형태')]:
        form.append([label, situation.get(key + '_status', 'UNKNOWN'), situation.get(key + '_notes', '')])
    form.append(['근로자 참여', participation.get('status'), participation.get('role'), participation.get('participant_code'), participation.get('method'), participation.get('date'), participation.get('notes')])
    form.append(['다. 작업조건 조사 · 작업별 작업방법/작업자세/작업시간 및 11개 기준 근거'])
    form.append(['작업명', '작업자 코드', '작업방법', '작업자세', '작업시간(분/일)', '작업강도(A)', '작업빈도(B)', 'A×B', '11개 기준 계산 초안', '관찰 유해요인·발생 원인', '개선 제안', '근거 메모'])
    analyses = {x['task_id']: x for x in snapshot.get('analyses', [])}
    for task in snapshot.get('tasks', []):
        analysis = analyses.get(task['task_id'])
        conditions = task.get('conditions') or {}
        statuses = []
        for item in ((analysis or {}).get('judgment') or {}).get('judgments', []):
            statuses.append(f'{item.get('criterion_no')}호 {item.get('effective_status') or item.get('computed_status')}')
        intensity = conditions.get('work_intensity')
        frequency = conditions.get('work_frequency')
        score = intensity * frequency if isinstance(intensity, int) and isinstance(frequency, int) else None
        hazards = ' / '.join((x for x in (conditions.get('hazard_factors'), conditions.get('hazard_causes')) if x))
        form.append([task.get('task_name'), conditions.get('worker_code'), conditions.get('method_description'), conditions.get('work_posture'), conditions.get('daily_work_min'), intensity, frequency, score, ', '.join(statuses) or '분석 전/미출력', hazards, conditions.get('improvement_recommendations'), (conditions.get('evidence') or {}).get('text')])
    form.append(['라. 유해요인 조사 결과'])
    form.append(['조사자 결과 요약', symptom.get('summary')])
    form.append(['주의', '이 시트는 입력한 값으로 작성한 별지 서식 기반 출력입니다. 빈칸 확인, 서명/날인, 원본 별지 양식과의 대조 및 최종 적합성 검토가 필요합니다.'])
    responses = symptom.get('responses') or []
    symptom_ws = workbook.create_sheet('별지2 증상조사표')
    symptom_ws.append(['근골격계질환 증상조사표 · 고용노동부고시 제2020-12호 별지 제2호서식 기반'])
    symptom_ws.append(['조사', survey.get('title'), '조사일', symptom.get('date'), '조사 방법', symptom.get('method'), '응답자 수', symptom.get('respondent_count')])
    sites = [('neck', '목'), ('shoulder', '어깨'), ('arm_elbow', '팔/팔꿈치'), ('hand_wrist', '손/손목/손가락'), ('back', '허리'), ('leg_foot', '다리/발')]
    headers = ['응답 코드', '성명(선택)', '연령', '성별', '현 직장 경력(개월)', '작업부서', '수행작업', '현재 작업 내용', '현재 작업 기간(개월)', '1일 근무시간', '휴식시간', '이전 작업', '이전 작업 기간(개월)', '여가·취미', '가사노동', '의사진단 이력', '사고·상해 이력', '상해 부위', '육체적 부담(1~5)', '지난 1년 작업관련 증상']
    for _, label in sites:
        headers.extend([f'{label} 방향', f'{label} 지속기간', f'{label} 통증정도', f'{label} 지난 1주'])
    symptom_ws.append(headers)
    for response in responses:
        row = [response.get('response_code'), response.get('name'), response.get('age_years'), response.get('sex'), response.get('workplace_tenure_months'), response.get('department'), response.get('job'), response.get('current_work'), response.get('current_work_months'), response.get('daily_work_hours'), response.get('rest_breaks'), response.get('previous_work'), response.get('previous_work_months'), response.get('leisure_activities'), response.get('household_work'), response.get('medical_history'), response.get('injury_history'), ', '.join(response.get('injury_sites') or []), response.get('physical_burden'), response.get('work_related_symptoms_last_year')]
        regions = response.get('regions') or {}
        for site, _ in sites:
            detail = regions.get(site) or {}
            row.extend([detail.get('side'), detail.get('duration'), detail.get('severity'), detail.get('last_week')])
        symptom_ws.append(row)
    symptom_ws.append(['기재 요령', '응답 코드 사용을 권장합니다. 개인 식별정보·건강정보가 포함된 파일은 접근을 제한하세요. 진단을 대신하거나 질환을 판정하는 설문이 아닙니다.'])
    for ws in (form, symptom_ws):
        ws.freeze_panes = 'A6' if ws is form else 'B4'
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name='맑은 고딕', size=10, bold=cell.row in (1, 5, 6) or (cell.column == 1 and ws is form), color='FFFFFF' if cell.row in (1, 5, 6) else '25324A')
                cell.fill = PatternFill('solid', fgColor='1F3A5F' if cell.row in (1, 5, 6) else 'F4F6F8' if cell.row % 2 == 0 else 'FFFFFF')
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for column in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(column)].width = 24 if ws is form else 18
        for row in range(1, ws.max_row + 1):
            ws.row_dimensions[row].height = 32 if row in (1, 5, 6) else 24
