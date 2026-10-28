"""Render Korean workbooks from an immutable export snapshot.

Every sheet is written in everyday Korean and is also what the app previews. A table may ask for form-style drawing
(no dark header row, label cells, merged cells, borders, fixed row heights); the preview reads the same keys.
"""
import math
from io import BytesIO

from .display import MISSING_CELLS, REVIEW_CELLS
from .overview import overview
from .plain import kst, plain_sheets

FONT = '맑은 고딕'
DARK, SOFT, MISSING_FILL, REVIEW_FILL, GRAY, LABEL_FILL, EDGE = '1F3A5F', 'F4F6F8', 'FCE4D6', 'FFF2CC', '5A6578', 'E6EBF2', 'B3BCC9'


def _units(text):
    """Rough display width of a text in Excel column units (a Korean letter takes two)."""
    return sum(2 if ord(char) > 127 else 1 for char in text)


def _lines(text, width):
    width = max(4, width - 2)
    return sum(max(1, math.ceil(_units(part) / width)) for part in str(text).split('\n'))


def _merge_map(table):
    """{(row, col): (last_row, last_col)} for the first cell of every merged block, and the set of covered cells."""
    starts, covered = {}, set()
    for r1, c1, r2, c2 in table.get('merges') or []:
        starts[(r1, c1)] = (r2, c2)
        covered.update((r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1) if (r, c) != (r1, c1))
    return starts, covered


def _write_sheet(wb, sheet, subtitle, anchors=None):
    """Write one sheet description. `anchors` collects {task_id: (sheet, cell)} for tables that reserve room for a photo."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.pagebreak import Break
    ws = wb.create_sheet(sheet['name'])
    if sheet.get('tab_color'):
        ws.sheet_properties.tabColor = sheet['tab_color']
    cursor = [0]

    def put(values):
        """Append a row and return its number. Counting here keeps merged blocks on the right rows even after an empty spacer row."""
        ws.append(values)
        cursor[0] += 1
        return cursor[0]

    put([f"LOCO 작업 자세 분석 보고서 — {sheet['title']}"])
    ws['A1'].font = Font(name=FONT, size=17, bold=True, color=DARK)
    for text in (sheet.get('lead'), subtitle):
        if text:
            ws.cell(put([text]), 1).font = Font(name=FONT, size=11, color=GRAY)
    put([])
    widths = {}
    for table in sheet['tables']:
        for index, width in enumerate(table.get('widths') or [], 1):
            widths[index] = max(widths.get(index, 0), width)
    ncols = max([len(t['columns']) for t in sheet['tables']] + [1])
    plain_table = len(sheet['tables']) == 1 and not sheet['tables'][0].get('heading') and sheet['tables'][0].get('header', True)
    header_row = last_row = None
    edge = Side(style='thin', color=EDGE)
    for table in sheet['tables']:
        if table.get('break'):
            ws.row_breaks.append(Break(id=cursor[0]))
        if table.get('heading'):
            ws.cell(put([table['heading']]), 1).font = Font(name=FONT, size=13, bold=True, color=DARK)
        bordered = bool(table.get('bordered'))
        columns = len(table['columns'])
        if table.get('header', True):
            start = put(table['columns'])
            header_row = header_row or start
            tallest = 1
            for column in range(1, columns + 1):
                cell = ws.cell(start, column)
                cell.font = Font(name=FONT, size=11, bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor=DARK)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                if bordered:
                    cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
                if isinstance(cell.value, str) and cell.value:
                    tallest = max(tallest, _lines(cell.value, widths.get(column, 16)))
            if tallest > 1:
                ws.row_dimensions[start].height = min(400, 15.5 * tallest + 4)
        labels, label_rows = set(table.get('labels') or []), set(table.get('label_rows') or [])
        label_cells = {tuple(c) for c in table.get('label_cells') or []}
        starts, covered = _merge_map(table)
        heights = {int(k): v for k, v in (table.get('heights') or {}).items()}
        first = cursor[0] + 1
        for number, row in enumerate(table['rows']):
            r = put(row)
            tallest = 1
            for column in range(1, len(row) + 1):
                cell = ws.cell(r, column)
                label = (column - 1) in labels or number in label_rows or (number, column - 1) in label_cells
                cell.font = Font(name=FONT, size=11, bold=label, color=DARK if label else '25324A')
                zebra = SOFT if (number + 1) % 2 == 0 and not bordered else 'FFFFFF'
                cell.fill = PatternFill('solid', fgColor=LABEL_FILL if label else zebra)
                cell.alignment = Alignment(vertical='center' if label else 'top', wrap_text=True)
                if bordered:
                    cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
                if cell.value in MISSING_CELLS:
                    cell.fill = PatternFill('solid', fgColor=MISSING_FILL)
                elif cell.value in REVIEW_CELLS:
                    cell.fill = PatternFill('solid', fgColor=REVIEW_FILL)
                if isinstance(cell.value, str) and cell.value and (number, column - 1) not in covered:
                    last_column = starts.get((number, column - 1), (0, column - 1))[1]
                    wide = sum(widths.get(i + 1, 16) for i in range(column - 1, last_column + 1))
                    tallest = max(tallest, _lines(cell.value, wide))
            ws.row_dimensions[r].height = heights.get(number, min(400, max(20, 15.5 * tallest + 4)))
            photo = table.get('photo')
            if photo and anchors is not None and number == photo.get('row', 0):
                anchors[photo['task_id']] = (ws, f'A{r}')
        for r1, c1, r2, c2 in table.get('merges') or []:
            ws.merge_cells(start_row=first + r1, start_column=c1 + 1, end_row=first + r2, end_column=c2 + 1)
        last_row = cursor[0]
        put([])
    for note in sheet.get('notes') or []:
        r = put([note])
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        ws.cell(r, 1).font = Font(name=FONT, size=10, italic=True, color=GRAY)
        ws.cell(r, 1).alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[r].height = max(18, 14 * _lines(note, sum(widths.get(i, 16) for i in range(1, ncols + 1))) + 4)
    for index in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(index)].width = widths.get(index, 16)
    if plain_table and header_row:
        ws.freeze_panes = ws.cell(header_row + 1, 1)
        if last_row and last_row > header_row:
            ws.auto_filter.ref = f'A{header_row}:{get_column_letter(ncols)}{last_row}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def _add_representatives(snapshot, ov, anchors):
    """The representative photo goes into the room the form reserved for it, only when the person asked for images."""
    if not snapshot.get('include_representative') or not anchors:
        return
    from openpyxl.drawing.image import Image as ExcelImage
    from backend.engine.images import read_image
    analyses = {a['task_id']: a for a in snapshot.get('analyses', [])}
    for item in ov['tasks']:
        a = analyses.get(item['task_id'])
        representative = ((a or {}).get('reba') or {}).get('summary', {}).get('representative_slot_no')
        if a is None or representative is None:
            continue
        source = a['snapshot']
        slot = next((s for s in (source.get('capture') or {}).get('slots', []) if s.get('slot_no') == representative), None)
        target = anchors.get(item['task_id'])
        if slot is None or target is None:
            continue
        ws, cell = target
        photo = next((p for p in source.get('photos', []) if p['photo_id'] == slot['photo_id']), None)
        try:
            if not photo:
                raise FileNotFoundError('대표 사진 파일 정보를 찾지 못했어요')
            with read_image(photo['file_path']) as image:
                image.thumbnail((260, 200))
                stream = BytesIO()
                image.save(stream, format='PNG')
                stream.seek(0)
            ws.add_image(ExcelImage(stream), cell)
        except Exception as error:
            ws[cell] = f'대표 이미지를 넣지 못했어요: {error}'


def render(snapshot, destination):
    from openpyxl import Workbook
    ov = overview(snapshot)
    survey = snapshot['survey']
    wb = Workbook()
    wb.remove(wb.active)
    subtitle = f"조사: {survey.get('title') or ''}  ·  조사일 {kst(survey.get('survey_date'), True) or ''}  ·  출력 {kst(snapshot.get('exported_at')) or ''}"
    anchors = {}
    for sheet in plain_sheets(snapshot, ov):
        _write_sheet(wb, sheet, subtitle, anchors)
    _add_representatives(snapshot, ov, anchors)
    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    wb.save(destination)
    return dict(file_path=destination, sheets=list(wb.sheetnames), overview=ov)
