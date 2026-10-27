"""PPTX 글자 넘침·겹침 점검기 (PowerPoint/LibreOffice 가 없는 PC 에서 쓰는 대용 검사).

맑은 고딕 글꼴 파일(C:/Windows/Fonts/malgun*.ttf)의 실제 글자 폭으로, 공백 단위(어절) 줄바꿈을 흉내 낸다.
- 칸 높이를 넘는 글자, 한 어절이 칸 폭보다 긴 경우, 슬라이드 밖으로 나간 도형, 글상자끼리 겹침을 보고한다.
- 실제 PowerPoint 와 줄 나눔이 조금 다를 수 있으므로 '여유율'을 둔다. 결과는 참고용이다.

사용: python audit_fit.py deck.pptx [--slack 0.0]
"""
import sys
import re
from pptx import Presentation
from pptx.util import Emu
from PIL import ImageFont

FONT_REG = 'C:/Windows/Fonts/malgun.ttf'
FONT_BOLD = 'C:/Windows/Fonts/malgunbd.ttf'
SCALE = 20  # 폭 측정 정밀도(글자 크기 pt × SCALE 픽셀)
_cache = {}


def font(bold, size):
    key = (bold, round(size * SCALE))
    if key not in _cache:
        _cache[key] = ImageFont.truetype(FONT_BOLD if bold else FONT_REG, round(size * SCALE))
    return _cache[key]


def char_w(ch, bold, size):
    return font(bold, size).getlength(ch) / SCALE / 72.0  # 인치


def natural_line(size, bold):
    asc, desc = font(bold, size).getmetrics()
    return (asc + desc) / SCALE / 72.0


def para_runs(p, default_size, force_bold=False):
    """문단 → [(글자, pt, bold)]"""
    out = []
    for r in p.runs:
        size = r.font.size.pt if r.font.size else default_size
        bold = bool(r.font.bold) or force_bold
        for ch in r.text.replace('\v', '\n'):
            out.append((ch, size, bold))
    return out


def wrap(chars, width):
    """어절(공백) 단위 줄바꿈. 반환: [(줄 글자 목록)] 와 '강제 분해' 여부."""
    lines, cur, cur_w, forced = [], [], 0.0, False
    tokens, tok = [], []
    for c in chars:
        tok.append(c)
        if c[0] in ' \n':
            tokens.append(tok)
            tok = []
    if tok:
        tokens.append(tok)
    for t in tokens:
        tw = sum(char_w(c[0], c[2], c[1]) for c in t)
        trail = t[-1][0] == ' '
        visible_w = tw - (char_w(' ', t[-1][2], t[-1][1]) if trail else 0)
        if t[-1][0] == '\n':
            cur.extend(t[:-1])
            lines.append(cur)
            cur, cur_w = [], 0.0
            continue
        if cur and cur_w + visible_w > width + 1e-6:
            lines.append(cur)
            cur, cur_w = [], 0.0
        if visible_w > width + 1e-6:  # 한 어절이 줄보다 길다 → 글자 단위로 쪼갬
            forced = True
            for c in t:
                cw = char_w(c[0], c[2], c[1])
                if cur and cur_w + cw > width + 1e-6:
                    lines.append(cur)
                    cur, cur_w = [], 0.0
                cur.append(c)
                cur_w += cw
        else:
            cur.extend(t)
            cur_w += tw
    if cur or not lines:
        lines.append(cur)
    return lines, forced


def shape_text_height(shape, report_forced):
    tf = shape.text_frame
    bodyPr = tf._txBody.find('{http://schemas.openxmlformats.org/drawingml/2006/main}bodyPr')
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    l = Emu(int(bodyPr.get('lIns', 91440))).inches
    r = Emu(int(bodyPr.get('rIns', 91440))).inches
    t = Emu(int(bodyPr.get('tIns', 45720))).inches
    b = Emu(int(bodyPr.get('bIns', 45720))).inches
    width = Emu(shape.width).inches - l - r
    total = 0.0
    any_forced = False
    for p in tf.paragraphs:
        pPr = p._p.find('a:pPr', ns)
        marL = int(pPr.get('marL', 0)) / 914400 if pPr is not None else 0.0
        mult = 1.0
        bef = aft = 0.0
        if pPr is not None:
            sp = pPr.find('a:lnSpc/a:spcPct', ns)
            if sp is not None:
                mult = int(sp.get('val')) / 100000
            sb = pPr.find('a:spcBef/a:spcPts', ns)
            sa = pPr.find('a:spcAft/a:spcPts', ns)
            bef = int(sb.get('val')) / 100 / 72 if sb is not None else 0.0
            aft = int(sa.get('val')) / 100 / 72 if sa is not None else 0.0
        is_title = shape.is_placeholder and shape.placeholder_format.type is not None and 'TITLE' in str(shape.placeholder_format.type)
        chars = para_runs(p, 18, force_bold=is_title)
        if not chars:
            # 빈 문단: 기본 크기 한 줄
            total += natural_line(18, False) * mult + bef + aft
            continue
        lines, forced = wrap(chars, width - marL)
        any_forced = any_forced or forced
        for ln in lines:
            if ln:
                h = max(natural_line(c[1], c[2]) for c in ln)
            else:
                h = natural_line(chars[0][1], chars[0][2])
            total += h * mult
        total += bef + aft
    return total, Emu(shape.height).inches - t - b, any_forced, width


def main():
    path = sys.argv[1]
    slack = 0.0
    if '--slack' in sys.argv:
        slack = float(sys.argv[sys.argv.index('--slack') + 1])
    prs = Presentation(path)
    sw, sh = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    problems = 0
    for idx, slide in enumerate(prs.slides, 1):
        boxes = []
        for shape in slide.shapes:
            x, y, w, h = (Emu(v).inches for v in (shape.left, shape.top, shape.width, shape.height))
            name = shape.name
            if x < -0.01 or y < -0.01 or x + w > sw + 0.01 or y + h > sh + 0.01:
                print(f'[슬라이드 {idx}] 범위 밖: {name} ({x:.2f},{y:.2f},{w:.2f},{h:.2f})')
                problems += 1
            if shape.has_text_frame and shape.text_frame.text.strip() and '<a:fld' not in shape._element.xml:
                need, have, forced, width = shape_text_height(shape, True)
                snippet = shape.text_frame.text.strip().replace('\n', ' / ')[:34]
                ratio = need / have if have > 0 else 9
                flag = ''
                if need > have * (1 + slack) + 1e-6:
                    flag = '넘침'
                elif ratio > 0.92:
                    flag = '빠듯'
                if forced:
                    flag += ' 어절이 칸보다 김'
                if flag:
                    print(f'[슬라이드 {idx}] {flag.strip()}: “{snippet}” 필요 {need:.2f}in / 칸 {have:.2f}in ({ratio * 100:.0f}%) 폭 {width:.2f}in')
                    if '넘침' in flag or '김' in flag:
                        problems += 1
                boxes.append((name, snippet, x, y, w, h, need, have))
        # 글상자끼리 겹침(실제 글자가 차지하는 높이 기준)
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                ax0, ay0, ax1, ay1 = a[2], a[3], a[2] + a[4], a[3] + min(a[5], max(a[6], 0.1))
                bx0, by0, bx1, by1 = b[2], b[3], b[2] + b[4], b[3] + min(b[5], max(b[6], 0.1))
                ox = min(ax1, bx1) - max(ax0, bx0)
                oy = min(ay1, by1) - max(ay0, by0)
                if ox > 0.05 and oy > 0.05:
                    print(f'[슬라이드 {idx}] 글상자 겹침: “{a[1]}” ↔ “{b[1]}” (가로 {ox:.2f}in, 세로 {oy:.2f}in)')
                    problems += 1
    print(f'점검 끝: 문제 {problems}건')


if __name__ == '__main__':
    main()
