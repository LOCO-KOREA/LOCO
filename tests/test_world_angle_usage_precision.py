"""Real 66.88395999428462 -> 66.9 knee references retain usage in every output."""
import copy
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook,load_workbook
from ergo.gpt.payload import build_numeric_summary,_posture_results
from ergo.report.plain import _posture
from ergo.report.renderer import _write_sheet


def recorded_knee():
    # Values recorded in actual-five/run-02, source 5234f9b. The raw flag is
    # availability metadata; selection lives in current_score.measurement_refs.
    angle=dict(side='LEFT',angle_code='KNEE_FLEX_3D',value_deg=66.88395999428462,
               coordinate_space='INFERRED_WORLD_3D',used_for_reba=False,min_confidence=.9847931861877441)
    reference=dict(angle,value_deg=66.9,used_for_reba=True)
    measurement=dict(scope='COMMON',field='knee_band',value='OVER_60',angles=[reference],
                     estimated=True,measurement_method='INFERRED_WORLD_3D')
    slot=dict(slot_no=1,photo_id=18,label='LEFT',analysis_target=True)
    return dict(snapshot=dict(schema_version='2.0',capture=dict(slots=[slot]),settings=dict(pose_model=dict(confidence_min=.5))),
                pose=dict(slots=[dict(slot_no=1,photo_id=18,status='VALID',world_angles=[angle],angles=[],points=[])]),
                reba=dict(slots=[dict(slot_no=1,photo_id=18,current_score=dict(total=8,risk='HIGH',kind='WITH_ASSUMPTIONS',measurement_refs=[measurement]))]))


def sheet(value):
    return _posture(dict(tasks=[dict(name='Recorded 20kg task',analysis=value,source=value['snapshot'])]))


def used_in_ai(value):
    photos=_posture_results(value,value['snapshot']['capture']['slots'])
    return any(angle['used_for_reba'] for photo in photos for angle in photo['inferred_world_angles'])


def used_in_sheet(value):
    result=sheet(value)
    return any('현재 REBA 초안에 사용' in str(cell) for table in (result or {}).get('tables',[]) for row in table['rows'] for cell in row)


class WorldAngleUsagePrecision(unittest.TestCase):
    def test_recorded_knee_selection_reaches_ai_and_physical_excel_without_mutation(self):
        value=recorded_knee();before=copy.deepcopy(value)
        self.assertTrue(used_in_ai(value))
        with tempfile.TemporaryDirectory() as temporary:
            file=Path(temporary)/'3d-usage.xlsx';book=Workbook();book.remove(book.active)
            _write_sheet(book,sheet(value),'Recorded single-photo regression');book.save(file);book.close()
            restored=load_workbook(file,read_only=True)
            rows=list(restored['자세 분석(REBA)'].iter_rows(values_only=True));restored.close()
            row=next(row for row in rows if len(row)>5 and row[3]=='3D 추정 무릎 굽힘')
            self.assertEqual(row[4],66.9)
            self.assertIn('현재 REBA 초안에 사용',row[5])
            self.assertNotIn('사용하지 않음',row[5])
        self.assertEqual(value,before,'output projection must not rewrite angles, selected refs or scores')

    def test_usage_never_crosses_photo_side_code_coordinate_selection_or_target(self):
        cases=[('photo',lambda v:v['reba']['slots'][0].update(photo_id=999)),
               ('side',lambda v:v['pose']['slots'][0]['world_angles'][0].update(side='RIGHT')),
               ('angle',lambda v:v['pose']['slots'][0]['world_angles'][0].update(angle_code='ELBOW_FLEX_3D')),
               ('raw coordinate',lambda v:v['pose']['slots'][0]['world_angles'][0].update(coordinate_space='IMAGE_2D')),
               ('selected 2D',lambda v:v['reba']['slots'][0]['current_score']['measurement_refs'][0]['angles'][0].update(coordinate_space='IMAGE_2D')),
               ('not selected',lambda v:v['reba']['slots'][0]['current_score']['measurement_refs'][0]['angles'][0].update(used_for_reba=False)),
               ('confirmed manual',lambda v:v['reba']['slots'][0]['current_score'].update(kind='CONFIRMED'))]
        for name,change in cases:
            with self.subTest(name=name):
                value=recorded_knee();change(value)
                self.assertFalse(used_in_ai(value));self.assertFalse(used_in_sheet(value))
        value=recorded_knee();value['snapshot']['capture']['slots'][0]['analysis_target']=False
        self.assertEqual(build_numeric_summary(value)['posture_results'],[])
        self.assertFalse(used_in_sheet(value))

    def test_reference_must_equal_the_raw_value_or_its_stored_one_decimal_rounding(self):
        for raw,reference,expected in [(66.88395999428462,66.89,False),(66.84,66.9,False),
                                       (66.84,66.8,True),(66.25,66.2,True),(66.25,66.3,False),
                                       (66.55,66.5,True),(66.55,66.6,False)]:
            with self.subTest(raw=raw,reference=reference):
                value=recorded_knee();value['pose']['slots'][0]['world_angles'][0]['value_deg']=raw
                value['reba']['slots'][0]['current_score']['measurement_refs'][0]['angles'][0]['value_deg']=reference
                self.assertEqual(used_in_ai(value),expected);self.assertEqual(used_in_sheet(value),expected)


if __name__=='__main__':unittest.main()
