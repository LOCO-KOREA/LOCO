"""Scope cut (owner, 2026-10-27: "점수 계산에 안 쓰이면 보여주지마. 받지도 마."): LOCO calculates REBA and hands the values over as an Excel file.

Four record groups never fed a score and are no longer asked, shown or received: the reason of an occasional survey, the worksite situation,
the worker participation and the symptom survey. Rows an older version saved still hold them; those rows must load and the groups must reach no output."""
import inspect
import json
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from openpyxl import load_workbook

from backend.engine import quick_assess
from backend.engine.survey_defaults import normalize_survey
from backend.engine.task_defaults import now, task_input
from backend.intake import contract
from backend.report.form_mapping import map_snapshot
from backend.service.application import create_service

REMOVED = ('occasional_reason', 'occasional_cause', 'occasional_cause_source', 'workplace_situation', 'participation', 'symptom_survey')
KEPT = ('survey_type', 'survey_date', 'investigator', 'department', 'process', 'title', 'previous_survey_date')
# What the old Excel printed from the removed groups, and what the old editor wrote into them.
MARKERS = ('수시조사 사유', '새 포장 설비 도입', '질환자가 발생했거나', '근로자 참여', '참여함', 'W-1', '작업 방법 의견',
           '작업장 상황', '작업설비', '작업속도', '업무형태', '설비 교체함', '물량 증가', '교대 변경',
           '증상조사 요약', '증상조사 기록', '증상조사 방법', '익명 설문 1건', '익명-001')


def base(**changes):
    value = dict(survey_type='REGULAR', survey_date='2026-10-23', investigator='조사자', department='팀', process='공정', title='범위 검증')
    value.update(changes)
    return value


def removed_groups():
    """Every one of the six removed keys filled in, the way the old editor saved them."""
    return dict(
        occasional_reason='새 포장 설비 도입', occasional_cause='DISEASE_CASE', occasional_cause_source='USER',
        workplace_situation=dict(equipment_status='YES', equipment_notes='설비 교체함', workload_status='YES', workload_notes='물량 증가',
                                 speed_status='NO', speed_notes='', work_pattern_status='YES', work_pattern_notes='교대 변경'),
        participation=dict(status='YES', role='AFFECTED_WORKER', participant_code='W-1', method='INTERVIEW', date='2026-10-23', notes='작업 방법 의견'),
        symptom_survey=dict(method='QUESTIONNAIRE', date='2026-10-23', respondent_count=1, summary='익명 설문 1건', responses=[dict(
            response_code='익명-001', name='', work_related_symptoms_last_year='YES',
            regions=dict(neck=dict(side='BOTH', duration='WEEK_TO_MONTH', severity='MODERATE', last_week='YES')))]))


class Normalization(unittest.TestCase):
    def test_only_the_kept_fields_survive_whatever_is_sent(self):
        sent = normalize_survey(base(**removed_groups()))
        self.assertEqual(sent, normalize_survey(base()), 'the removed groups change nothing')
        self.assertFalse(set(REMOVED) & set(sent), 'none of the six keys is part of a survey any more')
        self.assertEqual(sent['survey_type'], 'REGULAR')
        self.assertEqual(sent['title'], '범위 검증')

    def test_values_of_the_removed_groups_are_not_even_checked(self):
        broken = base(participation={'status': 'MAYBE'}, occasional_cause='NOPE',
                      symptom_survey={'responses': [{'response_code': 'a'}, {'response_code': 'a'}], 'respondent_count': -3},
                      workplace_situation={'equipment_status': 'SOMETIMES'})
        self.assertFalse(set(REMOVED) & set(normalize_survey(broken)))

    def test_an_occasional_survey_needs_no_reason(self):
        self.assertEqual(normalize_survey(base(survey_type='OCCASIONAL'))['survey_type'], 'OCCASIONAL')
        self.assertEqual(list(inspect.signature(normalize_survey).parameters), ['value'], 'there is no switch that asks for the reason any more')

    def test_the_kept_fields_are_still_checked_and_kept(self):
        for broken in (base(survey_type='WEEKLY'), base(previous_survey_date='어제'), base(survey_date=20261023), base(title=7)):
            with self.subTest(broken=broken), self.assertRaises(ValueError):
                normalize_survey(broken)
        kept = normalize_survey(base(previous_survey_date='2026-10-01'))
        self.assertEqual({key: kept[key] for key in KEPT if key in kept}['previous_survey_date'], '2026-10-01')
        self.assertTrue(set(KEPT) <= set(kept), 'the fields the Excel header prints are still there')


class Questions(unittest.TestCase):
    def test_the_question_list_has_none_of_the_four_groups(self):
        questions = quick_assess.questions({}, [], {})
        ids = {q['id'] for q in questions}
        self.assertFalse(ids & {'participation', 'participation_method', 'symptom', 'situation_equipment', 'situation_workload', 'situation_speed', 'situation_work_pattern'})
        self.assertFalse({q['target']['type'] for q in questions} & {'situation', 'participation', 'symptom'})
        self.assertFalse({q['group'] for q in questions} & {'작업장 상황', '참여·증상'})
        self.assertEqual([q['id'] for q in questions if q['group'] == '조사 정보'], ['task_name', 'template', 'investigator', 'department', 'process'])
        self.assertEqual(contract.LATER_GROUPS, ('조사 정보',))
        self.assertFalse(hasattr(quick_assess, 'SITUATION'))


class Service(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.db = str(Path(self.temp.name) / 'case.db')
        self.ctx = create_service(self.db)
        self.ctx.dispatch('list_surveys', {})

    def stored(self, survey_id):
        with closing(sqlite3.connect(self.db)) as connection:
            return json.loads(connection.execute('SELECT metadata_json FROM SURVEY WHERE survey_id=?', (survey_id,)).fetchone()[0])

    def old_row(self, **changes):
        """A row an older version saved: written straight to the table, past every check of today's service."""
        metadata = dict(base(title='옛 조사', survey_type='OCCASIONAL', previous_survey_date='2025-10-01'), survey_type_source='USER', **removed_groups())
        metadata.update(changes)
        review = dict(reviewed=False, reviewed_analysis_ids=[], reviewed_metadata_revision=None, reviewer=None, reviewed_at=None, note=None)
        return self.ctx.store('create_survey', metadata, review, now(), now())

    def add_task(self, survey_id, name='상자 포장'):
        return self.ctx.dispatch('create_task', {'survey_id': survey_id, 'draft': task_input({'task_name': name})})

    def export(self, survey_id):
        destination = str(Path(self.temp.name) / f'report-{survey_id}.xlsx')
        self.ctx.dispatch('export_excel', {'survey_id': survey_id, 'spec': {'analysis_ids': [], 'include_representative': False, 'comparisons': [], 'destination': destination}})
        book = load_workbook(destination, data_only=False)
        self.addCleanup(book.close)
        return book

    def test_nothing_of_the_four_groups_is_received_or_stored(self):
        created = self.ctx.dispatch('create_survey', {'draft': base(**removed_groups())})
        self.assertFalse(set(REMOVED) & set(created))
        self.assertFalse(set(REMOVED) & set(self.stored(created['survey_id'])), 'not even saved')
        again = self.ctx.dispatch('update_survey', {'survey_id': created['survey_id'], 'draft': dict(created, **removed_groups(), title='다른 제목')})
        self.assertEqual(again['title'], '다른 제목')
        self.assertFalse(set(REMOVED) & set(again))
        self.assertFalse(set(REMOVED) & set(self.stored(created['survey_id'])))
        self.assertFalse(set(REMOVED) & set(self.ctx.dispatch('get_survey', {'survey_id': created['survey_id']})))

    def test_an_occasional_survey_is_saved_without_a_reason(self):
        created = self.ctx.dispatch('create_survey', {'draft': base(survey_type='OCCASIONAL')})
        self.assertEqual((created['survey_type'], created['survey_type_source']), ('OCCASIONAL', 'USER'))
        changed = self.ctx.dispatch('update_survey', {'survey_id': created['survey_id'], 'draft': base(survey_type='OCCASIONAL', title='이름만 변경')})
        self.assertEqual((changed['title'], changed['survey_type']), ('이름만 변경', 'OCCASIONAL'))

    def test_the_overview_has_no_survey_followups_and_still_names_real_gaps(self):
        survey = self.ctx.dispatch('create_survey', {'draft': base()})
        self.add_task(survey['survey_id'])
        overview = self.ctx.dispatch('preview_excel', {'survey_id': survey['survey_id']})
        self.assertNotIn('survey_followups', overview)
        self.assertEqual(set(overview['state']), {'reviewed', 'review_current', 'has_incomplete', 'has_stale'})
        self.assertEqual([(gap['module'], gap['kind']) for gap in overview['tasks'][0]['next_inputs']], [('ANALYSIS', 'NOT_CALCULATED')],
                         'what the calculation still needs is told as before')
        self.assertTrue(overview['state']['has_incomplete'])

    def test_an_old_row_loads_without_the_removed_groups_and_keeps_the_rest(self):
        identity = self.old_row()
        loaded = self.ctx.dispatch('get_survey', {'survey_id': identity})
        self.assertFalse(set(REMOVED) & set(loaded))
        self.assertEqual((loaded['title'], loaded['survey_type'], loaded['previous_survey_date']), ('옛 조사', 'OCCASIONAL', '2025-10-01'))
        self.assertEqual([item['survey_id'] for item in self.ctx.dispatch('list_surveys', {})], [identity])
        reviewed = self.ctx.dispatch('set_survey_review', {'survey_id': identity, 'review': dict(
            reviewed=False, reviewed_analysis_ids=[], reviewed_metadata_revision=None, reviewer=None, reviewed_at=None, note=None)})
        self.assertFalse(set(REMOVED) & set(reviewed), 'every handler that returns a survey leaves them out')
        saved = self.ctx.dispatch('update_survey', {'survey_id': identity, 'draft': loaded})
        self.assertFalse(set(REMOVED) & set(saved))
        kept = self.stored(identity)
        self.assertEqual({key: kept[key] for key in REMOVED}, {key: removed_groups()[key] for key in REMOVED},
                         'saving the row keeps what the old version stored, untouched (existing rows are preserved)')
        again = self.ctx.dispatch('update_survey', {'survey_id': identity, 'draft': dict(loaded, title='새 제목', participation={'status': 'NO'})})
        self.assertEqual(again['title'], '새 제목')
        self.assertFalse(set(REMOVED) & set(again), 'still never handed out')
        self.assertEqual(self.stored(identity)['participation'], removed_groups()['participation'], 'a new value sent for a removed group is not stored over the old one')

    def test_an_old_row_reaches_neither_the_excel_nor_the_preview_nor_the_form_mapping(self):
        identity = self.old_row()
        self.add_task(identity)
        snapshot = self.ctx.export_snapshot(identity, [])
        self.assertFalse(set(REMOVED) & set(snapshot['survey']), 'every output is built from this snapshot')
        book = self.export(identity)
        self.assertNotIn('증상조사 요약', book.sheetnames)
        self.assertEqual(book.sheetnames, ['조사 결과'])
        text = '\n'.join(str(cell.value) for sheet in book for row in sheet.iter_rows() for cell in row if cell.value not in (None, ''))
        for marker in MARKERS:
            self.assertNotIn(marker, text, marker)
        self.assertIn('이전 조사일', text, 'the kept header fields are still printed')
        preview = json.dumps(self.ctx.dispatch('preview_excel_sheets', {'survey_id': identity}), ensure_ascii=False)
        for marker in MARKERS:
            self.assertNotIn(marker, preview, marker)
        model = map_snapshot(snapshot)
        self.assertTrue(all(row['state'] == 'UNKNOWN' and row['note'] == '' for row in model['situation']))
        self.assertFalse([warning for warning in model['warnings'] if '변경 있음' in warning])

    def test_an_old_row_whose_removed_values_would_not_pass_any_check_still_loads_and_exports(self):
        identity = self.old_row(participation={'status': 'MAYBE'}, occasional_cause='NOPE', workplace_situation={'equipment_status': 'SOMETIMES'},
                                symptom_survey={'responses': [{'response_code': 'a'}, {'response_code': 'a'}], 'respondent_count': -3})
        self.assertFalse(set(REMOVED) & set(self.ctx.dispatch('get_survey', {'survey_id': identity})))
        self.assertEqual(self.export(identity).sheetnames, ['조사 결과'])

    def test_an_old_row_with_nothing_else_to_say_still_exports(self):
        identity = self.old_row(survey_type='REGULAR', previous_survey_date=None)
        book = self.export(identity)
        self.assertEqual(book.sheetnames, ['조사 결과'])


if __name__ == '__main__':
    unittest.main()
