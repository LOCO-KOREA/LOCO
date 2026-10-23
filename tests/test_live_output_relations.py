import unittest
from openpyxl import Workbook
from tests.live_output_relations import lifting_pairs,workbook_lifting_pairs

class RelationalOracle(unittest.TestCase):
    def test_equal_value_sets_cannot_hide_crossed_conditions(self):
        def facts(counts):return [dict(target=target,condition=condition,value=value) for condition,mass,count in zip(('light','heavy'),(5,30),counts) for target,value in [('card.lift.load_mass_kg',mass),('card.lift.lift_count_day',count)]]
        good=facts((99,1));crossed=facts((1,99))
        self.assertEqual(sorted(f['value'] for f in good if f['target'].endswith('lift_count_day')),sorted(f['value'] for f in crossed if f['target'].endswith('lift_count_day')))
        self.assertEqual(lifting_pairs(good),[[30,1],[5,99]])
        self.assertNotEqual(lifting_pairs(crossed),[[30,1],[5,99]])

    def test_workbook_rows_are_joined_by_condition_not_sorted_independently(self):
        book=Workbook();sheet=book.active;sheet.title='입력한 값'
        for index,(label,value) in enumerate([('물체 중량','5 kg'),('하루 들기 횟수','99 회/일'),('물체 중량','30 kg'),('하루 들기 횟수','1 회/일')],1):
            sheet.cell(index,3,f'25kg 이상 들기: {label} (조건 {(index-1)//2+1})');sheet.cell(index,4,value)
        self.assertEqual(workbook_lifting_pairs(book),[[30,1],[5,99]])
        sheet.cell(2,4,'1 회/일');sheet.cell(4,4,'99 회/일')
        self.assertNotEqual(workbook_lifting_pairs(book),[[30,1],[5,99]])
        book.close()

if __name__=='__main__':unittest.main()
