"""Counterexamples for the independent E2E oracle itself, not application behavior."""
import unittest
from openpyxl import Workbook
from tests.live_output_audit import schema_errors,audit_preview,json_equal,submitted_message

class OutputOracle(unittest.TestCase):
    def test_intake_boundary_trims_outer_whitespace_without_hiding_lost_content(self):
        messages=[dict(message_id=1,text='12.5kg 상자를 하루 40회 옮겨요.\n총 7.5시간 일해요.')]
        self.assertIs(submitted_message(messages,' \n12.5kg 상자를 하루 40회 옮겨요.\n총 7.5시간 일해요. \t'),messages[0])
        self.assertIsNone(submitted_message(messages,'12.5kg 상자를 하루 40회 옮겨요. 총 7.5시간 일해요.'))
        self.assertIsNone(submitted_message(messages,'12.5kg 상자를 하루 41회 옮겨요.\n총 7.5시간 일해요.'))

    def test_json_numeric_spelling_is_equal_but_booleans_and_null_are_not_numbers(self):
        self.assertTrue(json_equal(dict(minutes=[120.0,0.0,12.5]),dict(minutes=[120,0,12.5])))
        self.assertFalse(json_equal(dict(value=False),dict(value=0)))
        self.assertFalse(json_equal(dict(value=None),dict(value=0)))
        self.assertFalse(json_equal(dict(value='120'),dict(value=120)))
    def test_zero_null_boolean_and_required_are_distinct(self):
        schema=dict(type='object',required=['value'],additionalProperties=False,
            properties=dict(value=dict(anyOf=[dict(type='number',minimum=0,maximum=100),dict(type='null')])))
        for value in (0,None,12.5):self.assertFalse(schema_errors(dict(value=value),schema))
        for value in (False,True,-1,101,'0'):self.assertTrue(schema_errors(dict(value=value),schema))
        self.assertTrue(schema_errors({},schema));self.assertTrue(schema_errors(dict(value=1,extra=2),schema))

    def test_preview_oracle_detects_a_wrong_exact_cell_and_wrong_unit(self):
        book=Workbook();sheet=book.active;sheet.title='입력한 값'
        sheet.cell(6,1,'하루 실제 작업시간');sheet.cell(6,2,'420 분/일')
        preview=[dict(name='입력한 값',lead='입력',tables=[dict(columns=['항목','값'],rows=[['하루 실제 작업시간','420 분/일']])])]
        checks=[];defects=[];audit_preview('test',preview,book,checks,defects);self.assertFalse(defects)
        sheet.cell(6,2,'420 시간');checks=[];defects=[];audit_preview('test',preview,book,checks,defects)
        self.assertEqual(len(defects),1);self.assertEqual(defects[0]['cell'],'B6')

    def test_strict_calculation_operands_cannot_be_empty_or_partial(self):
        schema=dict(type='array',minItems=2,maxItems=3,items=dict(type='object',additionalProperties=False,
            required=['quote','minutes'],properties=dict(quote=dict(type='string',minLength=1),minutes=dict(type='number',minimum=0))))
        valid=[dict(quote='2시간',minutes=120),dict(quote='6시간',minutes=360)]
        self.assertFalse(schema_errors(valid,schema));self.assertTrue(schema_errors(valid[:1],schema))
        self.assertTrue(schema_errors([dict(quote='',minutes=0),valid[1]],schema))

if __name__=='__main__':unittest.main()
