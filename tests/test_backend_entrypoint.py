"""Exercise the renamed package in fresh processes without the old package path."""
import json
import os
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from contextlib import closing
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]


class BackendEntrypoint(unittest.TestCase):
    def test_file_and_module_entrypoints_store_and_export_without_ergo(self):
        for isolated in (True, False):
            with self.subTest(isolated_file=isolated), tempfile.TemporaryDirectory() as directory:
                folder = Path(directory)
                database = folder / 'worker.db'
                output = folder / 'report.xlsx'
                env = {key: value for key, value in os.environ.items()
                       if not key.startswith(('LOCO_', 'OPENAI_', 'PYTHON'))}
                env.update(LOCO_DB_PATH=str(database), PYTHONIOENCODING='utf-8',
                           PYTHONDONTWRITEBYTECODE='1')
                args = (['-I', str(ROOT / 'backend/service/json_lines.py')] if isolated
                        else ['-B', '-m', 'backend.service.json_lines'])
                requests = [
                    ('create_survey', {'draft': {'title': 'rename smoke', 'survey_type': 'REGULAR'}}),
                    ('create_task', {'survey_id': 1, 'draft': {}}),
                    ('export_excel', {'survey_id': 1, 'spec': {
                        'analysis_ids': [], 'include_representative': False,
                        'destination': str(output)}}),
                    ('nonexistent_method', {}),
                    ('list_surveys', {}),
                ]
                lines = '\n'.join(json.dumps(dict(request_id=i, method=method, params=params))
                                  for i, (method, params) in enumerate(requests, 1)) + '\n'
                completed = subprocess.run([sys.executable, *args], input=lines,
                    cwd=folder if isolated else ROOT, env=env, text=True,
                    encoding='utf-8', capture_output=True, timeout=30)
                self.assertEqual(completed.returncode, 0, completed.stderr)
                responses = {row['request_id']: row['result']
                             for row in map(json.loads, completed.stdout.splitlines())}
                self.assertEqual(set(responses), set(range(1, 6)))
                for identity in (1, 2, 3, 5):
                    self.assertTrue(responses[identity]['ok'], responses[identity])
                self.assertFalse(responses[4]['ok'])
                workbook = load_workbook(output, read_only=True)
                try:
                    self.assertTrue(workbook.sheetnames)
                finally:
                    workbook.close()
                with closing(sqlite3.connect(database)) as connection:
                    tables = {row[0] for row in connection.execute(
                        "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")}
                self.assertEqual(tables, {'SURVEY', 'TASK', 'PHOTO', 'ANALYSIS', 'APP_SETTINGS', 'EXPORT_LOG'})

    @unittest.skipUnless(sys.platform == 'win32' and sys.version_info[:2] == (3, 12),
                         'Windows CPython 3.12 build contract')
    def test_freeze_command_includes_renamed_providers_and_resources(self):
        from scripts import build_windows

        if not all((ROOT / definition['model_path']).is_file()
                   for definition in build_windows.MODELS.values()):
            self.skipTest('Registered pose models are not installed')

        class FreezeCommandCaptured(Exception):
            pass

        with tempfile.TemporaryDirectory() as directory:
            with patch.dict(os.environ, {'LOCO_BUILD_WORK_ROOT': directory}), patch(
                    'scripts.build_windows.subprocess.run', side_effect=FreezeCommandCaptured) as freeze:
                with self.assertRaises(FreezeCommandCaptured):
                    build_windows.main()
            command = freeze.call_args.args[0]
            self.assertEqual(Path(command[-1]), ROOT / 'backend/service/json_lines.py')
            for provider in ('mediapipe', 'movenet'):
                self.assertIn('backend.pose.providers.' + provider, command)
            for resource in ('db/schema_v1.sql', 'report/templates', 'report/assets'):
                source = ROOT / 'backend' / resource
                self.assertTrue(source.exists(), source)
                destination = 'backend/db' if source.is_file() else 'backend/' + resource
                self.assertIn(str(source) + ';' + destination, command)
