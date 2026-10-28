"""The workbook for ordinary readers: everyday words only, no detailed record for experts, and a preview that matches the file."""
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from backend.engine.quick_assess import NO_BURDEN
from backend.engine.task_defaults import ARM, COMMON, task_input
from backend.report import workbook
from backend.report.plain import _inputs, _legal_rows, _posture, _explanations
from backend.service import application
from tests.report_helpers import CODE, SURVEY_KEYS, lifting_input, make_survey, run_analysis, update_survey

PLAIN_SHEETS = ['조사 결과', '부담작업 해당 여부', '입력한 값']


def blocks(sheet):
    """{heading: the rows under it} for the 부담작업 해당 여부 sheet. A block runs until the next empty row."""
    result, rows = {}, list(sheet.iter_rows(values_only=True))
    for index, row in enumerate(rows):
        if row[0] and '→ 최종 판단:' in str(row[0]):
            body = []
            for later in rows[index + 2:]:
                if all(value in (None, '') for value in later):
                    break
                body.append(later)
            result[row[0]] = body
    return result


class PlainWorkbook(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        patcher = patch.object(application, 'DB', str(Path(self.temp.name) / 'loco.db'))
        patcher.start()
        self.addCleanup(patcher.stop)
        self.addCleanup(self.temp.cleanup)
        self.survey = make_survey()
        self.task = application.dispatch('create_task', {'survey_id': self.survey['survey_id'], 'draft': lifting_input()})
        self.analysis = run_analysis(self.task)

    def export(self, comparisons=None):
        destination = str(Path(self.temp.name) / 'report.xlsx')
        application.dispatch('export_excel', {'survey_id': self.survey['survey_id'], 'spec': {
            'analysis_ids': [], 'include_representative': False, 'comparisons': comparisons or [], 'destination': destination}})
        return load_workbook(destination, data_only=False)

    def cells(self, book, name):
        return [cell.value for row in book[name].iter_rows() for cell in row if cell.value not in (None, '')]

    def test_the_workbook_holds_only_the_plain_sheets(self):
        book = self.export()
        self.assertEqual(book.sheetnames, PLAIN_SHEETS)
        self.assertNotIn('전문가용 상세기록', book.sheetnames, 'the detailed record for experts was removed on request')

    def test_plain_sheets_contain_no_codes_or_internal_words(self):
        book = self.export()
        for name in PLAIN_SHEETS:
            for value in self.cells(book, name):
                if isinstance(value, str):
                    self.assertIsNone(CODE.search(value), f'{name}: {value!r}')

    def test_no_ids_versions_or_schema_on_the_first_sheet(self):
        text = ' '.join(str(v) for v in self.cells(self.export(), '조사 결과'))
        for word in ('조사 ID', '분석 ID', '작업 ID', '스키마', '보고서 버전', 'T0', '.000000'):
            self.assertNotIn(word, text)


    def test_burden_sheet_contains_only_recorded_criteria_and_calculated_results(self):
        book = self.export()
        sheet = book['부담작업 해당 여부']
        rows = list(sheet.iter_rows(values_only=True))
        header = next(index for index, row in enumerate(rows) if row[:3] == ('작업', '기준', '적용 결과'))
        rows = rows[header:]
        text = '\n'.join(str(value) for row in rows for value in row if value not in (None, ''))
        self.assertIn('키보드·마우스 집중 입력', text)
        self.assertIn('부담작업', text)
        self.assertIn('25kg 이상 들기', text)
        self.assertIn('부담작업 아님', text)
        for marker in ('입력 필요', '입력 안 됨', '아직 답하지 않음', '미확정'):
            self.assertNotIn(marker, text)

    def test_explicit_unknown_answer_is_kept_in_inputs_but_default_unknown_is_omitted(self):
        application.dispatch('delete_task', {'task_id': self.task['task_id']})
        task = application.dispatch('create_task', {'survey_id': self.survey['survey_id'],
            'draft': task_input({'task_name': '일회성 작업 여부 확인'})})
        application.dispatch('intake_session', {'task_id': task['task_id'], 'views': {}})
        saved = application.dispatch('intake_turn', {'task_id': task['task_id'],
            'answers': [{'target': 'one_off', 'value': None, 'condition': 'default'}]})
        fact = next(f for f in saved['facts'] if f['target'] == 'one_off')
        self.assertEqual((fact['state'], fact['source']), ('UNKNOWN', 'USER'))
        book = self.export()
        self.assertIn('입력한 값', book.sheetnames)
        values = self.cells(book, '입력한 값')
        self.assertTrue(any('모름' in str(value) for value in values))

    def test_inputs_sheet_lists_what_was_entered_with_units_and_how(self):
        rows = list(self.export()['입력한 값'].iter_rows(values_only=True, min_row=6))
        values = {row[2]: row[3] for row in rows}
        self.assertEqual(len(rows[0]), 4, 'provenance stays with its value; no sparse source/note columns')
        self.assertIn('240 분/일 · 면담', values['키보드·마우스 집중 입력: 집중 키보드·마우스 사용 시간'])
        self.assertIn('20 kg · 면담', values['25kg 이상 들기: 물체 중량'])
        self.assertEqual(values['하루 실제 작업시간'], '480 분/일')

    def test_empty_input_provenance_columns_are_omitted(self):
        sheet = _inputs({'tasks': [{'name': '기록 작업', 'source': {
            'conditions': {'daily_work_min': 480}, 'legal': {}, 'capture': {'slots': []}, 'manual': []}}]})
        table = sheet['tables'][0]
        self.assertEqual(table['columns'], ['작업', '구분', '항목', '값'])
        self.assertEqual(table['rows'], [['기록 작업', '작업 조건', '하루 실제 작업시간', '480 분/일']])
        self.assertEqual(sheet['lead'], '계산에 쓴 값이에요.')

    def test_only_exact_user_or_ai_unknown_answers_survive_excel_projection(self):
        from backend.intake.contract import SESSION_KEY

        for origin in ('USER', 'AI'):
            with self.subTest(origin=origin):
                source = {
                    'conditions': {'daily_work_min': 'UNKNOWN', SESSION_KEY: {'facts': {
                        'work_type.keyboard@default': {
                            'target': 'work_type.keyboard', 'condition': 'default', 'value': None,
                            'source': origin, 'quote': '작업 종류를 모르겠어요'}
                    }}},
                    'legal': {'criteria': [
                        {'criterion_no': 1, 'condition_present': 'UNKNOWN', 'evidence': None},
                        {'criterion_no': 2, 'condition_present': 'UNKNOWN', 'evidence': None}], 'exposures': []},
                    'capture': {'slots': []},
                    'manual': [{'slot_no': 1, 'scope': 'COMMON', 'field': 'trunk_base', 'value': 'UNKNOWN',
                                'source': 'POSE', 'evidence': None}],
                }
                sheet = _inputs({'tasks': [{'name': '기록 작업', 'source': source}]})
                table = sheet['tables'][0]
                self.assertEqual(table['columns'], ['작업', '구분', '항목', '값'])
                labels = [row[2] for row in table['rows']]
                self.assertIn('키보드·마우스 집중 입력: 이 작업을 한다', labels)
                self.assertNotIn('같은 동작 반복: 이 작업을 한다', labels,
                                 'a related parent answer cannot fill a missing child UNKNOWN')
                self.assertNotIn('하루 실제 작업시간', labels)
                self.assertFalse(any('몸통' in label for label in labels), 'untraced posture UNKNOWN is internal')
                value = table['rows'][labels.index('키보드·마우스 집중 입력: 이 작업을 한다')][3]
                self.assertIn('모름', value)
                self.assertIn('작업 종류를 모르겠어요', value)
                self.assertIn('직접 입력' if origin == 'USER' else '작업 설명', value)

    def test_internal_row_completeness_marker_is_not_exported_as_a_user_answer(self):
        source = {'legal': {'criteria': [{'criterion_no': 8, 'condition_present': 'UNKNOWN',
                                          'rows_complete': 'YES', 'rows': [], 'evidence': None}],
                            'exposures': []}}
        rows = _legal_rows('미확인 작업', source)
        self.assertEqual(rows, [], 'rows_complete is a calculation gate, not an investigator answer')

    def test_legacy_unchecked_work_type_is_not_exported_as_a_confirmed_no(self):
        source = {'legal': {'criteria': [{'criterion_no': 2, 'condition_present': 'NO',
                                          'evidence': {'source': None, 'text': '빠른 분석: "이 작업의 특징"에서 선택하지 않음(해당 없음)'}}],
                            'exposures': []}}
        self.assertEqual(_legal_rows('과거 간편 분석', source), [])

    def test_work_type_no_without_provenance_is_omitted_but_quickflow_answer_is_kept(self):
        source = {'legal': {'criteria': [{'criterion_no': 2, 'condition_present': 'NO', 'evidence': None}],
                            'exposures': []}}
        self.assertEqual(_legal_rows('미확인 작업', source), [])

        source['legal']['criteria'][0]['evidence'] = {
            'source': 'INTERVIEW', 'text': '빠른 분석: 조사자가 이 기준을 직접 확인해 아니오로 선택'}
        rows = _legal_rows('직접 확인한 작업', source)
        row = next(row for row in rows if row[2].endswith('이 작업을 한다'))
        self.assertEqual(row[3], '아니오')
        self.assertIn('조사자가 이 기준을 직접 확인해 아니오로 선택', row[5])

    def test_explicit_user_and_ai_work_type_facts_remain_in_excel(self):
        from backend.intake.contract import SESSION_KEY

        for origin in ('USER', 'AI'):
            with self.subTest(origin=origin):
                source = {
                    'conditions': {SESSION_KEY: {'facts': {
                        'work_type.keyboard@default': {
                            'target': 'work_type.keyboard', 'condition': 'default', 'value': 'NO',
                            'source': origin, 'quote': '키보드 작업을 하지 않는다고 확인했습니다'}
                    }}},
                    'legal': {'criteria': [{'criterion_no': 1, 'condition_present': 'NO', 'evidence': None}],
                              'exposures': []}
                }
                rows = _legal_rows('질문한 작업', source)
                row = next(row for row in rows if row[2].endswith('이 작업을 한다'))
                self.assertEqual(row[3], '아니오')
                self.assertIn('직접 입력' if origin == 'USER' else '작업 설명', row[4])
                self.assertIn('키보드 작업을 하지 않는다고 확인했습니다', row[5])

    def test_time_is_written_as_korean_wall_time_not_iso_text(self):
        import re
        book = self.export()
        text = ' '.join(str(v) for v in self.cells(book, '조사 결과'))
        self.assertRegex(text, r'20\d\d-\d\d-\d\d \d\d:\d\d')
        self.assertNotIn('T', ' '.join(re.findall(r'\d{4}-\d\d-\d\dT\S+', text)))


    def test_workbook_has_no_followup_sheet_and_does_not_take_worksite_answers(self):
        self.assertNotIn('확인사항', self.export().sheetnames)
        sent = {key: value for key, value in application.dispatch('get_survey', {'survey_id': self.survey['survey_id']}).items() if key in SURVEY_KEYS}
        application.dispatch('update_survey', {'survey_id': self.survey['survey_id'], 'draft': dict(sent, workplace_situation=dict(
            equipment_status='YES', workload_status='YES', speed_status='YES', work_pattern_status='YES',
            equipment_notes='', workload_notes='', speed_notes='', work_pattern_notes=''))})
        book = self.export()
        main = '\n'.join(str(value) for value in self.cells(book, '조사 결과'))
        for label in ('작업설비', '작업량', '작업속도', '업무형태'):
            self.assertNotIn(label, main)
        self.assertNotIn('입력 필요', main)
        self.assertNotIn('MoveNet', main)
        self.assertNotIn('기술태그', main)


    def test_blank_defaults_do_not_create_excel_rows_or_missing_markers(self):
        application.dispatch('delete_task', {'task_id': self.task['task_id']})
        task = application.dispatch('create_task', {'survey_id': self.survey['survey_id'],
            'draft': task_input({'task_name': '빈 입력 작업'})})
        run_analysis(task)
        book = self.export()
        self.assertEqual(book.sheetnames, ['조사 결과'])
        text = '\n'.join(str(value) for sheet in book for value in self.cells(book, sheet.title))
        for marker in ('입력 필요', '입력 안 됨', '입력되지 않음', '아직 답하지 않음', '미확정'):
            self.assertNotIn(marker, text)
        self.assertNotIn('확인사항', book.sheetnames)
        self.assertIn('빈 입력 작업', text)

    def test_tag_only_posture_data_does_not_create_an_empty_reba_table(self):
        analysis = {'snapshot': {'settings': {'pose_model': {'confidence_min': 0.5}}},
                    'reba': {}, 'pose': {'slots': []},
                    'tags': [{'tag': 'TRUNK_FLEX', 'yes_count': 1, 'observed_count': 1}]}
        sheet = _posture({'tasks': [{'analysis': analysis, 'source': {'capture': {'slots': []}},
                                    'name': '태그 작업'}]})
        self.assertIsNotNone(sheet)
        headings = [table.get('heading') for table in sheet['tables']]
        self.assertIn('자세 표시 (사진 몇 장에서 보였는지, 하루 시간이 아니에요)', headings)
        self.assertNotIn('사진별 REBA 점수 (점수가 높을수록 몸에 부담이 커요)', headings)

    def test_failed_pose_hides_machine_paths_and_does_not_turn_unobserved_tags_into_zero(self):
        model_path = r'C:\private\models\pose_landmarker_heavy.task'
        analysis = {
            'snapshot': {'settings': {'pose_model': {'confidence_min': 0.5}}},
            'reba': {},
            'pose': {'status': 'FAILED', 'slots': [{
                'slot_no': 1, 'status': 'ERROR', 'points': [], 'angles': [],
                'diagnostics': [{'code': 'MODEL_ERROR', 'message': f'model not found: {model_path}'}],
            }]},
            # This malformed/legacy summary says an angle tag was observed despite failed pose.
            'slot_tags': [],
            'tags': [{'tag': 'TRUNK_FLEX', 'yes_count': 0, 'observed_count': 1, 'manual_observed_count': 0}],
        }
        sheet = _posture({'tasks': [{'analysis': analysis, 'source': {'capture': {'slots': []}}, 'name': '실패 작업'}]})
        values = [value for table in sheet['tables'] for row in table['rows'] for value in row]
        text = ' '.join(str(value) for value in values)
        self.assertNotIn(model_path, text)
        self.assertIn('자세 분석을 끝내지 못했어요.', text)
        self.assertNotIn('0%', text)
        self.assertFalse(any(table.get('heading', '').startswith('자세 표시') for table in sheet['tables']))

    def test_successful_zero_percent_tag_keeps_its_observed_denominator_and_source(self):
        analysis = {
            'snapshot': {'settings': {'pose_model': {'confidence_min': 0.5}}},
            'reba': {},
            'pose': {'status': 'DONE', 'valid_count': 2, 'slots': []},
            'tags': [{'tag': 'TRUNK_FLEX', 'yes_count': 0, 'observed_count': 2, 'manual_observed_count': 0}],
        }
        sheet = _posture({'tasks': [{'analysis': analysis, 'source': {'capture': {'slots': []}}, 'name': '정상 자세 작업'}]})
        values = [value for table in sheet['tables'] for row in table['rows'] for value in row]
        self.assertIn('0%', values)
        self.assertIn(2, values)
        self.assertIn('사진 각도에서 확인', values)

    def test_failed_pose_keeps_explicit_manual_zero_tag_with_its_provenance(self):
        analysis = {
            'snapshot': {'settings': {'pose_model': {'confidence_min': 0.5}}},
            'reba': {},
            'pose': {'status': 'FAILED', 'slots': []},
            'slot_tags': [{'slot_no': 1, 'tag': 'TRUNK_FLEX', 'value': 'NO', 'source': 'MANUAL'}],
            'tags': [{'tag': 'TRUNK_FLEX', 'yes_count': 0, 'observed_count': 1, 'manual_observed_count': 1}],
        }
        sheet = _posture({'tasks': [{'analysis': analysis, 'source': {'capture': {'slots': []}}, 'name': '직접 확인 작업'}]})
        values = [value for table in sheet['tables'] for row in table['rows'] for value in row]
        self.assertIn('0%', values)
        self.assertIn('조사자가 직접 확인', values)

    def test_explanation_without_confirmation_time_is_not_exported(self):
        task = {'task_id': 7, 'input_revision': 3, 'latest_analysis_id': 12}
        missing = _explanations({'tasks': [{'name': '설명 작업', 'task': task, 'analysis': {
            'analysis_id': 12, 'input_revision': 3, 'ai': [
                {'is_active': True, 'applied_text': '미확정 시각 설명', 'applied_kind': 'INVESTIGATOR', 'applied_at': None}]}}]})
        self.assertIsNone(missing, 'active state without the apply receipt is malformed and is not a confirmed explanation')
        confirmed = _explanations({'tasks': [{'name': '설명 작업', 'task': task, 'analysis': {
            'analysis_id': 12, 'input_revision': 3, 'ai': [
                {'is_active': True, 'source_current': True, 'applied_text': '확정한 문장', 'applied_kind': 'INVESTIGATOR', 'applied_at': '2026-10-24T00:00:00+00:00'}]}}]})
        values = [value for table in confirmed['tables'] for row in table['rows'] for value in row]
        self.assertIn('확정한 문장', values)
        self.assertNotIn('미확정 시각 설명', values)

    def test_stale_applied_explanation_is_omitted_from_preview_and_excel(self):
        settings = application.dispatch('get_settings', {})
        settings.update(gpt_enabled=True, gpt_model='test-model')
        application.dispatch('update_settings', {'settings': settings})
        with patch('backend.gpt.client.generate', return_value=(None, 'NO_API_KEY')):
            generated = application.dispatch('generate_ai_text', {
                'analysis_id': self.analysis['analysis_id'], 'purpose': 'TASK_EXPLANATION'})
        application.dispatch('edit_ai_text', {'analysis_id': self.analysis['analysis_id'],
            'text_id': generated['text']['text_id'], 'text': '현재 작업 입력에 확정한 설명입니다.'})
        application.dispatch('apply_ai_text', {'analysis_id': self.analysis['analysis_id'],
            'text_id': generated['text']['text_id'], 'use_edited': True})

        current_preview = application.dispatch('preview_excel_sheets', {'survey_id': self.survey['survey_id']})
        self.assertIn('결과 설명', [sheet['name'] for sheet in current_preview['sheets']])
        task = application.dispatch('get_task', {'task_id': self.task['task_id']})['task']
        updated = application.task_input(task)
        updated['task_name'] = '설명 확정 뒤 바뀐 작업'
        saved = application.dispatch('save_task_input', {'task_id': self.task['task_id'], 'input': updated})
        self.assertGreater(saved['input_revision'], self.analysis['input_revision'])

        stale_preview = application.dispatch('preview_excel_sheets', {'survey_id': self.survey['survey_id']})
        stale_names = [sheet['name'] for sheet in stale_preview['sheets']]
        self.assertNotIn('결과 설명', stale_names)
        self.assertIn('부담작업 해당 여부', stale_names)
        book = self.export()
        self.assertNotIn('결과 설명', book.sheetnames)
        self.assertIn('부담작업 해당 여부', book.sheetnames)

    def test_human_review_sheet_requires_current_review_and_is_shared_by_preview_and_excel(self):
        survey = application.dispatch('get_survey', {'survey_id': self.survey['survey_id']})
        review = {
            'reviewed': True,
            'reviewed_analysis_ids': [self.analysis['analysis_id']],
            'reviewed_metadata_revision': survey['metadata_revision'],
            'reviewer': '현장 검토자',
            'reviewed_at': '2026-10-24T01:30:00Z',
            'note': '입력 내용과 분석 결과를 확인함',
        }
        application.dispatch('set_survey_review', {'survey_id': self.survey['survey_id'], 'review': review})

        preview = application.dispatch('preview_excel_sheets', {'survey_id': self.survey['survey_id']})
        people = next(sheet for sheet in preview['sheets'] if sheet['name'] == '사람 검토')
        self.assertEqual(people['tables'][0]['columns'], ['작업명', '검토 상태', '검토자', '검토 시각', '메모'])
        self.assertEqual(people['tables'][0]['rows'], [[
            '상자 들기', '검토 완료', '현장 검토자', '2026-10-24 10:30', '입력 내용과 분석 결과를 확인함']])
        book = self.export()
        self.assertIn('사람 검토', book.sheetnames)
        cells = self.cells(book, '사람 검토')
        for value in ('상자 들기', '검토 완료', '현장 검토자', '2026-10-24 10:30', '입력 내용과 분석 결과를 확인함'):
            self.assertIn(value, cells)

        task = application.dispatch('get_task', {'task_id': self.task['task_id']})['task']
        updated = application.task_input(task)
        updated['task_name'] = '검토 후 변경된 작업'
        saved = application.dispatch('save_task_input', {'task_id': self.task['task_id'], 'input': updated})
        self.assertGreater(saved['input_revision'], self.analysis['input_revision'])
        stale_preview = application.dispatch('preview_excel_sheets', {'survey_id': self.survey['survey_id']})
        self.assertNotIn('사람 검토', [sheet['name'] for sheet in stale_preview['sheets']])
        self.assertNotIn('사람 검토', self.export().sheetnames)

    def test_human_review_excludes_unseen_second_task_from_preview_and_excel(self):
        second = application.dispatch('create_task', {'survey_id': self.survey['survey_id'],
            'draft': lifting_input('표시되지 않은 작업 B')})
        second_analysis = run_analysis(second)
        survey = application.dispatch('get_survey', {'survey_id': self.survey['survey_id']})
        review = {
            'reviewed': True,
            'reviewed_analysis_ids': [self.analysis['analysis_id']],
            'reviewed_metadata_revision': survey['metadata_revision'],
            'reviewer': '작업 A 검토자',
            'reviewed_at': '2026-10-24T01:30:00Z',
            'note': '표시된 작업 A만 확인함',
        }
        application.dispatch('set_survey_review', {'survey_id': self.survey['survey_id'], 'review': review})

        overview = application.dispatch('preview_excel', {'survey_id': self.survey['survey_id']})
        statuses = {item['task_name']: item['review_status'] for item in overview['tasks']}
        self.assertEqual(statuses['상자 들기'], 'REVIEWED_CURRENT')
        self.assertEqual(statuses['표시되지 않은 작업 B'], 'UNREVIEWED')
        preview = application.dispatch('preview_excel_sheets', {'survey_id': self.survey['survey_id']})
        people = next(sheet for sheet in preview['sheets'] if sheet['name'] == '사람 검토')
        self.assertEqual(people['tables'][0]['rows'], [[
            '상자 들기', '검토 완료', '작업 A 검토자', '2026-10-24 10:30', '표시된 작업 A만 확인함']])
        self.assertNotIn(second_analysis['analysis_id'], review['reviewed_analysis_ids'])
        self.assertNotIn('표시되지 않은 작업 B', str(people['tables'][0]['rows']))

        book = self.export()
        self.assertIn('사람 검토', book.sheetnames)
        values = self.cells(book, '사람 검토')
        self.assertIn('상자 들기', values)
        self.assertNotIn('표시되지 않은 작업 B', values)

    def test_preview_changes_nothing(self):
        before = application.dispatch('list_export_logs', {'survey_id': self.survey['survey_id']})
        application.dispatch('preview_excel_sheets', {'survey_id': self.survey['survey_id']})
        self.assertEqual(application.dispatch('list_export_logs', {'survey_id': self.survey['survey_id']}), before)

    def test_applied_explanation_gets_its_own_sheet_and_an_unapplied_one_does_not(self):
        self.assertNotIn('결과 설명', self.export().sheetnames)
        settings = application.dispatch('get_settings', {})
        settings.pop('gpt_key_present', None)
        settings.update(gpt_enabled=True, gpt_model='test-model')
        application.dispatch('update_settings', {'settings': settings})
        with patch('backend.gpt.client.generate', return_value=(None, 'NO_API_KEY')):
            generated = application.dispatch('generate_ai_text', {'analysis_id': self.analysis['analysis_id'], 'purpose': 'TASK_EXPLANATION'})
        self.assertNotIn('결과 설명', self.export().sheetnames, 'a generated draft must not appear in Excel')
        application.dispatch('edit_ai_text', {'analysis_id': self.analysis['analysis_id'], 'text_id': generated['text']['text_id'], 'text': '조사자가 확인하고 확정한 문장입니다.'})
        application.dispatch('apply_ai_text', {'analysis_id': self.analysis['analysis_id'], 'text_id': generated['text']['text_id'], 'use_edited': True})
        book = self.export()
        self.assertEqual(book.sheetnames, PLAIN_SHEETS + ['결과 설명'])
        cells=self.cells(book, '결과 설명')
        self.assertIn('조사자가 확인하고 확정한 문장입니다.', cells)
        self.assertIn('확정 시각', cells)
        self.assertIn('조사자가 고친 설명', cells)
        self.assertNotIn('기본 설명(AI를 쓰지 못했어요)', cells)

    def test_same_revision_applied_explanation_with_stale_numeric_summary_is_omitted(self):
        task = {'task_id': 7, 'input_revision': 3, 'latest_analysis_id': 12}
        sheet = _explanations({'tasks': [{'name': '설명 작업', 'task': task, 'analysis': {
            'analysis_id': 12, 'input_revision': 3, 'ai': [
                {'is_active': True, 'source_current': False, 'input_summary': {'criteria': []},
                 'applied_text': '같은 입력 revision이지만 요약이 낡은 설명', 'applied_kind': 'INVESTIGATOR',
                 'applied_at': '2026-10-24T00:00:00+00:00'}]}}]})
        self.assertIsNone(sheet, 'revision equality alone does not make an applied summary current')

    def test_before_after_sheet_is_in_words_and_hides_bookkeeping_changes(self):
        task = application.dispatch('get_task', {'task_id': self.task['task_id']})['task']
        updated = application.task_input(task)
        updated['legal']['exposures'][0]['minutes'] = 60
        updated['legal']['exposures'][0]['evidence'] = {'source': 'INTERVIEW', 'text': '다시 확인', 'observed_at': '2026-10-24T01:00:00Z'}
        saved = application.dispatch('save_task_input', {'task_id': task['task_id'], 'input': updated})
        after = run_analysis(saved)
        book = self.export([{'before_analysis_id': self.analysis['analysis_id'], 'after_analysis_id': after['analysis_id'], 'note': '시간 줄임'}])
        self.assertEqual(book.sheetnames, PLAIN_SHEETS + ['개선 전후 비교'])
        text = '\n'.join(str(v) for v in self.cells(book, '개선 전후 비교'))
        self.assertIn('집중 키보드·마우스 사용 시간', text)
        self.assertIn('240 분/일', text)
        self.assertIn('60 분/일', text)
        self.assertIn('시간 줄임', text)
        self.assertIsNone(CODE.search(text.replace('MOEL', '')), text)
        self.assertNotIn('observed_at', text)
        self.assertIn('계산에 영향이 없는 변경', text, 'the sheet says how many bookkeeping-only changes it left out')


@unittest.skipUnless(os.path.isfile(Path(__file__).resolve().parents[1] / '.local' / 'fixtures' / 'tensorflow-movenet-example.jpeg')
                     and os.path.isfile(Path(__file__).resolve().parents[1] / 'models' / 'movenet_lightning.tflite'), 'public example photo and official model required')
class PhotoWorkbook(unittest.TestCase):
    def test_photo_workbook_names_the_body_parts_places_the_photo_and_summarizes_the_reba_score(self):
        fixture = Path(__file__).resolve().parents[1] / '.local' / 'fixtures' / 'tensorflow-movenet-example.jpeg'
        with tempfile.TemporaryDirectory() as root, patch.object(application, 'DB', str(Path(root) / 'photo.db')):
            survey = make_survey('사진 보고서')
            task = application.dispatch('create_task', {'survey_id': survey['survey_id'], 'draft': task_input({'task_name': '사진 작업'})})
            photo = application.dispatch('import_photos', {'task_id': task['task_id'], 'file_paths': [str(fixture)]})['added'][0]
            draft = application.task_input(application.dispatch('get_task', {'task_id': task['task_id']})['task'])
            note = {'source': 'INTERVIEW', 'text': '테스트 입력', 'observed_at': '2026-10-23T01:00:00Z'}
            draft['capture']['slots'] = [dict(slot_no=1, photo_id=photo['photo_id'], label='오른쪽 사진', analysis_target=True, camera_side='RIGHT', forward_direction='SCREEN_RIGHT',
                                              side_view_verified='YES', observation=dict(load_state='UNLOADED', measured_load_kg=None, single_person='YES', full_frame='YES', normal_work='YES', evidence=dict(note)))]
            draft['manual'] = [dict(slot_no=1, photo_id=photo['photo_id'], scope=scope, field=field, value=NO_BURDEN[field], source='INTERVIEW', evidence=dict(note))
                               for scope, fields in (('COMMON', COMMON), ('LEFT', ARM), ('RIGHT', ARM)) for field in fields]
            saved = application.dispatch('save_task_input', {'task_id': task['task_id'], 'input': draft})
            analysis = run_analysis(saved, ('pose', 'tags', 'reba', 'judgment'))
            destination = str(Path(root) / 'photo.xlsx')
            application.dispatch('export_excel', {'survey_id': survey['survey_id'], 'spec': {
                'analysis_ids': [analysis['analysis_id']], 'include_representative': True, 'comparisons': [], 'destination': destination}})
            book = load_workbook(destination)
            sheet = book['자세 분석(REBA)']
            values = [c.value for row in sheet.iter_rows() for c in row if c.value not in (None, '')]
            for word in ('사진별 REBA 점수 (점수가 높을수록 몸에 부담이 커요)', '몸통', '위팔', '손목', '오른쪽 사진'):
                self.assertTrue(any(isinstance(v, str) and word in v for v in values), word)
            main = book['조사 결과']
            self.assertEqual(len(main._images), 1, 'the representative photo is embedded where the form asks for it')
            self.assertEqual(len(sheet._images), 0)
            text = '\n'.join(str(c.value) for row in main.iter_rows() for c in row if c.value not in (None, ''))
            for phrase in ('3-1. REBA 자세 평가 — 사진 작업', 'REBA(전신 자세 평가)', '대표 사진 · 오른쪽 사진'):
                self.assertIn(phrase, text)
            for value in values:
                if isinstance(value, str):
                    self.assertIsNone(CODE.search(value), value)


if __name__ == '__main__':
    unittest.main()
