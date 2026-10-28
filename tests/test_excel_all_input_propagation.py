"""Explicit posture answers survive report projection and the actual XLSX writer."""
import copy
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook,load_workbook
from backend.report.plain import _inputs,_manual_rows
from backend.report.renderer import _write_sheet


def source():
    unknown=dict(target='manual.COMMON.static_over_1min',condition='photo:7',photo_id=7,
        value=None,state='UNKNOWN',source='USER',quote='1분 이상 유지하는지는 모르겠어요',updated_at='2026-10-27T01:00:00Z')
    return dict(input_revision=3,conditions=dict(daily_work_min=60,annual_work_days=200,
        intake_session=dict(projected_revision=3,views=dict(RIGHT=7),facts={
            'manual.COMMON.static_over_1min@photo:7':unknown,
            'manual.COMMON.static_over_1min@default':dict(unknown,condition='default',alias_of='manual.COMMON.static_over_1min@photo:7'),
            'manual.COMMON.shock_force@photo:7':dict(unknown,target='manual.COMMON.shock_force',value='NO',state='KNOWN',quote='충격은 없어요'),
            'manual.COMMON.trunk_twisted@photo:7':dict(unknown,target='manual.COMMON.trunk_twisted',source='POSE',quote='측정할 수 없음'),
            'manual.LEFT.coupling@photo:999':dict(unknown,target='manual.LEFT.coupling',condition='photo:999',photo_id=999,quote='다른 사진의 응답')
        })),
        legal={},capture=dict(slots=[dict(slot_no=1,photo_id=7,label='RIGHT',analysis_target=True)]),
        manual=[dict(slot_no=1,photo_id=7,scope='COMMON',field='shock_force',value='NO',source='INTERVIEW')])


class ExcelAllInputPropagation(unittest.TestCase):
    def test_recorded_unknown_is_preserved_but_no_answer_is_not_invented(self):
        value=source();before=copy.deepcopy(value)
        rows=_manual_rows('합성 작업',value)
        unknown=[row for row in rows if row[3]=='모름']
        self.assertEqual(len(unknown),1,'an actual unknown answer must survive even when manual projection has no row')
        self.assertIn('같은 자세를 1분 넘게 유지',unknown[0][2])
        self.assertEqual(unknown[0][4],'직접 입력')
        self.assertIn('1분 이상 유지하는지는 모르겠어요',unknown[0][5])
        self.assertTrue(any('아니오로 기록한 항목' in row[2] and '충격' in row[3] for row in rows))
        self.assertFalse(any('몸통 비틀림' in row[2] or '다른 사진의 응답' in row[5] for row in rows))
        self.assertEqual(value,before,'reporting does not change an input or immutable snapshot')

    def test_stale_or_inactive_photo_unknown_facts_are_not_exported(self):
        for fault in ('revision','deleted_photo','inactive_photo','mismatched_photo'):
            with self.subTest(fault=fault):
                value=source()
                item=value['conditions']['intake_session']['facts']['manual.COMMON.static_over_1min@photo:7']
                if fault=='revision':value['conditions']['intake_session']['projected_revision']=2
                elif fault=='deleted_photo':value['capture']['slots']=[]
                elif fault=='inactive_photo':item['photo_active']=False
                else:item['photo_id']=8
                self.assertFalse(any(row[3]=='모름' for row in _manual_rows('합성 작업',value)))

    def test_an_actual_manual_row_wins_without_an_unknown_duplicate(self):
        value=source()
        value['manual'].append(dict(slot_no=1,photo_id=7,scope='COMMON',field='static_over_1min',value='YES',source='INTERVIEW'))
        rows=_manual_rows('합성 작업',value)
        self.assertFalse(any(row[3]=='모름' for row in rows))
        self.assertTrue(any('같은 자세를 1분 넘게 유지' in row[2] and row[3]=='예' for row in rows))

    def test_same_sheet_definition_writes_the_answer_and_units_into_xlsx(self):
        value=source()
        sheet=_inputs(dict(tasks=[dict(name='합성 작업',source=value)]))
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'입력 전파.xlsx'
            book=Workbook();book.remove(book.active)
            _write_sheet(book,sheet,'합성 검증')
            book.save(path)
            stored=load_workbook(path,read_only=True,data_only=False)
            try:
                rows=[row for row in stored['입력한 값'].iter_rows(values_only=True)]
                for expected in sheet['tables'][0]['rows']:
                    self.assertIn(tuple(expected),rows,'preview row and physical XLSX row must agree')
                text='\n'.join(str(cell) for row in rows for cell in row if cell is not None)
                for fragment in ('60 분/일','200 일/년','모름','충격'):
                    self.assertIn(fragment,text)
                self.assertNotIn('다른 사진의 응답',text)
            finally:stored.close()


if __name__=='__main__':unittest.main()
