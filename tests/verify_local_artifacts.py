"""Independently compare persisted SQLite export snapshots with actual XLSX cells."""
import json
import sqlite3
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from ergo.report.labels import CRITERIA, criterion_name

MARK = {'APPLICABLE': '○', 'NOT_APPLICABLE': 'X', 'NEEDS_INPUT': '입력 필요', 'NEEDS_REVIEW': '검토 필요', 'EXCLUDED': '제외'}


def _criterion_recorded(source, number):
    from ergo.report.input_trace import InputTrace
    trace = InputTrace(source)
    legal = source.get('legal') or {}
    criterion = next((item for item in legal.get('criteria', []) if item.get('criterion_no') == number), {})
    if trace.criterion_presence_recorded(criterion, number):
        return True
    if trace.lookup('complete', criterion.get('rows_complete'), criterion=number):
        return True
    if any(value not in (None, '') for row in criterion.get('rows') or []
           for key, value in row.items() if key not in ('row_id', 'label', 'evidence')):
        return True
    return any(item.get('condition_code', '').startswith(f'C{number}_') and item.get('minutes') is not None
               for item in legal.get('exposures', []))


def _response_has_answers(response):
    for key, value in response.items():
        if key in ('response_code', 'regions'):
            continue
        if value not in (None, '', 'UNKNOWN', []):
            return True
    return any(value not in (None, '', 'UNKNOWN')
               for region in (response.get('regions') or {}).values()
               for value in region.values())


def _reviewed_current_rows(snap):
    """Independently derive review rows from the frozen survey and analysis identities."""
    survey = snap['survey']
    review = survey.get('review') or {}
    if not review.get('reviewed') or review.get('reviewed_metadata_revision') != survey.get('metadata_revision'):
        return []
    reviewed_ids = set(review.get('reviewed_analysis_ids') or [])
    analyses = {analysis['task_id']: analysis for analysis in snap.get('analyses', [])}
    rows = []
    for task in snap.get('tasks', []):
        analysis = analyses.get(task['task_id'])
        if not analysis or analysis.get('analysis_id') not in reviewed_ids:
            continue
        if task.get('input_revision') != analysis.get('input_revision') or task.get('latest_analysis_id') != analysis.get('analysis_id'):
            continue
        task_name = task.get('task_name') or ''
        if task_name == '사진 분석 작업':
            task_name = ''
        reviewed_at = review.get('reviewed_at')
        if reviewed_at:
            text = str(reviewed_at)
            if len(text) == 10 and text[4] == '-' and text[7] == '-':
                shown_at = text
            else:
                try:
                    moment = datetime.fromisoformat(text.replace('Z', '+00:00'))
                except ValueError:
                    shown_at = text
                else:
                    kst = timezone(timedelta(hours=9))
                    moment = moment.astimezone(kst) if moment.tzinfo else moment + timedelta(hours=9)
                    shown_at = moment.strftime('%Y-%m-%d %H:%M')
        else:
            shown_at = ''
        rows.append((task_name, '검토 완료', review.get('reviewer') or '', shown_at, review.get('note') or ''))
    return rows


def expected_sheets(snap):
    """Independently derive data-backed sheet presence from the saved export snapshot."""
    survey = snap['survey']
    analyses = {a['task_id']: a for a in snap.get('analyses', [])}
    sources = [analyses[t['task_id']]['snapshot'] if t['task_id'] in analyses else t for t in snap.get('tasks', [])]
    names = ['조사 결과']

    has_burden_result = False
    for task in snap.get('tasks', []):
        analysis = analyses.get(task['task_id'])
        if not analysis:
            continue
        source = analysis.get('snapshot') or task
        for judgment in ((analysis.get('judgment') or {}).get('judgments') or []):
            status = judgment.get('effective_status') or judgment.get('computed_status')
            if status in ('APPLICABLE', 'NOT_APPLICABLE', 'EXCLUDED', 'NEEDS_REVIEW') and _criterion_recorded(source, judgment['criterion_no']):
                has_burden_result = True
                break
        if has_burden_result:
            break
    if has_burden_result:
        names.append('부담작업 해당 여부')

    symptom = survey.get('symptom_survey') or {}
    responses = symptom.get('responses') or []
    response_has_data = any(_response_has_answers(response) for response in responses)
    symptom_has_data = bool(response_has_data or symptom.get('method') not in (None, '', 'UNKNOWN', 'NOT_DONE')
                            or symptom.get('date') or symptom.get('summary')
                            or (symptom.get('respondent_count') is not None and (not responses or response_has_data)))
    if symptom_has_data:
        names.append('증상조사 요약')

    posture = False
    for analysis in analyses.values():
        posture |= any(any(slot.get(side) for side in ('left', 'right')) for slot in (analysis.get('reba') or {}).get('slots', []))
        for slot in (analysis.get('pose') or {}).get('slots', []):
            posture |= bool(slot.get('engine') or slot.get('angles') or any(item.get('message') for item in slot.get('diagnostics') or []))
            posture |= any(angle.get('value_deg') is not None for angle in slot.get('world_angles') or [])
        posture |= any(tag.get('observed_count') for tag in analysis.get('tags') or [])
    if posture:
        names.append('자세 분석(REBA)')

    input_fields = {'daily_work_min', 'annual_work_days', 'one_off', 'planned_duration_months', 'work_height_cm',
                    'worker_code', 'method_description', 'work_posture', 'work_intensity', 'work_frequency',
                    'hazard_factors', 'hazard_causes', 'improvement_recommendations'}
    has_input = False
    for source in sources:
        conditions = source.get('conditions') or {}
        facts = ((conditions.get('intake_session') or {}).get('facts') or {}).values()
        fact_targets = {fact.get('target') for fact in facts}
        has_input |= any(conditions.get(field) not in (None, '', 'UNKNOWN') or field in fact_targets for field in input_fields)
        has_input |= any(isinstance(target, str) and target.startswith('work_type.') for target in fact_targets)
        legal = source.get('legal') or {}
        from ergo.report.input_trace import InputTrace
        trace = InputTrace(source)
        has_input |= any(trace.criterion_presence_recorded(item, item.get('criterion_no'))
                         or trace.lookup('complete', item.get('rows_complete'), criterion=item.get('criterion_no'))
                         or any(value not in (None, '') for row in item.get('rows') or []
                                for key, value in row.items() if key not in ('row_id', 'label', 'evidence'))
                         for item in legal.get('criteria') or [])
        has_input |= any(item.get('minutes') is not None for item in legal.get('exposures') or [])
        has_input |= bool(source.get('manual'))
        for slot in (source.get('capture') or {}).get('slots', []):
            observation = slot.get('observation') or {}
            has_input |= bool(slot.get('photo_id') is not None and any(
                value not in (None, '') for value in
                (slot.get('camera_side'), slot.get('forward_direction'), slot.get('side_view_verified'),
                 slot.get('analysis_target'), observation.get('single_person'), observation.get('full_frame'),
                 observation.get('normal_work'), observation.get('load_state'), observation.get('measured_load_kg'))))
    if has_input:
        names.append('입력한 값')

    review_rows = _reviewed_current_rows(snap)
    if review_rows:
        names.append('사람 검토')

    if any(ai.get('is_active') and ai.get('applied_text') for a in analyses.values() for ai in a.get('ai') or []):
        names.append('결과 설명')
    if snap.get('comparison_specs'):
        names.append('개선 전후 비교')

    from ergo.intake.contract import SESSION_KEY
    for source in sources:
        photos = (((source.get('conditions') or {}).get(SESSION_KEY) or {}).get('photo_vision') or {}).get('photos') or []
        owners = {slot.get('photo_id') for slot in (source.get('capture') or {}).get('slots', [])}
        if any(photo.get('photo_id') in owners and photo.get('photo_active') is not False
               and any(isinstance(item, dict) for item in (photo.get('observations') or {}).values())
               for photo in photos):
            names.append('AI 사진 관찰')
            break
    return names

def verify(directory, minimum_exports=5):
    from ergo.report.form import FINAL
    from ergo.report.labels import criterion_name
    from ergo.report.plain import change_label, change_value
    directory = Path(directory).resolve()
    db = sqlite3.connect('file:' + (directory/'loco.db').as_posix() + '?mode=ro', uri=True)
    db.row_factory = sqlite3.Row
    tables = {r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")}
    assert tables == {'SURVEY','TASK','PHOTO','ANALYSIS','APP_SETTINGS','EXPORT_LOG'}, tables
    outputs=[]
    for log in db.execute('SELECT * FROM EXPORT_LOG ORDER BY export_id'):
        snap=json.loads(log['snapshot_json']); receipt=json.loads(log['receipt_json'])
        workbook=load_workbook(log['file_path'], data_only=False)
        expected=expected_sheets(snap)
        assert workbook.sheetnames==expected==receipt['sheets'], (workbook.sheetnames, expected, receipt['sheets'])
        comparison_analyses={a['analysis_id']:a for a in snap.get('comparison_analyses',[])}
        for analysis in comparison_analyses.values():
            stored=json.loads(db.execute('SELECT snapshot_json FROM ANALYSIS WHERE analysis_id=?',(analysis['analysis_id'],)).fetchone()[0])
            assert analysis['snapshot']==stored
        comparison_checks=[]
        for spec in snap.get('comparison_specs',[]):
            before=comparison_analyses[spec['before_analysis_id']];after=comparison_analyses[spec['after_analysis_id']]
            bmax=(before.get('reba') or {}).get('summary',{}).get('full_max');amax=(after.get('reba') or {}).get('summary',{}).get('full_max')
            if bmax is not None and amax is not None:
                delta=amax-bmax
                cells={str(cell.value) for row in workbook['개선 전후 비교'].iter_rows() for cell in row if cell.value is not None}
                assert (f'+{delta}' if delta>0 else str(delta)) in cells, 'the score change is missing from the before/after sheet'
                bm={(m['slot_no'],m['scope'],m['field']):m.get('value') for m in before['snapshot']['manual']}
                am={(m['slot_no'],m['scope'],m['field']):m.get('value') for m in after['snapshot']['manual']}
                for key in bm.keys()&am.keys():
                    if bm[key]!=am[key]:
                        path=f'manual.{key[0]}.{key[1]}.{key[2]}.value'
                        assert change_label(path) in cells and change_value(path,bm[key]) in cells and change_value(path,am[key]) in cells, path
                comparison_checks.append(dict(before_analysis_id=before['analysis_id'],after_analysis_id=after['analysis_id'],before_max=bmax,after_max=amax,delta=delta))
        expected_results=set()
        for analysis in snap['analyses']:
            stored=json.loads(db.execute('SELECT snapshot_json FROM ANALYSIS WHERE analysis_id=?',(analysis['analysis_id'],)).fetchone()[0])
            assert analysis['snapshot']==stored
            task=next(item for item in snap['tasks'] if item['task_id']==analysis['task_id'])
            source=analysis.get('snapshot') or task
            for judgment in ((analysis.get('judgment') or {}).get('judgments') or []):
                status=judgment.get('effective_status') or judgment.get('computed_status')
                number=judgment['criterion_no']
                if status in ('APPLICABLE','NOT_APPLICABLE','EXCLUDED','NEEDS_REVIEW') and _criterion_recorded(source,number):
                    expected_results.add((task.get('task_name') or task.get('name') or '', f"{criterion_name(number)}({number}호)",
                                          {'APPLICABLE':'부담작업','NOT_APPLICABLE':'부담작업 아님','EXCLUDED':'제외','NEEDS_REVIEW':'검토 필요'}[status]))
        if expected_results:
            assert '부담작업 해당 여부' in workbook.sheetnames
            rows=list(workbook['부담작업 해당 여부'].iter_rows(values_only=True))
            header=next(i for i,row in enumerate(rows) if row[:3]==('작업','기준','계산 결과'))
            criterion_labels={f'{criterion_name(number)}({number}호)' for number in CRITERIA}
            result_labels={'부담작업','부담작업 아님','제외','검토 필요'}
            actual_results={tuple(row[:3]) for row in rows[header+1:]
                            if row[0] not in (None,'') and row[1] in criterion_labels and row[2] in result_labels}
            assert actual_results==expected_results,(actual_results,expected_results)
        review_rows = _reviewed_current_rows(snap)
        if review_rows:
            ws = workbook['사람 검토']
            rows = list(ws.iter_rows(values_only=True))
            header = next(i for i, row in enumerate(rows) if row[:5] == ('작업명', '검토 상태', '검토자', '검토 시각', '메모'))
            actual_review_rows = [tuple(row[:5]) for row in rows[header + 1:] if row[1] == '검토 완료']
            assert actual_review_rows == review_rows, (actual_review_rows, review_rows)
        for sheet in workbook:
            all_values='\n'.join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value not in (None,''))
            for marker in ('입력 필요','입력되지 않음','미응답','확인사항'):
                assert marker not in all_values,(sheet.title,marker)
            assert all(cell.data_type!='f' for row in sheet for cell in row), 'Unexpected Excel formula'
        outputs.append(dict(export_id=log['export_id'],file=log['file_path'],sheets=expected,
                            analysis_count=len(snap['analyses']),metadata_revision=snap['survey']['metadata_revision'],comparisons=comparison_checks))
        workbook.close()
    assert len(outputs)>=minimum_exports, 'The UI run did not produce the requested number of workbooks'
    models=[]
    for row in db.execute('SELECT analysis_id,result_json FROM ANALYSIS'):
        result=json.loads(row['result_json'])
        for pose in (result.get('pose') or {}).get('slots',[]):
            if pose.get('points'):
                model=pose.get('engine') or {}
                expected=model.get('point_count',17)
                assert len(pose['points'])==expected
                if model.get('provider')=='mediapipe':
                    assert len(pose['world_points'])==expected
                    assert all(not a['used_for_reba'] for a in pose['world_angles'])
                models.append(dict(analysis_id=row['analysis_id'],slot=pose['slot_no'],status=pose['status'],model=model.get('model_id'),points=expected))
    assert models
    records={table:db.execute('SELECT COUNT(*) FROM '+table).fetchone()[0] for table in sorted(tables)}
    db.close()
    result=dict(kind='read-only independent SQLite/XLSX comparison',records=records,exports=outputs,model_results=models)
    (directory/'artifact-receipt.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(dict(directory=str(directory),records=records,workbooks=len(outputs),model_results=len(models),verdict='PASS'),ensure_ascii=False))


if __name__=='__main__':
    verify(sys.argv[1], int(sys.argv[2]) if len(sys.argv)>2 else 5)
