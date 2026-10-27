"""Excel output contains recorded survey inputs and data-backed analysis results only."""
import re
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from ergo.report.form import CHECK_RULE, _next_input_text, _support_overview
from ergo.report.labels import CRITERION_RULE
from ergo.service import application
from ergo.engine.task_defaults import task_input
from tests.report_helpers import lifting_input, make_survey, run_analysis, update_survey

LAW_TEXT = Path(__file__).parent / 'fixtures' / '고시_제2020-12호_본문.txt'
NUMBER = re.compile(r'\d+(?:\.\d+)?')


def text_of(sheet):
    return [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value not in (None, '')]


def row_with(sheet, first):
    return next((tuple(cell.value for cell in row) for row in sheet.iter_rows() if row[0].value == first), None)


class FormWorkbook(unittest.TestCase):
    def test_next_input_translation_still_groups_the_in_app_followup(self):
        fields = ['common.neck_base', 'common.trunk_base', 'common.legs_support', 'left.wrist_base', 'right.wrist_base']
        text = _next_input_text({'next_inputs': [dict(kind='MISSING_INPUT', slot_no=1, field_paths=fields, message='long')]})
        for label in ['목·몸통', '다리', '왼팔·손목', '오른팔·손목', '목 굽힘', '몸통 굽힘', '다리 지지']:
            self.assertIn(label, text)
        self.assertEqual(text.count('손목 꺾임'), 2)
        self.assertEqual(_next_input_text({'next_inputs': [dict(kind='REVIEW_REQUIRED', field_paths=fields, message='사진을 다시 검토하세요')]}), '사진을 다시 검토하세요')
        self.assertEqual(_next_input_text({'next_inputs': []}), '')

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        patcher = patch.object(application, 'DB', str(Path(self.temp.name) / 'loco.db'))
        patcher.start()
        self.addCleanup(patcher.stop)
        self.addCleanup(self.temp.cleanup)
        self.survey = make_survey('양식 검증', survey_type='OCCASIONAL')
        self.task = application.dispatch('create_task', {'survey_id': self.survey['survey_id'], 'draft': lifting_input()})
        self.analysis = run_analysis(self.task)

    def export(self):
        destination = str(Path(self.temp.name) / 'report.xlsx')
        application.dispatch('export_excel', {'survey_id': self.survey['survey_id'], 'spec': {
            'analysis_ids': [], 'include_representative': False, 'comparisons': [], 'destination': destination}})
        return load_workbook(destination, data_only=False)

    def preview(self):
        return application.dispatch('preview_excel_sheets', {'survey_id': self.survey['survey_id']})['sheets']

    def rerun(self, **conditions):
        task = application.dispatch('get_task', {'task_id': self.task['task_id']})['task']
        updated = application.task_input(task)
        updated['conditions'].update(conditions)
        run_analysis(application.dispatch('save_task_input', {'task_id': task['task_id'], 'input': updated}))

    def test_quickflow_machine_task_name_is_not_exported(self):
        task = application.dispatch('create_task', {'survey_id': self.survey['survey_id'],
            'draft': task_input({'task_name': '사진 분석 작업'})})
        run_analysis(task)
        book = self.export()
        text = '\n'.join(str(value) for sheet in book for value in text_of(sheet))
        self.assertNotIn('사진 분석 작업', text)

    def test_unmodified_survey_type_is_preserved_when_editor_omits_it(self):
        current = application.dispatch('get_survey', {'survey_id': self.survey['survey_id']})
        current.pop('survey_type', None)
        updated = application.dispatch('update_survey', {'survey_id': self.survey['survey_id'], 'draft': current})
        self.assertEqual(updated['survey_type'], 'OCCASIONAL')
        self.assertEqual(updated['survey_type_source'], 'USER')

    def test_unknown_legacy_survey_type_is_not_exported_as_user_answer(self):
        for survey_type in ('INITIAL', 'OCCASIONAL'):
            sheet=_support_overview({'survey': {'survey_type': survey_type, 'survey_type_source': 'UNKNOWN'}, 'tasks': []})
            self.assertIsNone(sheet, f'{survey_type} without user provenance must not appear in Excel')

    def test_backend_default_survey_type_is_not_exported_as_a_user_answer(self):
        survey = application.dispatch('create_survey', {'draft': {'title': '기본 유형 조사'}})
        self.assertEqual(survey['survey_type'], 'REGULAR')
        self.assertEqual(survey['survey_type_source'], 'DEFAULT')
        task = application.dispatch('create_task', {'survey_id': survey['survey_id'],
            'draft': task_input({'task_name': '기본 작업'})})
        run_analysis(task)
        destination = str(Path(self.temp.name) / 'default-type.xlsx')
        application.dispatch('export_excel', {'survey_id': survey['survey_id'], 'spec': {
            'analysis_ids': [], 'include_representative': False, 'comparisons': [], 'destination': destination}})
        book = load_workbook(destination, data_only=False)
        try:
            self.assertNotIn('조사 구분', text_of(book['조사 결과']))
        finally:
            book.close()

    def test_only_selected_survey_answers_are_written(self):
        book = self.export()
        self.assertIn('조사 결과', book.sheetnames)
        self.assertNotIn('유해요인조사표', book.sheetnames)
        text = '\n'.join(text_of(book['조사 결과']))
        self.assertIn('수시조사', text)
        self.assertNotIn('☐', text)
        self.assertNotIn('입력 필요', text)
        self.assertNotIn('이전 조사일', text)
        self.assertNotIn('근로자 참여', text)

    def test_partially_answered_scale_has_no_empty_unasked_column(self):
        self.rerun(work_intensity=3)
        book = self.export()
        values = text_of(book['조사 결과'])
        self.assertIn('작업강도', values)
        self.assertIn('3', values)
        self.assertNotIn('작업빈도', values)
        self.assertNotIn('총점수', values)
        self.assertNotIn('입력 필요', values)

    def test_burden_results_show_only_recorded_criteria(self):
        book = self.export()
        self.assertIn('부담작업 해당 여부', book.sheetnames)
        rows = list(book['부담작업 해당 여부'].iter_rows(values_only=True))
        header = next(index for index, row in enumerate(rows) if row[:3] == ('작업', '기준', '계산 결과'))
        rows = rows[header:]
        text = '\n'.join(str(value) for row in rows for value in row if value not in (None, ''))
        self.assertIn('키보드·마우스 집중 입력', text)
        self.assertIn('부담작업', text)
        self.assertIn('25kg 이상 들기', text)
        self.assertIn('부담작업 아님', text)
        for marker in ('입력 필요', '입력 안 됨', '아직 답하지 않음', '미확정'):
            self.assertNotIn(marker, text)

    def test_unanswered_criterion_defaults_do_not_create_a_burden_sheet(self):
        application.dispatch('delete_task', {'task_id': self.task['task_id']})
        task = application.dispatch('create_task', {'survey_id': self.survey['survey_id'],
            'draft': task_input({'task_name': '빈 입력 작업'})})
        run_analysis(task)
        book = self.export()
        self.assertNotIn('부담작업 해당 여부', book.sheetnames)
        self.assertNotIn('입력 필요', '\n'.join(value for name in book.sheetnames for value in text_of(book[name])))
        self.assertNotIn('확인사항', book.sheetnames)

    def test_the_workbook_has_no_blank_official_template(self):
        book = self.export()
        self.assertNotIn('증상조사표', book.sheetnames)
        self.assertNotIn('유해요인조사표', book.sheetnames)
        text = '\n'.join(value for name in book.sheetnames for value in text_of(book[name]))
        for marker in ('수기 작성용', '직접 작성', '입력 필요', '입력되지 않음'):
            self.assertNotIn(marker, text)

    def test_missing_photo_and_reba_sections_are_omitted(self):
        main = '\n'.join(text_of(self.export()['조사 결과']))
        self.assertNotIn('대표 사진이 없어요', main)
        self.assertNotIn('사진 또는 그림', main)
        self.assertNotIn('REBA 자세 점수', main)

    def test_preview_and_file_share_the_data_only_sheets_and_cells(self):
        preview = self.preview()
        book = self.export()
        self.assertEqual([sheet['name'] for sheet in preview], book.sheetnames)
        for sheet in preview:
            values = {str(value) for value in text_of(book[sheet['name']])}
            for table in sheet['tables']:
                for row in table['rows']:
                    for value in row:
                        if value not in ('', None):
                            self.assertIn(str(value), values, f"{sheet['name']}: {value!r}")

    def test_rule_phrases_carry_exactly_the_numbers_of_the_preserved_notice(self):
        body = LAW_TEXT.read_text(encoding='utf-8').split('=== 본문 ===\n', 1)[1]
        section = body[body.index('제3조(근골격계부담작업)'):body.index('제4조(유해요인조사 방법)')]
        items = {int(number): ' '.join(text.split()) for number, text in re.findall(r'^([0-9]+)\. (.+)$', section, re.M)}
        self.assertEqual(sorted(items), list(range(1, 12)))
        for number, item in items.items():
            law = sorted(NUMBER.findall(item))
            self.assertEqual(sorted(set(NUMBER.findall(CRITERION_RULE[number]))), sorted(set(law)), f'{number}호 기준 문구의 숫자')
            self.assertEqual(sorted(set(NUMBER.findall(CHECK_RULE[number]))), sorted(set(law)), f'{number}호 체크리스트 문구의 숫자')


if __name__ == '__main__':
    unittest.main()
