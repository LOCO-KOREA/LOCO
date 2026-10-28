"""A saved investigator conclusion is an entered result even without raw quantities."""
import copy
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from backend.service.application import create_service


class ReviewConclusionExport(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.ctx=create_service(str(Path(self.temp.name)/'review.db'))
        self.survey=self.ctx.dispatch('create_survey',dict(draft=dict(title='결론 출력 검증')))
        task=self.ctx.dispatch('create_task',dict(survey_id=self.survey['survey_id'],draft=dict(task_name='원자료 미완성 작업')))
        self.analysis=self.ctx.dispatch('recalculate_judgments',dict(task_id=task['task_id']))
        self.identity=self.analysis['analysis_id'];self.snapshot=copy.deepcopy(self.analysis['snapshot'])

    def select(self,number,status):
        return self.ctx.dispatch('set_investigator_conclusion',dict(analysis_id=self.identity,conclusion=dict(
            criterion_no=number,status=status,reason=None,evidence=dict(source='DIRECT_OBSERVATION',text='검토 화면 선택'),reviewer=None,reviewed_at=None)))

    def preview(self):
        return self.ctx.dispatch('preview_excel_sheets',dict(survey_id=self.survey['survey_id'],analysis_ids=[self.identity],comparisons=[]))

    def test_selected_yes_and_no_export_without_requiring_the_missing_raw_inputs(self):
        self.select(1,'APPLICABLE');saved=self.select(2,'NOT_APPLICABLE')
        self.assertEqual(saved['judgment']['judgments'][0]['computed_status'],'NEEDS_INPUT')
        sheet=next(s for s in self.preview()['sheets'] if s['name']=='부담작업 해당 여부')
        rows=sheet['tables'][0]['rows']
        self.assertEqual(len(rows),2)
        self.assertEqual([r[2] for r in rows],['부담작업','부담작업 아님'])
        destination=Path(self.temp.name)/'결론.xlsx'
        self.ctx.dispatch('export_excel',dict(survey_id=self.survey['survey_id'],spec=dict(analysis_ids=[self.identity],include_representative=False,comparisons=[],destination=str(destination))))
        book=load_workbook(destination,read_only=True)
        cells=list(book['부담작업 해당 여부'].iter_rows(values_only=True))
        for row in rows:self.assertTrue(any(tuple(values[:3])==tuple(row) for values in cells),row)
        book.close()
        self.assertEqual(self.ctx.require('analysis',self.identity)['snapshot'],self.snapshot)
        self.select(2,None)
        remaining=next(s for s in self.preview()['sheets'] if s['name']=='부담작업 해당 여부')['tables'][0]['rows']
        self.assertEqual(remaining,[rows[0]],'a withdrawn conclusion must not remain in the next export')

    def test_unasked_and_unselected_unknown_criteria_remain_absent(self):
        self.assertFalse(any(s['name']=='부담작업 해당 여부' for s in self.preview()['sheets']))

    def test_changing_reasoning_only_keeps_the_saved_review_current(self):
        task_id=self.analysis['task_id']
        settings=self.ctx.store('get_settings');settings['gpt_reasoning_effort']='xhigh'
        self.ctx.dispatch('update_settings',dict(settings=settings))
        self.ctx.dispatch('intake_session',dict(task_id=task_id))
        self.identity=self.ctx.dispatch('recalculate_judgments',dict(task_id=task_id))['analysis_id']
        self.select(1,'APPLICABLE');before=self.ctx.require('task',task_id)
        settings=self.ctx.store('get_settings');settings['gpt_reasoning_effort']='high'
        self.ctx.dispatch('update_settings',dict(settings=settings))
        reopened=self.ctx.dispatch('intake_session',dict(task_id=task_id))
        current=self.ctx.dispatch('get_analysis',dict(analysis_id=self.identity))
        self.assertEqual(reopened['task']['input_revision'],before['input_revision'])
        self.assertEqual(current['currentness'],'CURRENT')
        self.assertEqual(current['judgment']['judgments'][0]['investigator']['status'],'APPLICABLE')
        self.assertEqual(self.ctx.settings()['gpt_reasoning_effort'],'high')


if __name__=='__main__':unittest.main()
