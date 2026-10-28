"""End-to-end 2.0 acceptance path from direct input through Korean Excel output."""
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from backend.engine.task_defaults import task_input
from backend.service import application


class V2Acceptance(unittest.TestCase):
    def test_box_task_can_run_without_photos_without_exporting_unanswered_defaults(self):
        with tempfile.TemporaryDirectory() as root, patch.object(application, 'DB', str(Path(root) / 'loco.db')):
            survey = application.dispatch('create_survey', {'draft': {
                'survey_type': 'REGULAR', 'survey_date': '2026-10-23', 'investigator': '조사자',
                'department': '팀', 'process': '공정', 'title': '직접입력 검증'}})
            draft = task_input({'task_name': '=2+2', 'input_template': 'LIFTING'})
            task = application.dispatch('create_task', {'survey_id': survey['survey_id'], 'draft': draft})
            analysis = application.new_analysis(task, ['judgment'])
            application.run(analysis, ['judgment'], threading.Event(), application._context.analysis_token(analysis['analysis_id']))
            current = application.analysis(analysis['analysis_id'])
            self.assertEqual(current['status'], 'FINISHED')
            self.assertEqual(len(current['judgment']['judgments']), 11)
            self.assertTrue(all(x['computed_status'] == 'NEEDS_INPUT' for x in current['judgment']['judgments']))
            self.assertEqual(current['module_states']['pose'], 'NOT_REQUESTED')

            preview = application.dispatch('preview_excel', {'survey_id': survey['survey_id'], 'analysis_ids': []})
            self.assertEqual(preview['counts']['needs_input_count'], 11)
            destination = str(Path(root) / 'loco_report.xlsx')
            result = application.dispatch('export_excel', {'survey_id': survey['survey_id'], 'spec': {
                'analysis_ids': [], 'include_representative': False, 'comparisons': [], 'destination': destination}})
            self.assertEqual(result['sheets'], ['조사 결과'])
            workbook = load_workbook(destination, data_only=False)
            try:
                self.assertEqual(workbook.sheetnames, ['조사 결과'])
                all_text = ' '.join(str(cell.value) for sheet in workbook for row in sheet.iter_rows() for cell in row if cell.value is not None)
                self.assertNotIn('아직 답하지 않은 기준', all_text)
                for marker in ('입력 필요', '입력 안 됨', '입력되지 않음', '확인사항'):
                    self.assertNotIn(marker, all_text)
                # A task name that looks like a formula must stay text everywhere it is written.
                formula_like = [cell for sheet in workbook for row in sheet.iter_rows() for cell in row if cell.value == '=2+2']
                self.assertTrue(formula_like)
                self.assertTrue(all(cell.data_type == 's' for cell in formula_like))
                self.assertTrue(any(cell.value == '=2+2' for row in workbook['조사 결과'].iter_rows() for cell in row))
            finally:
                workbook.close()

    def test_before_after_analysis_keeps_criterion_changes_and_limits_reba_delta(self):
        with tempfile.TemporaryDirectory() as root, patch.object(application, 'DB', str(Path(root) / 'compare.db')):
            survey = application.dispatch('create_survey', {'draft': {
                'survey_type': 'REGULAR', 'survey_date': '2026-10-23', 'investigator': '조사자',
                'department': '팀', 'process': '공정', 'title': '전후 비교'}})
            task = application.dispatch('create_task', {
                'survey_id': survey['survey_id'], 'draft': task_input({'task_name': '조립 작업'})})
            before = application.new_analysis(task, ['judgment'])
            application.run(before, ['judgment'], threading.Event(), application._context.analysis_token(before['analysis_id']))

            task = application.dispatch('get_task', {'task_id': task['task_id']})['task']
            updated = application.task_input(task)
            updated['conditions'].update(one_off='NO', planned_duration_months=12,
                                         annual_work_days=220, daily_work_min=480)
            for exposure in updated['legal']['exposures']:
                exposure['minutes'] = 0
            for criterion in updated['legal']['criteria']:
                criterion['condition_present'] = 'NO'
                if criterion['criterion_no'] >= 6:
                    criterion['rows_complete'] = 'YES'
            task = application.dispatch('save_task_input', {'task_id': task['task_id'], 'input': updated})
            after = application.new_analysis(task, ['judgment'])
            application.run(after, ['judgment'], threading.Event(), application._context.analysis_token(after['analysis_id']))

            overview = application.dispatch('preview_excel', {
                'survey_id': survey['survey_id'], 'analysis_ids': [],
                'comparisons': [{'before_analysis_id': before['analysis_id'],
                                 'after_analysis_id': after['analysis_id'], 'note': '조건 입력 후 재검토'}],
            })
            comparison = overview['comparisons'][0]
            criterion_one = comparison['criteria'][0]
            self.assertEqual(criterion_one['before_computed_status'], 'NEEDS_INPUT')
            self.assertEqual(criterion_one['after_computed_status'], 'NOT_APPLICABLE')
            self.assertFalse(comparison['sampling_compatible'])
            self.assertIsNone(comparison['reba_max_delta'])
            self.assertTrue(any(item['field_path'] == 'conditions.daily_work_min'
                                for item in comparison['condition_changes']))


if __name__ == '__main__':
    unittest.main()
