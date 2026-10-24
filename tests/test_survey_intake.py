"""Adversarial checks for the investigation context and conditional appendix export."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from ergo.engine.survey_defaults import blank_response, normalize_survey
from ergo.report import workbook
from ergo.service import application


def base_survey(**updates):
    value={
        'survey_type':'REGULAR','survey_date':'2026-10-23','investigator':'조사자',
        'department':'팀','process':'공정','title':'설문 검증','occasional_reason':None,
    }
    value.update(updates)
    return value


class SurveyIntakeTests(unittest.TestCase):
    def test_new_occasional_survey_requires_a_cause_but_legacy_read_remains_possible(self):
        with self.assertRaisesRegex(ValueError,'사유 유형'):
            normalize_survey(base_survey(survey_type='OCCASIONAL'),require_occasional_cause=True)
        legacy=normalize_survey(base_survey(survey_type='OCCASIONAL'))
        self.assertIsNone(legacy['occasional_cause'])
        self.assertEqual(legacy['survey_type_source'], 'UNKNOWN')
        self.assertEqual(legacy['occasional_cause_source'], 'UNKNOWN')

    def test_health_questionnaire_rejects_duplicate_codes_and_invalid_region_values(self):
        first=blank_response(1)
        second=blank_response(2)
        survey=base_survey(symptom_survey={'responses':[first,second]})
        survey['symptom_survey']['responses'][1]['response_code']=first['response_code']
        with self.assertRaisesRegex(ValueError,'중복'):
            normalize_survey(survey)
        survey=base_survey(symptom_survey={'responses':[blank_response(1)]})
        survey['symptom_survey']['responses'][0]['regions']['neck']['severity']='DIAGNOSIS'
        with self.assertRaisesRegex(ValueError,'severity'):
            normalize_survey(survey)

    def test_symptom_sheet_contains_only_recorded_responses_for_any_survey_cause(self):
        for cause in ('NEW_TASK_OR_EQUIPMENT', 'DISEASE_CASE'):
            with self.subTest(cause=cause), tempfile.TemporaryDirectory() as folder:
                survey = normalize_survey(base_survey(survey_type='OCCASIONAL', occasional_cause=cause))
                survey['survey_type_source'] = 'USER'
                survey['occasional_cause_source'] = 'USER'
                survey['participation'] = {'status': 'YES', 'role': 'AFFECTED_WORKER', 'participant_code': 'W-1',
                    'method': 'INTERVIEW', 'date': '2026-10-23', 'notes': '작업 방법 의견'}
                survey['workplace_situation'] = {
                    'equipment_status': 'NO', 'equipment_notes': '', 'workload_status': 'YES', 'workload_notes': '물량 증가',
                    'speed_status': 'NO', 'speed_notes': '', 'work_pattern_status': 'UNKNOWN', 'work_pattern_notes': '',
                }
                response = blank_response(1)
                response.update(work_related_symptoms_last_year='YES')
                response['regions']['neck'].update(side='BOTH', duration='WEEK_TO_MONTH', severity='MODERATE', last_week='YES')
                survey['symptom_survey'] = {'method': 'QUESTIONNAIRE', 'date': '2026-10-23',
                    'respondent_count': 1, 'summary': '익명 설문 1건', 'responses': [response]}
                survey.update(survey_id=1, metadata_revision=0, review={'reviewed': False})
                destination = str(Path(folder) / 'report.xlsx')
                snapshot = {'survey': survey, 'tasks': [], 'analyses': [], 'unanalysed_task_ids': [],
                    'include_representative': False, 'comparison_specs': [], 'comparison_analyses': [],
                    'exported_at': '2026-10-23T00:00:00+00:00'}
                rendered = workbook.render(snapshot, destination)
                self.assertEqual(rendered['sheets'], ['조사 결과', '증상조사 요약'])
                book = load_workbook(destination, data_only=False)
                try:
                    main = [str(cell.value) for row in book['조사 결과'].iter_rows()
                            for cell in row if cell.value not in (None, '')]
                    text = '\n'.join(main)
                    self.assertIn('수시조사', text)
                    selected = '새로운 작업·설비를 도입한 경우' if cause == 'NEW_TASK_OR_EQUIPMENT' else '근골격계질환자가 발생했거나 업무상 질병으로 인정받은 경우'
                    self.assertIn(selected, text)
                    self.assertNotIn('☐', text)
                    self.assertIn('W-1', text)
                    self.assertIn('작업 방법 의견', text)
                    self.assertIn('물량 증가', text)
                    symptoms = [cell.value for row in book['증상조사 요약'].iter_rows() for cell in row if cell.value not in (None, '')]
                    self.assertIn('응답 1 · 익명-001', symptoms)
                    self.assertIn('양쪽', symptoms)
                    self.assertIn('1주일~1개월 미만', symptoms)
                    self.assertIn('중간 정도 통증', symptoms)
                    for marker in ('입력 필요', '입력되지 않음', '수기 작성용', '직접 작성'):
                        self.assertNotIn(marker, text + '\n'.join(map(str, symptoms)))
                    for code in ('BOTH', 'WEEK_TO_MONTH', 'MODERATE', 'UNKNOWN'):
                        self.assertNotIn(code, symptoms)
                finally:
                    book.close()

    def test_partial_symptom_response_does_not_export_default_unknown_fields(self):
        response = blank_response(1)
        response['regions']['neck']['side'] = 'LEFT'
        survey = normalize_survey(base_survey(symptom_survey={
            'method': 'QUESTIONNAIRE', 'respondent_count': 1, 'responses': [response]}))
        survey.update(survey_id=1, metadata_revision=0, review={'reviewed': False})
        with tempfile.TemporaryDirectory() as folder:
            destination = str(Path(folder) / 'partial-symptom.xlsx')
            snapshot = {'survey': survey, 'tasks': [], 'analyses': [], 'unanalysed_task_ids': [],
                'include_representative': False, 'comparison_specs': [], 'comparison_analyses': [],
                'exported_at': '2026-10-23T00:00:00+00:00'}
            workbook.render(snapshot, destination)
            book = load_workbook(destination, data_only=False)
            try:
                values = [cell.value for row in book['증상조사 요약'].iter_rows()
                          for cell in row if cell.value not in (None, '')]
                self.assertIn('왼쪽', values)
                self.assertNotIn('모름', values)
                self.assertNotIn('미응답', values)
                for site in ('어깨', '팔/팔꿈치', '손/손목/손가락', '허리', '다리/발'):
                    self.assertFalse(any(isinstance(value, str) and value.startswith(site + ' ·') for value in values))
            finally:
                book.close()

    def test_a_survey_without_responses_has_no_symptom_sheet(self):
        with tempfile.TemporaryDirectory() as folder:
            survey=normalize_survey(base_survey())
            survey.update(survey_id=1,metadata_revision=0,review={'reviewed':False})
            snapshot={'survey':survey,'tasks':[],'analyses':[],'unanalysed_task_ids':[],'include_representative':False,'comparison_specs':[],'comparison_analyses':[],
                      'exported_at':'2026-10-23T00:00:00+00:00'}
            rendered=workbook.render(snapshot,str(Path(folder)/'report.xlsx'))
            self.assertNotIn('증상조사표',rendered['sheets']); self.assertNotIn('증상조사 요약',rendered['sheets'])
            self.assertEqual(rendered['sheets'], ['조사 결과'])

    def test_incomplete_survey_context_is_reported_instead_of_silent_completion(self):
        survey=normalize_survey(base_survey())
        snapshot={'survey':survey,'tasks':[],'analyses':[],'unanalysed_task_ids':[],
                  'include_representative':False,'comparison_specs':[],'comparison_analyses':[],
                  'exported_at':'2026-10-23T00:00:00+00:00'}
        result=workbook.overview(snapshot)
        self.assertFalse(result['state']['survey_input_complete'])
        self.assertGreaterEqual(len(result['survey_followups']),6)


if __name__=='__main__':
    unittest.main()
