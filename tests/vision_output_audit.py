"""Read-only DB, immutable snapshot and exact per-photo vision worksheet audit."""
import argparse
import json
from pathlib import Path
import re
import sqlite3
from openpyxl import load_workbook

def run(directory):
    receipt=json.loads((directory/'receipt.json').read_text(encoding='utf-8'));vision=receipt['vision'];checks=[]
    def check(name,wanted,actual,**context):checks.append(dict(name=name,expected=wanted,actual=actual,pass_=wanted==actual,**context))
    with sqlite3.connect('file:'+(directory/'loco.db').as_posix()+'?mode=ro',uri=True) as db:
        task=json.loads(db.execute('SELECT input_json FROM TASK WHERE task_id=?',(receipt['task_id'],)).fetchone()[0]);snapshot=json.loads(db.execute('SELECT snapshot_json FROM ANALYSIS WHERE analysis_id=?',(receipt['analysis_id'],)).fetchone()[0]);export=json.loads(db.execute('SELECT snapshot_json FROM EXPORT_LOG ORDER BY export_id DESC LIMIT 1').fetchone()[0])
    check('physical TASK preserves exact photo observations',vision,task['conditions']['intake_session']['photo_vision'])
    check('physical immutable ANALYSIS preserves exact photo observations',vision,snapshot['conditions']['intake_session']['photo_vision'])
    check('physical EXPORT_LOG uses selected immutable analysis',snapshot,export['analyses'][0]['snapshot'])
    check('no image payload persisted in TASK/analysis/export JSON',False,'data:image/' in json.dumps([task,snapshot,export]))
    labels={row['key']:row['label'] for row in vision['catalogue']};book=load_workbook(receipt['workbook'],data_only=False)
    check('optional AI photo worksheet exists',True,'AI 사진 관찰' in book.sheetnames)
    sheet=book['AI 사진 관찰'];rows=[[cell.value for cell in row] for row in sheet];ledger=[]
    for photo in vision['photos']:
        for key,item in photo['observations'].items():
            candidates=[(number,row) for number,row in enumerate(rows,1) if len(row)>=6 and row[1]==photo['photo_label'] and row[2]==labels[key]]
            check('exactly one worksheet row for this photo and observation',1,len(candidates),photo_id=photo['photo_id'],field=key)
            if not candidates:continue
            number,row=candidates[0]
            if item.get('value'):check('exact original structured observation reaches Excel',item['value'],row[3],photo_id=photo['photo_id'],field=key,cell='D'+str(number))
            check('status is a human label',True,isinstance(row[4],str) and bool(row[4]) and row[4] not in ('OBSERVED','UNCERTAIN','NOT_VISIBLE'),photo_id=photo['photo_id'],field=key)
            for name in ('evidence','reason'):
                if item.get(name):check('original '+name+' reaches same photo row',True,item[name] in str(row[5]),photo_id=photo['photo_id'],field=key,cell='F'+str(number))
            ledger.append(dict(photo_id=photo['photo_id'],field=key,row=number,cells=dict(zip('ABCDEF',row[:6]))))
    text='\n'.join(str(cell.value) for sheet in book for row in sheet for cell in row if cell.value is not None)
    check('no internal source/status prefix in workbook',False,bool(re.search(r'\b(?:OBSERVED|UNCERTAIN|NOT_VISIBLE|USER|POSE|SAVED|VISION)\b',text)))
    check('no response IDs or base64 data in workbook',False,'data:image/' in text or bool(re.search(r'\bresp_[A-Za-z0-9]+\b',text)))
    book.close();result=dict(checks=checks,ledger=ledger,defects=[row for row in checks if not row['pass_']]);(directory/'output-audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(dict(rows=len(ledger),checks=len(checks),defects=len(result['defects']))))
    if result['defects']:raise SystemExit(1)

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('directory',type=Path);args=parser.parse_args();run(args.directory.resolve())
