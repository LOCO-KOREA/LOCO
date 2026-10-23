"""Read-only physical DB / API / proposal / export / exact XLSX-cell ledger."""
import hashlib
import json
from pathlib import Path
import sqlite3
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from openpyxl import load_workbook
from ergo.user_language import user_text

def read(path):return json.loads(path.read_text(encoding='utf-8'))
def write(path,value):path.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')

def submitted_message(messages,text):
    """The intake boundary trims outer whitespace; all interior text must survive."""
    expected=text.strip()
    return next((message for message in messages if message['text']==expected),None)

def json_equal(left,right):
    # JSON numbers have one type; JavaScript serializes 120.0 as 120. Booleans are distinct.
    if type(left) in (int,float) and type(right) in (int,float):return left==right
    if type(left) is not type(right):return False
    if isinstance(left,dict):return left.keys()==right.keys() and all(json_equal(value,right[key]) for key,value in left.items())
    if isinstance(left,list):return len(left)==len(right) and all(json_equal(a,b) for a,b in zip(left,right))
    return left==right

def schema_errors(value,schema,path='$'):
    if 'anyOf' in schema:
        return [] if any(not schema_errors(value,option,path) for option in schema['anyOf']) else [path+': no anyOf alternative matches']
    types=schema.get('type');types=types if isinstance(types,list) else [types] if types else []
    matches={'null':value is None,'object':isinstance(value,dict),'array':isinstance(value,list),'string':isinstance(value,str),'number':type(value) in (int,float),'integer':type(value) is int,'boolean':type(value) is bool}
    if types and not any(matches[t] for t in types):return [path+': type mismatch']
    errors=[]
    if 'enum' in schema and value not in schema['enum']:errors.append(path+': enum mismatch')
    if isinstance(value,dict):
        properties=schema.get('properties',{})
        errors.extend(path+'.'+key+': required' for key in schema.get('required',[]) if key not in value)
        if schema.get('additionalProperties') is False:errors.extend(path+'.'+key+': extra' for key in value if key not in properties)
        for key,item in value.items():
            if key in properties:errors.extend(schema_errors(item,properties[key],path+'.'+key))
    elif isinstance(value,list):
        if 'minItems' in schema and len(value)<schema['minItems']:errors.append(path+': minItems')
        if 'maxItems' in schema and len(value)>schema['maxItems']:errors.append(path+': maxItems')
        for index,item in enumerate(value):errors.extend(schema_errors(item,schema.get('items',{}),path+'['+str(index)+']'))
    elif isinstance(value,str):
        if 'minLength' in schema and len(value)<schema['minLength']:errors.append(path+': minLength')
        if 'maxLength' in schema and len(value)>schema['maxLength']:errors.append(path+': maxLength')
    elif type(value) in (int,float):
        if 'minimum' in schema and value<schema['minimum']:errors.append(path+': minimum')
        if 'maximum' in schema and value>schema['maximum']:errors.append(path+': maximum')
    return errors

def capture(directory,stage):
    with sqlite3.connect('file:'+(directory/'loco.db').as_posix()+'?mode=ro',uri=True) as db:
        db.row_factory=sqlite3.Row;db.execute('BEGIN')
        tables=[r[0] for r in db.execute("select name from sqlite_master where type='table' and name not like 'sqlite_%'")]
        result=dict(schema={},tables={})
        for table in tables:
            result['schema'][table]=[dict(r) for r in db.execute('pragma table_info('+table+')')]
            result['tables'][table]=[]
            for row in db.execute('select * from '+table):
                item=dict(row)
                for key,value in item.items():
                    if key.endswith('_json') and value is not None:item[key]=json.loads(value)
                result['tables'][table].append(item)
    write(directory/(stage+'-db.json'),result)

def verify(directory):
    receipt=read(directory/'receipt.json');case=read(directory/'case.json');checks=[];defects=[];ledger=[]
    def check(name,expected,actual,**context):
        def compact(value):
            raw=json.dumps(value,ensure_ascii=False,sort_keys=True)
            return dict(sha256=hashlib.sha256(raw.encode()).hexdigest(),characters=len(raw)) if len(raw)>1500 else value
        same=json_equal(expected,actual)
        row=dict(name=name,expected=compact(expected),actual=compact(actual),pass_=same,**context);checks.append(row)
        if not row['pass_']:defects.append(row)
    check('journey reached final restart',True,bool(receipt.get('completed')))
    stages=[]
    for stage in receipt['stages']:
        api=read(directory/(stage+'.json'));db=read(directory/(stage+'-db.json'));stages.append((stage,api,db))
        check('six physical tables',sorted(['SURVEY','TASK','PHOTO','ANALYSIS','APP_SETTINGS','EXPORT_LOG']),sorted(db['tables']),stage=stage)
        physical=next(t for t in db['tables']['TASK'] if t['task_id']==receipt['task_id'])
        for key,value in physical['input_json'].items():check('API and physical task field',value,api['bundle']['task'].get(key),stage=stage,path=key)
        check('facts JSON survives actual serialization',physical['input_json']['conditions']['intake_session']['facts'],{f['identity']:{k:v for k,v in f.items() if k not in ('identity','label','unit')} for f in api['session']['facts']},stage=stage)
    if not stages:
        write(directory/'audit.json',dict(checks=checks,defects=defects,ledger=ledger));return
    for stage,api,_ in stages:
        expected=case.get('expected',{}) if stage=='01-first-ai' else {**case.get('expected',{}),**case.get('after_followup',{})} if stage=='02-followup-ai' else case.get('edits',{}) if stage in ('04-direct-edits','05-restarted','06-after-correction-ai','11-final-restart') else {}
        for target,wanted in expected.items():
            selected=[f for f in api['session']['facts'] if f['target']==target]
            actual=sorted(f['value'] for f in selected) if isinstance(wanted,list) else selected[0]['value'] if selected else 'MISSING'
            check('declared input expectation survives this checkpoint',sorted(wanted) if isinstance(wanted,list) else wanted,actual,stage=stage,target=target)
        title=next((f for f in api['session']['facts'] if f['target']=='task_name'),None)
        if case['id']=='06-decimal-load' and stage=='01-first-ai' and title and title['source']=='SAVED':
            check('automatic task name preserves the decimal quantity',True,title['value'].startswith('12.5'),stage=stage)
    actions=[json.loads(line) for line in (directory/'actions.jsonl').read_text(encoding='utf-8').splitlines()]
    for action in actions:
        if action['action']!='natural-language-input':continue
        stage=next((api for name,api,_ in stages if name==action['stage']),None)
        if stage:check('submitted natural reply is stored',True,submitted_message(stage['session']['messages'],action['text']) is not None,stage=action['stage'])
    # Further semantic and exact-cell checks are independent of application calculations.
    final_db=stages[-1][2];contract=stages[-1][1]['session']['contract'];catalogue={f['target']:f for f in contract['fields']}
    from ergo.intake.contract import response_schema
    frozen_contract=directory.parent/'baseline-intake-contract.json'
    expected_schema=read(frozen_contract)['response_schema'] if frozen_contract.exists() else response_schema()
    for stage,api,db in stages:
        if stage not in ('03-initial-export','07-final-analysis'):continue
        latest=api['bundle']['task']['latest_analysis_id'];stored=next(a for a in db['tables']['ANALYSIS'] if a['analysis_id']==latest)
        task=next(t for t in db['tables']['TASK'] if t['task_id']==receipt['task_id'])
        for key,value in task['input_json'].items():check('analysis uses the same typed TaskInput field as physical DB',value,stored['snapshot_json'].get(key),stage=stage,path=key)
    for export in final_db['tables']['EXPORT_LOG']:
        snap=export['snapshot_json'];analysis=snap['analyses'][0];source=analysis['snapshot'];label=Path(export['file_path']).stem
        check('task snapshot schema version','2.0',source.get('schema_version'),export=label)
        check('export schema version','2.0',snap.get('schema_version'),export=label)
        stored=next(a for a in final_db['tables']['ANALYSIS'] if a['analysis_id']==analysis['analysis_id'])
        check('export keeps selected immutable snapshot',stored['snapshot_json'],source,export=label)
        book=load_workbook(export['file_path'],data_only=False);cells={s.title:[dict(cell=c.coordinate,value=c.value,type=c.data_type) for row in s for c in row if c.value is not None] for s in book}
        write(directory/(label+'-cells.json'),cells)
        audit_source(directory,label,source,catalogue,cells,checks,defects,ledger)
        preview_file=directory/(label+'-preview.json')
        if preview_file.exists():audit_preview(label,read(preview_file)['sheets'],book,checks,defects)
        for ai in analysis.get('ai',[]):
            if ai.get('is_active'):check('applied explanation is exact XLSX text',True,any(c['value']==user_text(ai['applied_text']) for c in cells.get('쉬운 설명',[])),export=label,text_id=ai['text_id'])
        book.close()
    trace_file=directory/'openai.jsonl';trace=[json.loads(line) for line in trace_file.read_text(encoding='utf-8').splitlines()] if trace_file.exists() else []
    check('actual SDK requests recorded',True,any(t['kind']=='REQUEST' for t in trace))
    ipc_file=directory/'ipc.jsonl'
    if ipc_file.exists():
        packets=[json.loads(line) for line in ipc_file.read_text(encoding='utf-8').splitlines()]
        replies={(e['pid'],e['packet']['request_id']):e['packet']['result'] for e in packets if e['kind']=='RESPONSE'}
        for event in packets:
            packet=event['packet'];params=packet.get('params',{})
            if event['kind']!='REQUEST' or packet.get('method')!='intake_turn' or not params.get('text'):continue
            reply=replies.get((event['pid'],packet['request_id']))
            check('actual submitted IPC turn has a response',True,reply is not None,request_id=packet['request_id'])
            if not reply or not reply.get('ok'):continue
            state=reply['value'];message=submitted_message(state['messages'],params['text'])
            check('actual IPC turn response contains submitted text',True,message is not None,request_id=packet['request_id'])
            if message:
                turn=next((t for t in state['history'] if t.get('message_id')==message['message_id']),None)
                check('submitted IPC text has its own actual extraction',True,bool(turn and any(r['phase']=='EXTRACT_AND_PLAN' and r['connection'].get('response_id') for r in turn['requests'])),request_id=packet['request_id'])
    request_ids={t['trace_id']:t['request'] for t in trace if t['kind']=='REQUEST'}
    final_session=stages[-1][1]['session'];request_receipts={r['connection'].get('response_id'):r for turn in final_session['history'] for r in turn['requests']}
    for event in trace:
        if event['kind']!='RESPONSE':continue
        response=event['response'];request=request_ids[event['trace_id']];check('actual response model','gpt-6-luna',response.get('model'),response_id=response.get('id'))
        if request.get('text',{}).get('format',{}).get('name')!='loco_adaptive_intake':continue
        check('actual SDK schema equals the canonical intake schema',expected_schema,request['text']['format']['schema'],response_id=response.get('id'))
        text=''.join(c.get('text','') for output in response.get('output',[]) if output.get('type')=='message' for c in output.get('content',[]) if c.get('type')=='output_text')
        if not text:continue
        raw=json.loads(text);saved=request_receipts.get(response['id']);check('raw AI proposal receipt exists',True,saved is not None,response_id=response['id'])
        if saved is not None:check('raw AI answers are preserved',raw.get('answers'),saved['answers'],response_id=response['id'])
        check('raw AI response matches the requested strict schema',[],schema_errors(raw,request['text']['format']['schema']),response_id=response['id'])
        payload=json.loads(request['input'][1]['content']);check('actual requested reasoning effort','xhigh',request.get('reasoning',{}).get('effort'),response_id=response['id'])
        check('no image bytes in AI request',False,'data:image' in json.dumps(request),response_id=response['id'])
        if payload['phase']=='FOLLOWUP':
            stage=next(((name,api) for name,api,_ in stages if any(any(r['connection'].get('response_id')==response['id'] for r in t['requests']) for t in api['session']['history'])),None)
            if stage:
                name,api=stage
                for proposal in raw['answers']:
                    field=catalogue[proposal['target']];value=proposal['number'] if field['kind']=='number' else proposal['text']
                    facts=[f for f in api['session']['facts'] if f['target']==proposal['target']]
                    if value is not None and value!='UNKNOWN' and field['ai_writable'] and not facts:
                        decisions=[d for turn in api['session']['history'] for index,r in enumerate(turn['requests']) if r['connection'].get('response_id')==response['id'] for d in turn['decisions'] if d.get('request_index')==index and d['target']==proposal['target']]
                        check('novel FOLLOWUP proposal receives validation instead of phase discard',True,any(d['outcome']!='IGNORED' for d in decisions),stage=name,response_id=response['id'],target=proposal['target'],proposal=proposal)
        if payload['phase'] in ('EXTRACT_AND_PLAN','FOLLOWUP'):
            stage=next(((name,api) for name,api,_ in stages if any(any(r['connection'].get('response_id')==response['id'] for r in t['requests']) for t in api['session']['history'])),None)
            if stage:
                name,api=stage
                for turn in api['session']['history']:
                    for index,request_receipt in enumerate(turn['requests']):
                        if request_receipt['connection'].get('response_id')!=response['id']:continue
                        for proposal_index,proposal in enumerate(raw['answers']):
                            decision=next((d for d in turn['decisions'] if d.get('request_index')==index and d.get('proposal_index')==proposal_index),None)
                            check('each raw proposal has a recorded decision',True,decision is not None,response_id=response['id'],proposal_index=proposal_index)
                            if not decision or decision['outcome'] not in ('ACCEPTED','UNCHANGED'):continue
                            identity=proposal['target']+'@'+decision['condition'];fact=next((f for f in api['session']['facts'] if f['identity']==identity),None)
                            check('accepted proposal has a stored fact',True,fact is not None,response_id=response['id'],identity=identity)
                            if fact and decision['outcome']=='ACCEPTED':
                                field=catalogue[proposal['target']];value=proposal['number'] if field['kind']=='number' else proposal['text']
                                if value=='UNKNOWN':value=None
                                if field['kind']=='choice' and value is not None:value=next((c['value'] for c in field['choices'] if str(c['value'])==str(value) or c['label']==value),value)
                                check('accepted proposal value equals serialized fact',value,fact['value'],response_id=response['id'],identity=identity)
    # The declared natural-language oracle is independent of whatever AI chose to extract.
    for expected in case.get('rows',[]):
        for stage,api,_ in stages:
            if stage not in ('01-first-ai','02-followup-ai'):continue
            rows=api['bundle']['task']['legal']['criteria'][expected['criterion']-1]['rows']
            wanted={k:v for k,v in expected.items() if k!='criterion'}
            check('same-condition weight/count pair from initial statement',True,any(all(row.get(k)==v for k,v in wanted.items()) for row in rows),stage=stage,pair=wanted)
    write(directory/'audit.json',dict(id=case['id'],checks=checks,defects=defects,ledger=ledger,counts=dict(checks=len(checks),defects=len(defects),fields=len(ledger)),schema_note='AI proposals and TaskInput have distinct shapes; verified catalogue mapping and selected snapshot identity.'))

def audit_source(directory,label,source,catalogue,cells,checks,defects,ledger):
    # Independent paths and expected display values; do not call facts.project/read_values or report formatters.
    from ergo.report.labels import CRITERIA,EXPOSURES,INPUTS,MANUAL,RESULT_LABELS
    words={'YES':'예','NO':'아니오','UNKNOWN':'모름','GOOD':'좋음','FAIR':'보통','POOR':'나쁨','UNACCEPTABLE':'잡을 곳 없음','NOT_APPLICABLE':'해당 없음','BILATERAL_STABLE':'두 발 안정','UNILATERAL_OR_UNSTABLE':'한 발·불안정','UNDER_30':'거의 안 굽힘(30° 미만)','FROM_30_TO_60':'30~60°','OVER_60':'60° 초과'}
    choices={'trunk_base':{1:'똑바로',2:'조금(20° 이내)',3:'많이(20~60°)',4:'아주 많이(60° 초과)'},'neck_base':{1:'똑바로·조금 숙임(20° 이내)',2:'20° 넘게 숙임','2_EXTENSION':'뒤로 젖힘'},'upper_arm_base':{1:'몸 옆(20° 이내)',2:'20~45°',3:'45~90°',4:'90° 초과'},'lower_arm_base':{1:'60~100°',2:'그 밖(60° 미만·100° 초과)'},'wrist_base':{1:'15° 이내',2:'15° 넘게'}}
    names={'daily_work_min':'하루 실제 작업시간','annual_work_days':'1년에 작업하는 날 수','one_off':'한 번만 하고 끝나는 일회성 작업인가','planned_duration_months':'예정 작업기간'}
    units={'daily_work_min':'분/일','annual_work_days':'일/년','planned_duration_months':'개월','load_mass_kg':'kg','measured_load_kg':'kg','pinch_mass_kg':'kg','single_hand_mass_kg':'kg','pinch_force_kg_equiv':'kg 상당','grip_force_kg_equiv':'kg 상당','minutes':'분/일','lift_rate_min':'회/분','impact_rate_hour':'회/시간'}
    lift={8:{'load_mass_kg','lift_count_day'},9:{'load_mass_kg','below_knee_count_day','above_shoulder_count_day','extended_arm_count_day','union_count_day'},10:{'load_mass_kg','lift_rate_min','minutes'}}
    criteria={c['criterion_no']:c for c in source['legal']['criteria']};slots=source['capture']['slots'];primary=next((s for s in slots if s['analysis_target']),None)
    sheet=cells.get('입력한 값',[]);rowmap={}
    for cell in sheet:
        number=''.join(c for c in cell['cell'] if c.isdigit());column=''.join(c for c in cell['cell'] if c.isalpha());rowmap.setdefault(number,{})[column]=cell
    rows=[{k:v['value'] for k,v in columns.items()}|{'cell':columns.get('D',columns.get('C'))['cell']} for columns in rowmap.values() if 'C' in columns]
    def format_value(value,key='',unit=None):
        if value is None:return ''
        if key in choices and value in choices[key]:return choices[key][value]
        if isinstance(value,(int,float)) and not isinstance(value,bool):text=str(int(value)) if float(value).is_integer() else str(value)
        else:text=words.get(value,str(value))
        unit=unit if unit is not None else ('회/일' if key.endswith('count_day') else units.get(key,''))
        return text+(' '+unit if unit else '')
    def finding(name,expected,actual,identity,**extra):
        item=dict(name=name,export=label,identity=identity,expected=expected,actual=actual,pass_=expected==actual,**extra);checks.append(item)
        if not item['pass_']:defects.append(item)
    def lineage(fact):
        session=source['conditions']['intake_session'];messages={m.get('message_id'):m for m in session.get('messages',[])};proposals=[];direct=[]
        for ti,turn in enumerate(session.get('turns',[])):
            message=messages.get(turn.get('message_id'),{})
            for ai,answer in enumerate(message.get('answers',[])):
                if answer['target']==fact['target'] and (answer.get('condition') or 'default')==fact['condition']:
                    direct.append(dict(turn_id=turn['turn_id'],message_id=turn.get('message_id'),value=answer.get('value'),path=f'intake_session.messages[message_id={turn.get("message_id")}].answers[{ai}]'))
            for ri,request in enumerate(turn.get('requests',[])):
                for pi,proposal in enumerate(request.get('answers',[])):
                    if proposal['target']!=fact['target']:continue
                    decision=next((d for d in turn['decisions'] if d.get('request_index')==ri and d.get('proposal_index')==pi),None)
                    condition=(decision or proposal).get('condition') or 'default'
                    if condition!=fact['condition']:continue
                    proposals.append(dict(turn_id=turn['turn_id'],phase=request['phase'],response_id=request['connection'].get('response_id'),
                        raw=proposal,decision=decision,path=f'intake_session.turns[{ti}].requests[{ri}].answers[{pi}]'))
        texts=[dict(message_id=m.get('message_id'),text=m['text']) for m in session.get('messages',[]) if fact.get('quote') and fact['quote'] in m.get('text','')]
        return dict(proposals=proposals,direct_answers=direct,source_texts=texts)
    for identity,fact in source['conditions']['intake_session']['facts'].items():
        target=fact['target'];field=catalogue[target];d=field['destination'];kind=d['type'];value=fact['value'];projection=[];display=[];display_class='business-value';coverage='checked'
        add=lambda path,expected,actual:projection.append(dict(path=path,expected=expected,actual=actual))
        def display_row(matches,wanted,representation='value'):
            display.append(dict(expected=wanted,representation=representation,cells=[dict(cell=r['cell'],value=r.get('D'),source=r.get('E'),quote=r.get('F')) for r in matches]))
            finding('business value reaches an identified Excel row',True,bool(matches),identity,expected_display=wanted)
            if matches and value is not None and representation=='value':finding('exact Excel value and unit',True,any(str(r.get('D'))==wanted for r in matches),identity,cells=display[-1]['cells'])
        if kind=='conditions':
            key=d['field'];expected='UNKNOWN' if key=='one_off' and value is None else value;add('conditions.'+key,expected,source['conditions'].get(key));matches=[r for r in rows if r.get('C')==names.get(key)]
            if value is not None:display_row(matches,format_value(value,key))
            else:
                display_class='explicit-null';finding('null is not rendered as numeric zero',False,any(str(r.get('D','')).startswith('0 ') for r in matches),identity)
                display_row(matches,'모름')
                if matches:finding('explicit null is displayed as unknown',True,any(r.get('D')=='모름' for r in matches),identity)
            if matches and value is not None:finding('output input has a source',True,any(bool(r.get('E')) for r in matches),identity,cells=[r['cell'] for r in matches])
        elif kind=='task':
            add(d['field'],value,source.get(d['field']));finding('task name reaches workbook',True,any(c['value']==value for c in sheet),identity)
        elif kind in ('work_type','complete'):
            key='condition_present' if kind=='work_type' else 'rows_complete'
            for no in d['criteria']:add(f'legal.criteria[{no}].{key}',value or 'UNKNOWN',criteria[no].get(key))
            if kind=='work_type':
                for no in d['criteria']:
                    if value=='YES':matches=[r for r in rows if r.get('C')==CRITERIA[no]+': 이 작업을 한다'];display_row(matches,'예')
                    else:
                        title='해당 없음으로 답한 기준' if value=='NO' else '아직 답하지 않은 기준';matches=[r for r in rows if r.get('C')==CRITERIA[no]+': 이 작업을 한다']
                        if matches:display_row(matches,words.get(value,'모름'))
                        else:matches=[r for r in rows if r.get('C')==title and CRITERIA[no] in str(r.get('D'))];display_row(matches,words.get(value,'모름'),'grouped criterion')
            else:
                # Completeness is a calculation input; explicitly record whether it has a visible representation.
                display_class='completeness'
                for no in d['criteria']:
                    matches=[r for r in rows if r.get('C')==CRITERIA[no]+': 하루의 모든 해당 조건 행 기록 여부']
                    if value is not None:display_row(matches,format_value(value))
        elif kind=='criterion':
            mapping=d['mapping']
            if 'exposure' in mapping:
                code=mapping['exposure'];actual=next((e.get('minutes') for e in source['legal']['exposures'] if e['condition_code']==code),'MISSING');add('legal.exposures['+code+'].minutes',value,actual)
                matches=[r for r in rows if EXPOSURES[code]+' 시간' in str(r.get('C'))]
                if value is not None:display_row(matches,format_value(value,'minutes'))
            elif 'criterion' in mapping:
                key=mapping['criterion'];no=d['criteria'][0];add(f'legal.criteria[{no}].{key}',value or 'UNKNOWN',criteria[no].get(key));matches=[r for r in rows if r.get('C')==CRITERIA[no]+': '+INPUTS.get(key,key)]
                if value is not None:display_row(matches,format_value(value,key))
            else:
                key=mapping.get('lift') or mapping.get('row')
                for no in d['criteria']:
                    if 'lift' in mapping and key not in lift[no]:continue
                    listed=criteria[no].get('rows',[]);position=next((i for i,r in enumerate(listed) if r['row_id']==fact['condition']),None);row=listed[position] if position is not None else {}
                    add(f'legal.criteria[{no}].rows[{fact["condition"]}].{key}',value,row.get(key,'MISSING'))
                    suffix=' (조건 '+str(position+1)+')' if len(listed)>1 and position is not None else ''
                    matches=[r for r in rows if r.get('C')==CRITERIA[no]+': '+INPUTS.get(key,RESULT_LABELS.get(key,key))+suffix]
                    if value is not None:display_row(matches,format_value(value,key))
        elif kind=='manual':
            entry=next((m for m in source['manual'] if primary and m['photo_id']==primary['photo_id'] and m['scope']==d['scope'] and m['field']==d['field']),None)
            applicable=primary and value is not None and fact['state']!='ESTIMATED' and (fact['source'] not in ('POSE','MOVENET') or primary['observation']['normal_work']=='YES')
            add(f'manual[{d["scope"]}.{d["field"]}]',value if applicable else None,entry['value'] if entry else None)
            if applicable:
                side={'LEFT':'왼쪽','RIGHT':'오른쪽','COMMON':''}[d['scope']];text=MANUAL.get(d['field'],d['field']);matches=[r for r in rows if text in str(r.get('C')) and (not side or side in r['C'] or '양쪽' in r['C'])]
                if value=='NO':matches=[r for r in rows if '아니오로 기록' in str(r.get('C')) and text in str(r.get('D')) and (not side or side in str(r.get('D')) or '양쪽' in str(r.get('D')))];display_row(matches,'아니오','grouped NO')
                else:display_row(matches,format_value(value,d['field']))
            else:display_class='pose candidate or null';coverage='snapshot only; not an accepted manual input'
        elif kind in ('context','load','load_kg','view'):
            observation=(primary or {}).get('observation',{})
            if kind=='context':
                for key in ('single_person','full_frame','normal_work'):add('capture.primary.observation.'+key,value or 'UNKNOWN',observation.get(key))
                matches=[r for r in rows if str(r.get('C')).endswith('사진 확인')];display_row(matches,'사진 확인','composite photo context')
            elif kind=='load':
                add('capture.primary.observation.load_state',value or 'UNKNOWN',observation.get('load_state'));matches=[r for r in rows if str(r.get('C')).endswith('들고 있는 물건')];display_row(matches,'없음' if value=='UNLOADED' else '있음','load description')
            elif kind=='load_kg':
                if observation.get('load_state')=='MEASURED':add('capture.primary.observation.measured_load_kg',value,observation.get('measured_load_kg'));matches=[r for r in rows if str(r.get('C')).endswith('들고 있는 물건')];finding('photo load and kg reach exact photo row',True,any(format_value(value,'measured_load_kg') in str(r.get('D')) for r in matches),identity)
                else:display_class='unlinked load candidate';coverage='snapshot only'
            else:
                slot=next((s for s in slots if s.get('label')==d['view']),None)
                if slot:add('capture['+d['view']+'].forward_direction',value,slot.get('forward_direction'))
                display_class='photo direction'
        else:display_class='auxiliary '+kind;coverage='fact retained inside export snapshot; no direct worksheet field'
        for item in projection:finding('fact value matches canonical calculation input',item['expected'],item['actual'],identity,path=item['path'])
        if fact['source'] in ('AI','USER') and value is not None and display_class=='business-value' and kind in ('conditions','criterion','manual'):
            represented=[cell for row in display for cell in row['cells']]
            def own_quote(cell):
                quote=cell.get('quote') or ''
                for prefix in ('AI: ','USER: ','SAVED: ','POSE: ','MOVENET: ','직접 입력: ','저장된 입력: ','사진 자동 측정: ','AI 사진 관찰: '):
                    if quote.startswith(prefix):quote=quote[len(prefix):];break
                return quote
            if represented:finding('each output value keeps its own source quote',True,any(own_quote(cell)==fact.get('quote') for cell in represented),identity,expected_quote=fact.get('quote'),cells=represented)
        ledger.append(dict(export=label,identity=identity,target=target,condition=fact['condition'],value=value,source=fact['source'],state=fact['state'],quote=fact.get('quote'),unit=field['unit'],destination=d,projection=projection,display=display,display_class=display_class,coverage=coverage,lineage=lineage(fact)))

def audit_preview(label,sheets,book,checks,defects):
    for sheet in sheets:
        ws=book[sheet['name']];cursor=1+bool(sheet.get('lead'))+1+1
        for table in sheet['tables']:
            cursor+=bool(table.get('heading'))+bool(table.get('header',True))
            covered={(r,c) for r1,c1,r2,c2 in table.get('merges',[]) for r in range(r1,r2+1) for c in range(c1,c2+1) if (r,c)!=(r1,c1)}
            for r,row in enumerate(table['rows']):
                cursor+=1
                for c,value in enumerate(row):
                    if (r,c) in covered:continue
                    actual=ws.cell(cursor,c+1).value;expected=None if value=='' else value
                    item=dict(name='preview table cell equals physical XLSX cell',export=label,sheet=sheet['name'],cell=ws.cell(cursor,c+1).coordinate,expected=expected,actual=actual,pass_=expected==actual);checks.append(item)
                    if not item['pass_']:defects.append(item)
            cursor+=1

if __name__=='__main__':
    mode=sys.argv[1];directory=Path(sys.argv[2]).resolve()
    if mode=='capture':capture(directory,sys.argv[3])
    elif mode=='verify':verify(directory)
