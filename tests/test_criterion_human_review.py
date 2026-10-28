"""Typed non-binding suggestions become conclusions only after explicit human review."""
import copy
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import load_workbook
from backend.contracts import validate_request
from backend.gpt import client,payload
from backend.service.application import create_service
from test_v2_gpt import content_fixture


class CriterionHumanReview(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.ctx=create_service(str(Path(self.temp.name)/'review.db'))
        self.survey=self.ctx.dispatch('create_survey',dict(draft=dict(title='Human review')))
        self.task=self.ctx.dispatch('create_task',dict(survey_id=self.survey['survey_id'],draft=dict(task_name='상자')))
        self.analysis=self.ctx.dispatch('recalculate_judgments',dict(task_id=self.task['task_id']))
        self.identity=self.analysis['analysis_id'];self.snapshot=copy.deepcopy(self.analysis['snapshot'])

    def generated(self,status='UNDETERMINED',overrides=None):
        analysis=self.ctx.dispatch('get_analysis',dict(analysis_id=self.identity))
        summary=payload.build_numeric_summary(analysis)
        content=content_fixture(summary)
        for item in content['criteria']:
            item['proposed_status']=(overrides or {}).get(item['criterion_no'],status)
        response=SimpleNamespace(status='completed',id='resp_fixture',model='fixture',reasoning=dict(effort='high'),
                                 output_text=json.dumps(content,ensure_ascii=False))
        settings=self.ctx.settings();settings.update(gpt_enabled=True,gpt_model='fixture')
        self.ctx.dispatch('update_settings',dict(settings=settings))
        with patch.dict('os.environ',OPENAI_API_KEY='fixture'),patch('backend.gpt.client.create_response',return_value=response):
            result=self.ctx.dispatch('generate_ai_text',dict(analysis_id=self.identity))
        self.assertEqual(result['text']['result'],'GENERATED',result['text'].get('fallback_reason'))
        return result['text']

    def confirm(self,choices,source=None,revision=None):
        return self.ctx.dispatch('confirm_criterion_review',dict(analysis_id=self.identity,
            expected_input_revision=self.analysis['input_revision'] if revision is None else revision,
            conclusions=[dict(criterion_no=number,status=status) for number,status in choices],source_text_id=source,reviewer='검토자'))

    def preview(self):
        return self.ctx.dispatch('preview_excel_sheets',dict(survey_id=self.survey['survey_id'],analysis_ids=[self.identity],comparisons=[]))

    def test_strict_new_schema_and_legacy_content_read_compatibility(self):
        summary=payload.build_numeric_summary(self.analysis);new=content_fixture(summary)
        self.assertTrue(payload.validate_content(new,summary,allow_legacy=False))
        schema=payload.explanation_schema()['properties']['criteria']['items']
        self.assertIn('proposed_status',schema['required'])
        self.assertEqual(schema['properties']['proposed_status']['enum'],['APPLICABLE','NOT_APPLICABLE','UNDETERMINED'])
        old=copy.deepcopy(new)
        for item in old['criteria']:item.pop('proposed_status')
        self.assertTrue(payload.validate_content(old,summary))
        self.assertFalse(payload.validate_content(old,summary,allow_legacy=False))
        bad=copy.deepcopy(new);bad['criteria'][0]['proposed_status']='NO'
        self.assertFalse(payload.validate_content(bad,summary))
        bad=copy.deepcopy(new);bad['criteria'][0]['proposed_status']='NOT_APPLICABLE'
        self.assertFalse(payload.validate_content(bad,summary,allow_legacy=False),'unknown-only evidence cannot invent NO')

    def test_current_typed_ai_is_draft_only_and_has_no_exported_conclusions(self):
        ai=self.generated();current=self.ctx.dispatch('get_analysis',dict(analysis_id=self.identity))
        self.assertEqual(len(current['criteria_proposals']),11)
        self.assertTrue(all(item['proposed_status']=='UNDETERMINED' for item in current['criteria_proposals']))
        self.assertFalse(current['confirmation_current'])
        self.assertTrue(all(item['investigator']['status'] is None for item in current['judgment']['judgments']))
        self.assertFalse(any(sheet['name']=='부담작업 해당 여부' for sheet in self.preview()['sheets']))
        self.assertEqual(current['snapshot'],self.snapshot)
        self.assertEqual(current['ai'][0]['input_summary'],ai['input_summary'])

    def test_ai_free_review_hold_and_explicit_clear_are_distinct(self):
        chosen=self.confirm([(1,'APPLICABLE')]);old=copy.deepcopy(chosen['judgment']['judgments'][0])
        held=self.confirm([(1,'UNDETERMINED'),(2,'UNDETERMINED')])
        self.assertEqual(held['judgment']['judgments'][0],old)
        self.assertTrue(held['confirmation_current'])
        rows=next(sheet for sheet in self.preview()['sheets'] if sheet['name']=='부담작업 해당 여부')['tables'][0]['rows']
        self.assertEqual([row[2] for row in rows],['판단 보류','판단 보류'])
        cleared=self.confirm([(1,None)])
        self.assertIsNone(cleared['judgment']['judgments'][0]['investigator']['status'])
        self.assertEqual(cleared['judgment']['judgments'][0]['effective_status'],old['computed_status'])
        self.assertEqual(cleared['snapshot'],self.snapshot)
        self.assertEqual(len(cleared['criterion_review_history']),2)

    def test_linked_confirmation_preserves_ai_original_and_supports_explicit_text_apply_and_excel(self):
        ai=self.generated();before=copy.deepcopy(ai)
        confirmed=self.confirm([(1,'NOT_APPLICABLE'),(8,'NOT_APPLICABLE')],ai['text_id'])
        saved=confirmed['ai'][0]
        self.assertFalse(saved['source_current'])
        self.assertTrue(saved['confirmation_current'])
        self.assertEqual(saved['input_summary'],before['input_summary']);self.assertEqual(saved['content'],before['content'])
        self.assertEqual(saved['generated_text'],before['generated_text'])
        self.assertEqual(len(confirmed['criteria_proposals']),11)
        self.assertTrue(all(not item['source_current'] and item['confirmation_current'] for item in confirmed['criteria_proposals']))
        text=ai['generated_text']+'\n검토 수정 문장'
        self.ctx.dispatch('edit_ai_text',dict(analysis_id=self.identity,text_id=ai['text_id'],text=text))
        self.ctx.dispatch('apply_ai_text',dict(analysis_id=self.identity,text_id=ai['text_id'],use_edited=True))
        final=self.ctx.dispatch('get_analysis',dict(analysis_id=self.identity))
        self.assertTrue(final['confirmation_current']);self.assertTrue(final['ai'][0]['confirmation_current'])
        self.assertEqual(final['ai'][0]['applied_text'],text)
        destination=Path(self.temp.name)/'review.xlsx'
        self.ctx.dispatch('export_excel',dict(survey_id=self.survey['survey_id'],spec=dict(analysis_ids=[self.identity],include_representative=False,comparisons=[],destination=str(destination))))
        book=load_workbook(destination,read_only=True)
        self.assertIn('결과 설명',book.sheetnames)
        cells='\n'.join(str(cell) for row in book['결과 설명'].iter_rows(values_only=True) for cell in row if cell is not None)
        self.assertIn('검토 수정 문장',cells);book.close()
        self.assertEqual(final['snapshot'],self.snapshot)

    def test_partial_confirmation_carries_prior_explicit_human_conclusions(self):
        self.ctx.dispatch('set_investigator_conclusion',dict(analysis_id=self.identity,conclusion=dict(criterion_no=2,status='APPLICABLE',reason='직접 확인',evidence={},reviewer=None,reviewed_at=None)))
        confirmed=self.confirm([(8,'NOT_APPLICABLE')])
        self.assertEqual({item['criterion_no']:item['status'] for item in confirmed['criterion_review']['conclusions']},{2:'APPLICABLE',8:'NOT_APPLICABLE'})
        rows=next(sheet for sheet in self.preview()['sheets'] if sheet['name']=='부담작업 해당 여부')['tables'][0]['rows']
        self.assertEqual(len(rows),2)

    def test_changed_conclusion_score_inputs_and_current_analysis_invalidate_receipts(self):
        self.confirm([(1,'NOT_APPLICABLE')])
        self.ctx.dispatch('set_investigator_conclusion',dict(analysis_id=self.identity,conclusion=dict(criterion_no=1,status='APPLICABLE',reason='수정',evidence={},reviewer=None,reviewed_at=None)))
        stale=self.ctx.dispatch('get_analysis',dict(analysis_id=self.identity))
        self.assertFalse(stale['confirmation_current'])
        rows=next(sheet for sheet in self.preview()['sheets'] if sheet['name']=='부담작업 해당 여부')['tables'][0]['rows']
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0][0],'상자');self.assertIn('1호',rows[0][1]);self.assertEqual(rows[0][2],'부담작업',
                         'the stale batch receipt is ignored, while the newer direct investigator conclusion remains exportable')
        self.confirm([(1,'NOT_APPLICABLE')])
        self.ctx.edit_result(self.identity,lambda result:result['module_states'].update(judgment='FAILED'))
        self.assertFalse(self.ctx.dispatch('get_analysis',dict(analysis_id=self.identity))['confirmation_current'])
        self.ctx.dispatch('set_task_conditions',dict(task_id=self.task['task_id'],conditions=dict(self.task['conditions'],daily_work_min=61)))
        self.assertFalse(self.ctx.dispatch('get_analysis',dict(analysis_id=self.identity))['confirmation_current'])
        with self.assertRaises(ValueError):self.confirm([(1,'NOT_APPLICABLE')])

    def test_bad_batch_or_stale_token_never_partially_writes(self):
        before=self.ctx.require('analysis',self.identity)
        for choices in [[],[dict(criterion_no=1,status='APPLICABLE')]*2,[dict(criterion_no=12,status='APPLICABLE')],
                        [dict(criterion_no=1,status='UNKNOWN')],[dict(criterion_no=True,status=None)]]:
            with self.subTest(choices=choices),self.assertRaises(ValueError):
                self.ctx.dispatch('confirm_criterion_review',dict(analysis_id=self.identity,expected_input_revision=self.analysis['input_revision'],conclusions=choices))
        with self.assertRaises(ValueError):self.confirm([(1,'APPLICABLE')],revision=999)
        with self.assertRaises(ValueError):self.confirm([(1,'APPLICABLE')],source='foreign-text')
        self.assertEqual(self.ctx.require('analysis',self.identity),before)

    def test_legacy_or_stale_ai_never_populates_fake_proposals(self):
        ai=self.generated()
        def legacy(result):
            for item in result['ai'][0]['content']['criteria']:item.pop('proposed_status')
        self.ctx.edit_result(self.identity,legacy)
        self.assertEqual(self.ctx.dispatch('get_analysis',dict(analysis_id=self.identity))['criteria_proposals'],[])
        with self.assertRaises(ValueError):self.confirm([(1,'NOT_APPLICABLE')],ai['text_id'])

    def test_ai_conflict_is_a_warning_and_never_changes_computation_or_facts(self):
        self.ctx.dispatch('intake_turn',dict(task_id=self.task['task_id'],plan=False,answers=[
            dict(target='work_type.lift',value='YES'),dict(target='complete.lift',value='YES'),
            dict(target='card.lift.load_mass_kg',value=20),dict(target='card.lift.lift_count_day',value=120)]))
        self.analysis=self.ctx.dispatch('recalculate_judgments',dict(task_id=self.task['task_id']));self.identity=self.analysis['analysis_id']
        before=copy.deepcopy(self.analysis)
        self.generated(overrides={8:'APPLICABLE'})
        after=self.ctx.dispatch('get_analysis',dict(analysis_id=self.identity))
        suggestion=next(item for item in after['criteria_proposals'] if item['criterion_no']==8)
        self.assertEqual(suggestion['proposed_status'],'APPLICABLE');self.assertTrue(suggestion['conflicts_with_computed'])
        self.assertEqual(after['judgment'],before['judgment']);self.assertEqual(after['snapshot'],before['snapshot'])

    def test_new_analysis_without_input_change_still_invalidates_previous_confirmation(self):
        self.confirm([(1,'NOT_APPLICABLE')])
        latest=self.ctx.dispatch('recalculate_judgments',dict(task_id=self.task['task_id']))
        self.assertNotEqual(latest['analysis_id'],self.identity)
        self.assertFalse(self.ctx.dispatch('get_analysis',dict(analysis_id=self.identity))['confirmation_current'])
        with self.assertRaises(ValueError):self.confirm([(1,'NOT_APPLICABLE')])


if __name__=='__main__':unittest.main()
