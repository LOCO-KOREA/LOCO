"""Explicit live AI -> isolated service/SQLite -> analysis snapshot -> actual XLSX.

Synthetic statements only. No Electron, pose inference, UI or OS input is used.
"""
import argparse
from concurrent.futures import ThreadPoolExecutor, as_completed
from contextlib import closing
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import re
import sqlite3
import subprocess
import sys
import threading
import time

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from backend import ai_config
from backend.gpt import client as ai_client
from backend.service.application import create_service
from openpyxl import load_workbook

CASES=(
    ('and', '하루 작업은 손·팔을 높이 드는 작업 2시간 및 키보드 작업 6시간이 전부예요.',480),
    ('list', '하루 작업: 손·팔을 높이 드는 작업 2시간. 키보드 작업 6시간. 이 두 작업이 전부예요.',480),
    ('label_and_annotations', '하루 총 작업은 손·팔 들기 2시간(120분), 키보드 작업 6시간(360분)이 전부예요.',480),
    ('compound', '하루 전체 실제 작업시간은 한 시간 삼십분이에요.',90),
)
TRACE=threading.local()


def verify(case,folder,defaults):
    name,text,expected=case;directory=folder/name;directory.mkdir(parents=True)
    TRACE.directory=directory;TRACE.requests=0
    db=directory/'loco.db';ctx=create_service(db)
    survey=ctx.dispatch('create_survey',dict(draft=dict(title='합성 시간 표현 검증')))
    task=ctx.dispatch('create_task',dict(survey_id=survey['survey_id'],draft=dict(task_name='시간 표현 검증',
        analysis_scope=dict(judgment=True,pose=False,tags=False,reba=False))))
    settings=ctx.dispatch('get_settings',{})
    for key in ('gpt_contract','pose_contract','gpt_key_present'):settings.pop(key,None)
    settings.update(gpt_enabled=True,gpt_model=os.environ.get('LOCO_LIVE_MODEL',defaults['gpt_model']),
        gpt_reasoning_effort=os.environ.get('LOCO_LIVE_REASONING_EFFORT',defaults['gpt_reasoning_effort']))
    ctx.dispatch('update_settings',dict(settings=settings))
    started=time.monotonic()
    state=ctx.dispatch('intake_turn',dict(task_id=task['task_id'],text=text))
    assert state['ai']['status']=='CONNECTED',state['ai']
    assert state['ai']['model']==settings['gpt_model'],state['ai']
    assert state['ai']['reasoning_effort']==settings['gpt_reasoning_effort'],state['ai']
    assert state['ai'].get('response_id'),state['ai']
    assert state['task']['conditions']['daily_work_min']==expected,dict(case=name,conditions=state['task']['conditions'],dropped=state['dropped'])
    assert 'daily_work_min' not in {question['target'] for question in state['questions']}
    saved=next(fact for fact in state['facts'] if fact['target']=='daily_work_min')
    assert saved['source']=='AI',saved
    if name=='label_and_annotations':assert saved['calculation']['operation']=='SUM_STATED_DURATIONS',saved
    restored=create_service(db).dispatch('intake_session',dict(task_id=task['task_id']))
    assert restored['task']['conditions']['daily_work_min']==expected
    assert next(fact for fact in restored['facts'] if fact['target']=='daily_work_min')==saved
    request=ctx.dispatch('start_analysis',dict(task_id=task['task_id']))
    deadline=time.monotonic()+60
    while True:
        analysis=ctx.dispatch('get_analysis',dict(analysis_id=request['analysis_id']))
        if analysis['status']!='RUNNING':break
        assert time.monotonic()<deadline,'analysis did not settle'
        time.sleep(.05)
    assert analysis['status']=='FINISHED',analysis['status']
    assert analysis['snapshot']['conditions']['daily_work_min']==expected
    workbook=directory/'report.xlsx'
    receipt=ctx.dispatch('export_excel',dict(survey_id=survey['survey_id'],spec=dict(
        destination=str(workbook),analysis_ids=[analysis['analysis_id']],include_representative=False,comparisons=[])))
    with closing(sqlite3.connect(db)) as connection:
        tables=[row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")]
        assert len(tables)==6,tables
        frozen=json.loads(connection.execute('SELECT snapshot_json FROM ANALYSIS WHERE analysis_id=?',(analysis['analysis_id'],)).fetchone()[0])
        exported=json.loads(connection.execute('SELECT snapshot_json FROM EXPORT_LOG WHERE export_id=?',(receipt['export_id'],)).fetchone()[0])
        assert frozen['conditions']['daily_work_min']==expected
        assert next(task for task in exported['tasks'] if task['task_id']==state['task']['task_id'])['conditions']['daily_work_min']==expected
        assert text in json.dumps(exported,ensure_ascii=False)
    book=load_workbook(workbook,read_only=True,data_only=False)
    try:
        rows=[list(row) for row in book['입력한 값'].iter_rows(values_only=True) if '하루 실제 작업시간' in row]
        assert rows and any(any(re.fullmatch(re.escape(str(expected))+r'\s*분/일',str(value)) for value in row) for row in rows),rows
    finally:book.close()
    result=dict(case=name,expected_minutes=expected,actual_minutes=saved['value'],source=saved['source'],
        statement=text,calculation=saved.get('calculation'),ai=state['ai'],dropped=state['dropped'],
        seconds=round(time.monotonic()-started,3),database=str(db),workbook=str(workbook),
        workbook_sha256=hashlib.sha256(workbook.read_bytes()).hexdigest(),workbook_daily_rows=rows,
        no_duplicate_daily_question=True,restart=True,analysis_snapshot=True,export_snapshot=True,tables=tables)
    (directory/'receipt.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return result


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--directory',type=Path,default=ROOT/'.local'/('duration-grounding-'+str(time.time_ns())))
    parser.add_argument('--workers',type=int,default=2)
    parser.add_argument('--case',action='append',choices=[case[0] for case in CASES])
    args=parser.parse_args()
    if os.environ.get('LOCO_LIVE_AI')!='1' or not os.environ.get('OPENAI_API_KEY'):
        raise SystemExit('Set LOCO_LIVE_AI=1 and OPENAI_API_KEY for this live test.')
    os.environ.pop('OPENAI_BASE_URL',None)
    folder=args.directory.resolve();folder.mkdir(parents=True,exist_ok=False)
    defaults=ai_config.defaults()
    result=dict(kind='actual external AI/service/SQLite/analysis/XLSX, synthetic statements; no Electron, pose or OS input',
        source_base=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),
        source_dirty=True,observed_at=datetime.now(timezone.utc).isoformat(),cases=[],errors=[])
    original=ai_client.adaptive_intake
    def observe(payload,schema,model,key,settings=None):
        raw,reason=original(payload,schema,model,key,settings)
        TRACE.requests+=1
        trace=dict(phase=payload['phase'],response=raw,reason=reason)
        (TRACE.directory/('ai-request-'+str(TRACE.requests)+'.json')).write_text(json.dumps(trace,ensure_ascii=False,indent=2),encoding='utf-8')
        return raw,reason
    # Observe actual provider responses without replacing or changing them.
    ai_client.adaptive_intake=observe
    try:
        with ThreadPoolExecutor(max_workers=args.workers) as pool:
            pending={pool.submit(verify,case,folder,defaults):case[0] for case in CASES if not args.case or case[0] in args.case}
            for future in as_completed(pending):
                name=pending[future]
                try:
                    row=future.result();result['cases'].append(row)
                    print('PASS '+name+' -> '+str(row['actual_minutes'])+' minutes -> restart/snapshots/Excel',flush=True)
                except Exception as error:
                    result['errors'].append(dict(case=name,type=type(error).__name__,message=str(error)))
                    print('FAIL '+name+' '+str(error)[:500],flush=True)
    finally:ai_client.adaptive_intake=original
    (folder/'receipt.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(dict(passed=len(result['cases']),failed=len(result['errors']),receipt=str(folder/'receipt.json')),ensure_ascii=False),flush=True)
    return bool(result['errors'])


if __name__=='__main__':
    raise SystemExit(main())
